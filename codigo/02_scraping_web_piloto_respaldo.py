# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Fecha de extracción: prueba piloto con meses de control

# ============================================================
# 02_scraping_web.py
#
# FINANZAS I - UNIDAD I
#
# PILOTO SBS:
# 1. Localiza los XLS oficiales.
# 2. Comprueba los 6 meses de control.
# 3. Descarga solamente los 12 archivos piloto.
# 4. Conserva exactamente los bytes publicados por SBS.
# 5. Detecta el formato REAL mediante firma binaria:
#       XLS/OLE  -> xlrd
#       XLSX/ZIP -> openpyxl
# 6. Lee cada archivo con la librería correspondiente.
# 7. Muestra estructura y primeras filas.
#
# IMPORTANTE:
# - NO convierte archivos.
# - NO modifica archivos crudos.
# - NO cambia su extensión.
# - NO descarga todavía los 211 meses.
# ============================================================


from pathlib import Path
from urllib.parse import urljoin
import time

import requests
from bs4 import BeautifulSoup
import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. CONFIGURACIÓN GENERAL
# ============================================================

URL_REPORTE = (
    "https://www.sbs.gob.pe/app/stats_net/stats/"
    "EstadisticaSistemaFinancieroResultados.aspx?c={codigo}"
)


REPORTES = {

    "creditos": {
        "codigo": "B-2359",
        "nombre": (
            "Créditos Directos por Tipo, "
            "Modalidad y Moneda"
        ),
    },

    "depositos": {
        "codigo": "B-2318",
        "nombre": "Movimiento de los Depósitos",
    },
}


# ============================================================
# 2. MESES DE CONTROL
# ============================================================

MESES_CONTROL = [
    (2008, 6),     # junio 2008
    (2010, 12),    # diciembre 2010
    (2011, 12),    # diciembre 2011
    (2015, 6),     # junio 2015
    (2020, 6),     # junio 2020
    (2025, 12),    # diciembre 2025
]


CODIGO_MES = {
    1: "en",
    2: "fe",
    3: "ma",
    4: "ab",
    5: "my",
    6: "jn",
    7: "jl",
    8: "ag",
    9: "se",
    10: "oc",
    11: "no",
    12: "di",
}


NOMBRE_MES = {
    1: "enero",
    2: "febrero",
    3: "marzo",
    4: "abril",
    5: "mayo",
    6: "junio",
    7: "julio",
    8: "agosto",
    9: "setiembre",
    10: "octubre",
    11: "noviembre",
    12: "diciembre",
}


# ============================================================
# 3. CARPETAS
# ============================================================

RAIZ_PROYECTO = Path(
    __file__
).resolve().parent.parent


CARPETA_PILOTO = (
    RAIZ_PROYECTO
    / "datos_crudos"
    / "sbs_piloto"
)


CARPETA_PILOTO.mkdir(
    parents=True,
    exist_ok=True
)


# ============================================================
# 4. SESIÓN HTTP
# ============================================================

session = requests.Session()


session.headers.update({
    "User-Agent": (
        "Mozilla/5.0 "
        "(compatible; TrabajoAcademicoUNCP/1.0; "
        "Finanzas-I)"
    )
})


# ============================================================
# 5. CONSULTAR PÁGINA SBS
# ============================================================

def consultar_reporte(codigo):

    url = URL_REPORTE.format(
        codigo=codigo
    )

    respuesta = session.get(
        url,
        timeout=30
    )

    respuesta.raise_for_status()

    soup = BeautifulSoup(
        respuesta.text,
        "lxml"
    )

    return (
        url,
        soup,
        respuesta.status_code
    )


# ============================================================
# 6. OBTENER TEXTO DE LA PÁGINA
# ============================================================

def obtener_texto_pagina(soup):

    return soup.get_text(
        " ",
        strip=True
    )


# ============================================================
# 7. LOCALIZAR ENLACES XLS/XLSX
# ============================================================

def localizar_enlaces_xls(
    url_pagina,
    soup
):

    enlaces_xls = []

    for enlace in soup.find_all(
        "a",
        href=True
    ):

        href = enlace.get(
            "href",
            ""
        ).strip()

        href_minuscula = href.lower()

        # SBS puede publicar contenido XLSX
        # detrás de enlaces terminados en .xls.
        #
        # También aceptamos .xlsx si aparece
        # explícitamente en algún periodo.
        if (
            ".xls" not in href_minuscula
            and ".xlsx" not in href_minuscula
        ):
            continue

        url_xls = urljoin(
            url_pagina,
            href
        )

        texto = enlace.get_text(
            " ",
            strip=True
        )

        enlaces_xls.append({
            "texto": texto,
            "href_original": href,
            "url": url_xls,
        })

    return enlaces_xls


# ============================================================
# 8. IDENTIFICADOR DE MES
# ============================================================

def identificador_mes(
    codigo_reporte,
    year,
    month
):

    codigo_mes = CODIGO_MES[
        month
    ]

    return (
        f"{codigo_reporte}-"
        f"{codigo_mes}{year}"
    ).lower()


# ============================================================
# 9. BUSCAR MES DE CONTROL
# ============================================================

def buscar_mes_control(
    enlaces,
    codigo_reporte,
    year,
    month
):

    esperado = identificador_mes(
        codigo_reporte,
        year,
        month
    )

    coincidencias = []

    for item in enlaces:

        contenido = (
            item["url"]
            + " "
            + item["href_original"]
        ).lower()

        if esperado in contenido:

            coincidencias.append(
                item
            )

    return coincidencias


# ============================================================
# 10. NOMBRE LOCAL DEL ARCHIVO
# ============================================================

def crear_nombre_archivo(
    clave_reporte,
    year,
    month
):
    """
    Conservamos la extensión .xls porque corresponde
    al nombre/enlace publicado por SBS en estos reportes.

    NO significa que asumamos que internamente sea XLS.
    El formato real se detectará por firma binaria.
    """

    return (
        f"{clave_reporte}_"
        f"{year}_"
        f"{month:02d}.xls"
    )


# ============================================================
# 11. DESCARGAR ARCHIVO CRUDO
# ============================================================

def descargar_archivo(
    url,
    destino
):
    """
    Descarga los bytes exactamente como los entrega SBS.

    NO convierte.
    NO modifica.
    NO abre y vuelve a guardar.
    """

    print()
    print("Descargando:")
    print(url)

    respuesta = session.get(
        url,
        timeout=60
    )

    respuesta.raise_for_status()

    contenido = respuesta.content

    content_type = respuesta.headers.get(
        "Content-Type",
        ""
    )

    print(
        "HTTP:",
        respuesta.status_code
    )

    print(
        "Content-Type:",
        content_type
    )

    print(
        "Tamaño:",
        len(contenido),
        "bytes"
    )

    # --------------------------------------------------------
    # Protección contra páginas HTML
    # --------------------------------------------------------

    inicio = contenido[
        :500
    ].lower()

    if (
        b"<html" in inicio
        or b"<!doctype html" in inicio
    ):

        raise ValueError(
            "La SBS devolvió HTML en lugar "
            "de un archivo Excel."
        )

    # --------------------------------------------------------
    # Guardar EXACTAMENTE los bytes recibidos
    # --------------------------------------------------------

    with open(
        destino,
        "wb"
    ) as archivo:

        archivo.write(
            contenido
        )

    print(
        "Guardado sin modificar en:"
    )

    print(
        destino
    )

    # Pausa ética entre solicitudes
    time.sleep(1)

    return {
        "http": respuesta.status_code,
        "content_type": content_type,
        "bytes": len(contenido),
        "ruta": destino,
    }


# ============================================================
# 12. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato_excel(ruta):
    """
    Detecta el formato REAL mediante la firma binaria.

    XLS clásico:
        D0 CF 11 E0 A1 B1 1A E1
        -> contenedor OLE

    XLSX:
        comienza como archivo ZIP:
        PK...

    IMPORTANTE:
    La extensión del archivo NO se utiliza
    para decidir cómo leerlo.
    """

    with open(
        ruta,
        "rb"
    ) as archivo:

        cabecera = archivo.read(
            8
        )

    firma_hex = cabecera.hex()

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    # Los XLSX son archivos ZIP.
    # Las firmas ZIP comunes empiezan por:
    #
    # PK 03 04
    # PK 05 06
    # PK 07 08

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:

        formato = "XLS_OLE"

    elif cabecera.startswith(
        firmas_zip
    ):

        formato = "XLSX_ZIP"

    else:

        formato = "DESCONOCIDO"

    print()
    print(
        "Verificación:",
        ruta.name
    )

    print(
        "Firma hexadecimal:",
        firma_hex
    )

    print(
        "Formato real detectado:",
        formato
    )

    if formato == "XLS_OLE":

        print(
            "[OK] Archivo Excel XLS/OLE."
        )

    elif formato == "XLSX_ZIP":

        print(
            "[OK] Contenido Excel XLSX/ZIP "
            "detectado."
        )

        print(
            "     La extensión original se "
            "mantendrá sin cambios."
        )

    else:

        print(
            "[ADVERTENCIA] Formato binario "
            "no reconocido."
        )

    return formato


# ============================================================
# 13. LEER XLS/OLE CON XLRD
# ============================================================

def inspeccionar_xls_ole(
    ruta
):

    print()
    print("-" * 75)

    print(
        "LECTURA XLS/OLE:",
        ruta.name
    )

    try:

        libro = xlrd.open_workbook(
            filename=str(ruta),
            on_demand=True
        )

    except Exception as error:

        print(
            "[ERROR] xlrd no pudo "
            "abrir el archivo."
        )

        print(
            "Detalle:",
            error
        )

        return None

    nombres_hojas = libro.sheet_names()

    print(
        "Librería utilizada: xlrd"
    )

    print(
        "Número de hojas:",
        libro.nsheets
    )

    print(
        "Hojas:",
        nombres_hojas
    )

    resumen = []

    for nombre_hoja in nombres_hojas:

        hoja = libro.sheet_by_name(
            nombre_hoja
        )

        print()
        print(
            "Hoja:",
            nombre_hoja
        )

        print(
            "Filas:",
            hoja.nrows
        )

        print(
            "Columnas:",
            hoja.ncols
        )

        resumen.append({
            "hoja": nombre_hoja,
            "filas": hoja.nrows,
            "columnas": hoja.ncols,
        })

    libro.release_resources()

    return resumen


# ============================================================
# 14. LEER XLSX/ZIP CON OPENPYXL
# ============================================================

def inspeccionar_xlsx_zip(
    ruta
):
    """
    Abre el contenido XLSX directamente desde
    el archivo descargado.

    openpyxl normalmente revisa la extensión cuando
    recibe una ruta. Como SBS puede llamar .xls a un
    archivo cuyo contenido real es XLSX, pasamos un
    flujo binario abierto en lugar del nombre del archivo.

    Esto NO modifica el archivo original.
    """

    print()
    print("-" * 75)

    print(
        "LECTURA XLSX/ZIP:",
        ruta.name
    )

    archivo_binario = open(
        ruta,
        "rb"
    )

    try:

        libro = load_workbook(
            archivo_binario,
            read_only=True,
            data_only=True
        )

    except Exception as error:

        archivo_binario.close()

        print(
            "[ERROR] openpyxl no pudo "
            "abrir el archivo."
        )

        print(
            "Detalle:",
            error
        )

        return None

    print(
        "Librería utilizada: openpyxl"
    )

    print(
        "Número de hojas:",
        len(libro.sheetnames)
    )

    print(
        "Hojas:",
        libro.sheetnames
    )

    resumen = []

    for nombre_hoja in libro.sheetnames:

        hoja = libro[
            nombre_hoja
        ]

        print()
        print(
            "Hoja:",
            nombre_hoja
        )

        print(
            "Filas:",
            hoja.max_row
        )

        print(
            "Columnas:",
            hoja.max_column
        )

        resumen.append({
            "hoja": nombre_hoja,
            "filas": hoja.max_row,
            "columnas": hoja.max_column,
        })

    libro.close()
    archivo_binario.close()

    return resumen


# ============================================================
# 15. INSPECCIONAR SEGÚN FORMATO
# ============================================================

def inspeccionar_excel(
    ruta,
    formato
):

    if formato == "XLS_OLE":

        return inspeccionar_xls_ole(
            ruta
        )

    elif formato == "XLSX_ZIP":

        return inspeccionar_xlsx_zip(
            ruta
        )

    else:

        print()
        print(
            "[ERROR] No se intentará leer "
            "el archivo porque su formato "
            "no fue reconocido."
        )

        return None


# ============================================================
# 16. VISTA PREVIA XLS/OLE
# ============================================================

def vista_previa_xls_ole(
    ruta,
    max_filas=20,
    max_columnas=12
):

    libro = xlrd.open_workbook(
        filename=str(ruta),
        on_demand=True
    )

    for nombre_hoja in libro.sheet_names():

        hoja = libro.sheet_by_name(
            nombre_hoja
        )

        print()
        print(
            "HOJA:",
            nombre_hoja
        )

        print(
            f"Dimensión: "
            f"{hoja.nrows} filas x "
            f"{hoja.ncols} columnas"
        )

        print("-" * 75)

        limite_filas = min(
            max_filas,
            hoja.nrows
        )

        limite_columnas = min(
            max_columnas,
            hoja.ncols
        )

        for fila in range(
            limite_filas
        ):

            valores = []

            for columna in range(
                limite_columnas
            ):

                valor = hoja.cell_value(
                    fila,
                    columna
                )

                valores.append(
                    str(valor)
                )

            print(
                f"Fila {fila:02d}:",
                valores
            )

    libro.release_resources()


# ============================================================
# 17. VISTA PREVIA XLSX/ZIP
# ============================================================

def vista_previa_xlsx_zip(
    ruta,
    max_filas=20,
    max_columnas=12
):

    archivo_binario = open(
        ruta,
        "rb"
    )

    libro = load_workbook(
        archivo_binario,
        read_only=True,
        data_only=True
    )

    for nombre_hoja in libro.sheetnames:

        hoja = libro[
            nombre_hoja
        ]

        print()
        print(
            "HOJA:",
            nombre_hoja
        )

        print(
            f"Dimensión: "
            f"{hoja.max_row} filas x "
            f"{hoja.max_column} columnas"
        )

        print("-" * 75)

        limite_filas = min(
            max_filas,
            hoja.max_row
        )

        limite_columnas = min(
            max_columnas,
            hoja.max_column
        )

        # openpyxl usa índices desde 1
        for fila in range(
            1,
            limite_filas + 1
        ):

            valores = []

            for columna in range(
                1,
                limite_columnas + 1
            ):

                valor = hoja.cell(
                    row=fila,
                    column=columna
                ).value

                valores.append(
                    "" if valor is None
                    else str(valor)
                )

            # Mostramos índice desde 0 para que
            # visualmente sea comparable con xlrd.
            print(
                f"Fila {fila - 1:02d}:",
                valores
            )

    libro.close()
    archivo_binario.close()


# ============================================================
# 18. MOSTRAR VISTA PREVIA SEGÚN FORMATO
# ============================================================

def mostrar_vista_previa(
    ruta,
    formato,
    max_filas=20,
    max_columnas=12
):

    print()
    print("=" * 75)

    print(
        "VISTA PREVIA:",
        ruta.name
    )

    print(
        "Formato real:",
        formato
    )

    print("=" * 75)

    try:

        if formato == "XLS_OLE":

            vista_previa_xls_ole(
                ruta,
                max_filas=max_filas,
                max_columnas=max_columnas
            )

        elif formato == "XLSX_ZIP":

            vista_previa_xlsx_zip(
                ruta,
                max_filas=max_filas,
                max_columnas=max_columnas
            )

        else:

            print(
                "[ERROR] Formato no reconocido. "
                "No se mostrará vista previa."
            )

    except Exception as error:

        print(
            "[ERROR] No se pudo mostrar "
            "la vista previa."
        )

        print(
            "Detalle:",
            error
        )


# ============================================================
# 19. PROBAR UN REPORTE
# ============================================================

def probar_reporte(
    clave,
    configuracion
):
    """
    Mantiene la misma lógica de localización
    que ya produjo correctamente 12/12.
    """

    codigo = configuracion[
        "codigo"
    ]

    nombre_esperado = configuracion[
        "nombre"
    ]

    print()
    print("=" * 75)

    print(
        "REPORTE:",
        clave.upper()
    )

    print(
        "Código SBS:",
        codigo
    )

    print(
        "Nombre esperado:",
        nombre_esperado
    )

    print("=" * 75)

    try:

        url, soup, estado = consultar_reporte(
            codigo
        )

    except requests.RequestException as error:

        print()
        print(
            "ERROR AL CONSULTAR SBS"
        )

        print(
            error
        )

        return []

    print()
    print(
        "URL consultada:"
    )

    print(
        url
    )

    print()
    print(
        "Estado HTTP:",
        estado
    )

    # --------------------------------------------------------
    # Comprobar nombre
    # --------------------------------------------------------

    texto_pagina = obtener_texto_pagina(
        soup
    )

    palabras_clave = [
        palabra.lower()
        for palabra in nombre_esperado.split()
        if len(palabra) >= 5
    ]

    coincidencias_nombre = sum(
        palabra in texto_pagina.lower()
        for palabra in palabras_clave
    )

    if coincidencias_nombre > 0:

        print(
            "Nombre del reporte: "
            "aparentemente correcto."
        )

    else:

        print(
            "ADVERTENCIA: no se pudo confirmar "
            "el nombre esperado en el texto HTML."
        )

    # --------------------------------------------------------
    # Localizar enlaces
    # --------------------------------------------------------

    enlaces = localizar_enlaces_xls(
        url,
        soup
    )

    print()
    print(
        "Cantidad total de enlaces Excel encontrados:",
        len(enlaces)
    )

    if enlaces:

        print()
        print(
            "Primeros enlaces encontrados:"
        )

        for item in enlaces[:10]:

            print(
                "  ",
                item["texto"],
                "->",
                item["url"]
            )

    else:

        print()
        print(
            "ADVERTENCIA: no se encontraron "
            "enlaces Excel directamente en el HTML."
        )

    # --------------------------------------------------------
    # Buscar meses de control
    # --------------------------------------------------------

    print()
    print("-" * 75)

    print(
        "VERIFICACIÓN DE LOS 6 MESES DE CONTROL"
    )

    print("-" * 75)

    resultados = []

    for year, month in MESES_CONTROL:

        coincidencias = buscar_mes_control(
            enlaces,
            codigo,
            year,
            month
        )

        fecha_texto = (
            f"{NOMBRE_MES[month]} "
            f"{year}"
        )

        esperado = identificador_mes(
            codigo,
            year,
            month
        )

        if len(coincidencias) == 1:

            item = coincidencias[
                0
            ]

            print()
            print(
                f"[OK] {fecha_texto}"
            )

            print(
                "     Identificador:",
                esperado
            )

            print(
                "     URL:",
                item["url"]
            )

            resultados.append({
                "reporte": clave,
                "codigo": codigo,
                "year": year,
                "month": month,
                "estado": "OK",
                "url": item["url"],
            })

        elif len(coincidencias) == 0:

            print()
            print(
                f"[NO ENCONTRADO] "
                f"{fecha_texto}"
            )

            print(
                "     Se esperaba localizar:",
                esperado
            )

            resultados.append({
                "reporte": clave,
                "codigo": codigo,
                "year": year,
                "month": month,
                "estado": "NO ENCONTRADO",
                "url": None,
            })

        else:

            print()
            print(
                f"[MÚLTIPLES] "
                f"{fecha_texto}"
            )

            print(
                "     Número de coincidencias:",
                len(coincidencias)
            )

            for item in coincidencias:

                print(
                    "     ->",
                    item["url"]
                )

            resultados.append({
                "reporte": clave,
                "codigo": codigo,
                "year": year,
                "month": month,
                "estado": "MULTIPLES",
                "url": None,
            })

    return resultados


# ============================================================
# 20. DESCARGAR PILOTO
# ============================================================

def descargar_piloto(
    resultados
):

    archivos_descargados = []

    print()
    print("=" * 75)

    print(
        "ETAPA DE DESCARGA DEL PILOTO"
    )

    print("=" * 75)

    for resultado in resultados:

        if resultado[
            "estado"
        ] != "OK":

            continue

        nombre_archivo = crear_nombre_archivo(
            resultado["reporte"],
            resultado["year"],
            resultado["month"]
        )

        destino = (
            CARPETA_PILOTO
            / nombre_archivo
        )

        print()
        print("-" * 75)

        print(
            "Reporte:",
            resultado["reporte"]
        )

        print(
            "Fecha:",
            f"{resultado['year']}-"
            f"{resultado['month']:02d}"
        )

        # ----------------------------------------------------
        # Descargar exactamente como publica SBS
        # ----------------------------------------------------

        informacion = descargar_archivo(
            resultado["url"],
            destino
        )

        # ----------------------------------------------------
        # Detectar formato REAL
        # ----------------------------------------------------

        formato = detectar_formato_excel(
            destino
        )

        informacion[
            "formato"
        ] = formato

        informacion[
            "reporte"
        ] = resultado[
            "reporte"
        ]

        informacion[
            "year"
        ] = resultado[
            "year"
        ]

        informacion[
            "month"
        ] = resultado[
            "month"
        ]

        informacion[
            "url"
        ] = resultado[
            "url"
        ]

        archivos_descargados.append(
            informacion
        )

    return archivos_descargados


# ============================================================
# 21. FUNCIÓN PRINCIPAL
# ============================================================

def main():

    print()
    print("=" * 75)

    print(
        "PILOTO DE SCRAPING SBS"
    )

    print(
        "FINANZAS I - UNIDAD I"
    )

    print("=" * 75)

    print()
    print(
        "Los archivos se conservarán "
        "exactamente como los publica SBS."
    )

    print(
        "El formato real se detectará "
        "mediante la firma binaria."
    )

    print()
    print(
        "NO se descargarán todavía "
        "los 211 meses."
    )

    print()
    print(
        "Carpeta del proyecto:"
    )

    print(
        RAIZ_PROYECTO
    )

    print()
    print(
        "Carpeta de datos crudos:"
    )

    print(
        CARPETA_PILOTO
    )

    print()
    print(
        "Meses de control:"
    )

    for year, month in MESES_CONTROL:

        print(
            f" - "
            f"{NOMBRE_MES[month]} "
            f"{year}"
        )

    resultados_totales = []

    # --------------------------------------------------------
    # ETAPA 1: LOCALIZACIÓN
    # --------------------------------------------------------

    for clave, configuracion in REPORTES.items():

        resultados = probar_reporte(
            clave,
            configuracion
        )

        resultados_totales.extend(
            resultados
        )

    # --------------------------------------------------------
    # RESUMEN LOCALIZACIÓN
    # --------------------------------------------------------

    print()
    print("=" * 75)

    print(
        "RESUMEN DE LOCALIZACIÓN"
    )

    print("=" * 75)

    total_esperado = (
        len(REPORTES)
        * len(MESES_CONTROL)
    )

    encontrados = sum(
        resultado["estado"] == "OK"
        for resultado in resultados_totales
    )

    no_encontrados = sum(
        resultado["estado"] == "NO ENCONTRADO"
        for resultado in resultados_totales
    )

    multiples = sum(
        resultado["estado"] == "MULTIPLES"
        for resultado in resultados_totales
    )

    print()
    print(
        "Archivos esperados:",
        total_esperado
    )

    print(
        "Localizados correctamente:",
        encontrados
    )

    print(
        "No encontrados:",
        no_encontrados
    )

    print(
        "Con múltiples coincidencias:",
        multiples
    )

    # ========================================================
    # SOLO CONTINUAR SI SIGUE DANDO 12/12
    # ========================================================

    if encontrados != total_esperado:

        print()
        print(
            "[ERROR] No se localizaron "
            "los 12 archivos."
        )

        print(
            "NO se realizará ninguna descarga."
        )

        return

    print()
    print(
        "[OK] Los 12 archivos fueron localizados."
    )

    # ========================================================
    # ETAPA 2: DESCARGA
    # ========================================================

    try:

        archivos = descargar_piloto(
            resultados_totales
        )

    except Exception as error:

        print()
        print("=" * 75)

        print(
            "ERROR DURANTE LA DESCARGA"
        )

        print("=" * 75)

        print(
            error
        )

        print()
        print(
            "El programa se detendrá."
        )

        return

    # ========================================================
    # ETAPA 3: RESUMEN DE FORMATOS REALES
    # ========================================================

    print()
    print("=" * 75)

    print(
        "FORMATOS REALES DETECTADOS"
    )

    print("=" * 75)

    cantidad_ole = sum(
        archivo["formato"] == "XLS_OLE"
        for archivo in archivos
    )

    cantidad_xlsx = sum(
        archivo["formato"] == "XLSX_ZIP"
        for archivo in archivos
    )

    cantidad_desconocidos = sum(
        archivo["formato"] == "DESCONOCIDO"
        for archivo in archivos
    )

    print()
    print(
        "Total descargados:",
        len(archivos)
    )

    print(
        "XLS/OLE:",
        cantidad_ole
    )

    print(
        "XLSX/ZIP:",
        cantidad_xlsx
    )

    print(
        "Desconocidos:",
        cantidad_desconocidos
    )

    print()
    print(
        "Detalle:"
    )

    for archivo in archivos:

        print(
            f" - {archivo['ruta'].name}: "
            f"{archivo['formato']}"
        )

    # ========================================================
    # ETAPA 4: LECTURA CON LIBRERÍA ADECUADA
    # ========================================================

    print()
    print("=" * 75)

    print(
        "PRUEBA DE LECTURA"
    )

    print("=" * 75)

    lecturas_correctas = 0

    for archivo in archivos:

        resultado_lectura = inspeccionar_excel(
            archivo["ruta"],
            archivo["formato"]
        )

        if resultado_lectura is not None:

            lecturas_correctas += 1

    print()
    print("=" * 75)

    print(
        "RESUMEN DE LECTURA"
    )

    print("=" * 75)

    print()
    print(
        "Archivos abiertos correctamente:",
        lecturas_correctas,
        "de",
        len(archivos)
    )

    # ========================================================
    # ETAPA 5: VISTAS PREVIAS
    # ========================================================

    if (
        lecturas_correctas
        == len(archivos)
    ):

        print()
        print("=" * 75)

        print(
            "VISTAS PREVIAS"
        )

        print("=" * 75)

        print()
        print(
            "Se mostrarán las primeras "
            "20 filas y 12 columnas."
        )

        print(
            "Los archivos crudos NO serán "
            "modificados."
        )

        for archivo in archivos:

            mostrar_vista_previa(
                archivo["ruta"],
                archivo["formato"],
                max_filas=20,
                max_columnas=12
            )

    else:

        print()
        print(
            "[ADVERTENCIA]"
        )

        print(
            "No todos los archivos pudieron "
            "abrirse correctamente."
        )

    # ========================================================
    # RESULTADO FINAL
    # ========================================================

    print()
    print("=" * 75)

    print(
        "RESULTADO FINAL DEL PILOTO"
    )

    print("=" * 75)

    print()

    if (
        len(archivos) == 12
        and cantidad_desconocidos == 0
        and lecturas_correctas == 12
    ):

        print(
            "[OK] PILOTO COMPLETADO CORRECTAMENTE"
        )

        print()
        print(
            "1. Los 12 enlaces fueron localizados."
        )

        print(
            "2. Los 12 archivos fueron descargados."
        )

        print(
            "3. Todos los formatos reales "
            "fueron reconocidos."
        )

        print(
            "4. XLS/OLE se leyó con xlrd."
        )

        print(
            "5. XLSX/ZIP se leyó con openpyxl."
        )

        print(
            "6. Los archivos originales "
            "permanecieron sin modificar."
        )

        print()
        print(
            "SIGUIENTE PASO:"
        )

        print(
            "Comparar las estructuras de los "
            "reportes entre 2008 y 2025 para "
            "construir el parser."
        )

        print()
        print(
            "Todavía NO se ejecutará la "
            "descarga de los 211 meses."
        )

    else:

        print(
            "[ADVERTENCIA]"
        )

        print()
        print(
            "El piloto todavía presenta "
            "algún archivo problemático."
        )

        print(
            "NO debemos pasar aún a "
            "los 211 meses."
        )


# ============================================================
# 22. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":
    main()