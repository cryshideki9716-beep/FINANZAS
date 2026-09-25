# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Fecha de extracción: 2026-09-24

from pathlib import Path
from urllib.parse import urljoin, urlparse
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from datetime import datetime
import csv
import hashlib
import math
import re
import time
import unicodedata

import requests
from bs4 import BeautifulSoup
import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. CONFIGURACIÓN GENERAL
# ============================================================

CODIGO_ESTUDIANTE = "2024200485D"

ANIO_INICIO = 2015
MES_INICIO = 11
ANIO_FIN = 2025
MES_FIN = 12

PAUSA_DESCARGAS = 1.2
TIMEOUT_HTTP = 60
MAX_REINTENTOS_HTTP = 1
REANUDAR = True

# X1 conserva exactamente el control ya validado.
REL_TOL_X1 = 1e-14
ABS_TOL_X1 = 1e-6

URL_REPORTE = (
    "https://www.sbs.gob.pe/app/stats_net/stats/"
    "EstadisticaSistemaFinancieroResultados.aspx?c={codigo}"
)

REPORTES = {
    "Y": {
        "codigo": "B-2359",
        "nombre": "Créditos Directos por Tipo, Modalidad y Moneda",
    },
    "X1": {
        "codigo": "B-2318",
        "nombre": "Movimiento de los Depósitos",
    },
}

CODIGO_MES_SBS = {
    1: "en",
    2: "fe",
    3: "ma",
    4: "ab",
    5: "my",
    6: "jn",
    7: "jl",
    8: "ag",
    9: "se",
    10: "oc",
    11: "no",
    12: "di",
}


# ============================================================
# 2. RUTAS
# ============================================================

RAIZ = Path(__file__).resolve().parent.parent

CARPETA_DATOS_PROCESADOS = RAIZ / "datos_procesados"
CARPETA_DIAGNOSTICOS = RAIZ / "salidas" / "diagnosticos"

CARPETA_CRUDOS_Y = (
    RAIZ
    / "datos_crudos"
    / f"Y_B2359_{CODIGO_ESTUDIANTE}"
)

CARPETA_CRUDOS_X1 = (
    RAIZ
    / "datos_crudos"
    / f"X1_B2318_{CODIGO_ESTUDIANTE}"
)

CSV_Y = (
    CARPETA_DATOS_PROCESADOS
    / f"Y_dolarizacion_credito_{CODIGO_ESTUDIANTE}.csv"
)

CSV_X1 = (
    CARPETA_DATOS_PROCESADOS
    / f"X1_dolarizacion_depositos_{CODIGO_ESTUDIANTE}.csv"
)

CSV_AUDITORIA_Y = (
    CARPETA_DIAGNOSTICOS
    / f"auditoria_Y_dolarizacion_credito_{CODIGO_ESTUDIANTE}.csv"
)

CSV_AUDITORIA_X1 = (
    CARPETA_DIAGNOSTICOS
    / f"auditoria_X1_dolarizacion_depositos_{CODIGO_ESTUDIANTE}.csv"
)

TXT_RESUMEN = (
    CARPETA_DIAGNOSTICOS
    / f"resumen_Y_X1_produccion_{CODIGO_ESTUDIANTE}.txt"
)

for carpeta in (
    CARPETA_DATOS_PROCESADOS,
    CARPETA_DIAGNOSTICOS,
    CARPETA_CRUDOS_Y,
    CARPETA_CRUDOS_X1,
):
    carpeta.mkdir(
        parents=True,
        exist_ok=True,
    )


# ============================================================
# 3. COLUMNAS DE SALIDA
# ============================================================

COLUMNAS_Y = [
    "mes",
    "banco_original",
    "mn_soles_miles",
    "me_usd_miles",
    "total_soles_miles",
    "dolcred_pct",
    "problemas",
]

COLUMNAS_X1 = [
    "mes",
    "banco_original",
    "banco_original_me",
    "saldo_final_mn",
    "saldo_final_me",
    "doldep_pct",
    "problemas",
]

COLUMNAS_AUDITORIA_Y = [
    "mes",
    "estado",
    "url_fuente",
    "archivo_origen",
    "archivo_local",
    "sha256",
    "formato_real",
    "hoja",
    "fila_total_creditos",
    "cantidad_bancos",
    "cantidad_dolcred_validos",
    "cantidad_faltantes",

    # NUEVO:
    # Marca explícita de la regla individual vigente para
    # comparar Total vs MN y calcular DOLCRED.
    "regla_float_total_mn",

    "control_mn_estado",
    "control_mn_diferencia_absoluta_raw",
    "control_mn_diferencia_relativa_raw",
    "control_mn_formato_excel",
    "control_mn_decimales_mostrados",
    "control_mn_unidad_visible",
    "control_mn_tolerancia_fuente",
    "control_mn_suma_redondeada_half_up",
    "control_mn_total_redondeado_half_up",
    "control_me_estado",
    "control_me_diferencia_absoluta_raw",
    "control_me_diferencia_relativa_raw",
    "control_me_formato_excel",
    "control_me_decimales_mostrados",
    "control_me_unidad_visible",
    "control_me_tolerancia_fuente",
    "control_me_suma_redondeada_half_up",
    "control_me_total_redondeado_half_up",
    "control_total_estado",
    "control_total_diferencia_absoluta_raw",
    "control_total_diferencia_relativa_raw",
    "control_total_formato_excel",
    "control_total_decimales_mostrados",
    "control_total_unidad_visible",
    "control_total_tolerancia_fuente",
    "control_total_suma_redondeada_half_up",
    "control_total_total_redondeado_half_up",
    "fecha_hora_extraccion",
    "detalle",
]

COLUMNAS_AUDITORIA_X1 = [
    "mes",
    "estado",
    "url_fuente",
    "archivo_origen",
    "archivo_local",
    "sha256",
    "formato_real",
    "hoja_mn",
    "hoja_me",
    "cantidad_bancos",
    "cantidad_faltantes",
    "control_mn_estado",
    "control_mn_suma_bancos",
    "control_mn_total_sbs",
    "control_mn_diferencia_absoluta",
    "control_me_estado",
    "control_me_suma_bancos",
    "control_me_total_sbs",
    "control_me_diferencia_absoluta",
    "fecha_hora_extraccion",
    "detalle",
]


# ============================================================
# 4. EXCEPCIONES
# ============================================================

class ErrorTecnicoProduccion(Exception):
    pass


class ErrorEstructuralProduccion(Exception):
    pass


class MesSinDatoOficial(Exception):
    pass


class DatosFaltantesProduccion(Exception):
    pass


# ============================================================
# 5. SESIÓN HTTP
# ============================================================

session = requests.Session()

session.headers.update({
    "User-Agent": (
        "Mozilla/5.0 "
        "(compatible; TrabajoAcademicoUNCP/1.0; "
        "Finanzas-I; 2024200485D)"
    )
})


# ============================================================
# 6. UTILIDADES DE FECHA / MES
# ============================================================

def clave_mes(year, month):
    return f"{year:04d}-{month:02d}"


def generar_meses():

    meses = []

    year = ANIO_INICIO
    month = MES_INICIO

    while (
        year,
        month
    ) <= (
        ANIO_FIN,
        MES_FIN
    ):

        meses.append(
            (
                year,
                month,
            )
        )

        if month == 12:

            year += 1
            month = 1

        else:

            month += 1

    return meses


MESES_OBJETIVO = generar_meses()

if len(MESES_OBJETIVO) != 122:

    raise RuntimeError(
        "El período 2015-11 a 2025-12 "
        "debe contener exactamente 122 meses; "
        f"se obtuvieron {len(MESES_OBJETIVO)}."
    )


MESES_ESPERADOS = [
    clave_mes(
        year,
        month,
    )
    for year, month
    in MESES_OBJETIVO
]


def identificador_mes_sbs(
    codigo_reporte,
    year,
    month,
):

    return (
        f"{codigo_reporte}-"
        f"{CODIGO_MES_SBS[month]}"
        f"{year}"
    ).lower()


# ============================================================
# 7. UTILIDADES DE TEXTO
# ============================================================

def limpiar_texto(valor):

    if valor is None:
        return ""

    texto = str(
        valor
    )

    texto = texto.replace(
        "\xa0",
        " ",
    )

    texto = texto.replace(
        "\r",
        " ",
    )

    texto = texto.replace(
        "\n",
        " ",
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return texto.strip()


def normalizar_texto(valor):

    return limpiar_texto(
        valor
    ).lower()


def normalizar_comparacion(valor):

    texto = normalizar_texto(
        valor
    )

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(
            caracter
        )
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return texto.strip()


def sin_tildes(valor):

    return normalizar_comparacion(
        valor
    )


# ============================================================
# 8. CONVERSIÓN NUMÉRICA
# ============================================================

def convertir_numero(valor):

    if valor is None:
        return None

    if isinstance(
        valor,
        (
            int,
            float,
        )
    ):

        return float(
            valor
        )

    texto = str(
        valor
    ).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "n.d",
        "nd",
        "n/a",
        "na",
    }:

        return None

    try:

        return float(
            texto.replace(
                ",",
                "",
            )
        )

    except ValueError:

        return None


# ============================================================
# 9. CSV ATÓMICO / CHECKPOINT
# ============================================================

def leer_csv(ruta):

    if not ruta.exists():
        return []

    with ruta.open(
        "r",
        newline="",
        encoding="utf-8-sig",
    ) as archivo:

        return list(
            csv.DictReader(
                archivo
            )
        )


def escribir_csv_atomico(
    ruta,
    columnas,
    filas,
):

    ruta.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporal = ruta.with_suffix(
        ruta.suffix
        +
        ".tmp"
    )

    with temporal.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as archivo:

        escritor = csv.DictWriter(
            archivo,
            fieldnames=columnas,
            extrasaction="ignore",
        )

        escritor.writeheader()

        for fila in filas:

            escritor.writerow(
                fila
            )

    temporal.replace(
        ruta
    )


def cargar_estado():

    y_por_mes = defaultdict(
        list
    )

    for fila in leer_csv(
        CSV_Y
    ):

        mes = limpiar_texto(
            fila.get(
                "mes",
                "",
            )
        )

        if mes:

            y_por_mes[
                mes
            ].append(
                fila
            )

    x1_por_mes = defaultdict(
        list
    )

    for fila in leer_csv(
        CSV_X1
    ):

        mes = limpiar_texto(
            fila.get(
                "mes",
                "",
            )
        )

        if mes:

            x1_por_mes[
                mes
            ].append(
                fila
            )

    auditoria_y = {}

    for fila in leer_csv(
        CSV_AUDITORIA_Y
    ):

        mes = limpiar_texto(
            fila.get(
                "mes",
                "",
            )
        )

        if mes:

            auditoria_y[
                mes
            ] = fila

    auditoria_x1 = {}

    for fila in leer_csv(
        CSV_AUDITORIA_X1
    ):

        mes = limpiar_texto(
            fila.get(
                "mes",
                "",
            )
        )

        if mes:

            auditoria_x1[
                mes
            ] = fila

    return (
        y_por_mes,
        x1_por_mes,
        auditoria_y,
        auditoria_x1,
    )


def guardar_estado(
    y_por_mes,
    x1_por_mes,
    auditoria_y,
    auditoria_x1,
):

    filas_y = []

    for mes in sorted(
        y_por_mes
    ):

        filas_y.extend(
            y_por_mes[
                mes
            ]
        )

    filas_x1 = []

    for mes in sorted(
        x1_por_mes
    ):

        filas_x1.extend(
            x1_por_mes[
                mes
            ]
        )

    filas_aud_y = [
        auditoria_y[
            mes
        ]
        for mes
        in sorted(
            auditoria_y
        )
    ]

    filas_aud_x1 = [
        auditoria_x1[
            mes
        ]
        for mes
        in sorted(
            auditoria_x1
        )
    ]

    escribir_csv_atomico(
        CSV_Y,
        COLUMNAS_Y,
        filas_y,
    )

    escribir_csv_atomico(
        CSV_X1,
        COLUMNAS_X1,
        filas_x1,
    )

    escribir_csv_atomico(
        CSV_AUDITORIA_Y,
        COLUMNAS_AUDITORIA_Y,
        filas_aud_y,
    )

    escribir_csv_atomico(
        CSV_AUDITORIA_X1,
        COLUMNAS_AUDITORIA_X1,
        filas_aud_x1,
    )


# ============================================================
# 10. CHECKPOINT POR VARIABLE
# ============================================================

def _entero_seguro(valor):

    try:

        return int(
            str(
                valor
            ).strip()
        )

    except Exception:

        return None


def checkpoint_y_completo(
    mes,
    y_por_mes,
    auditoria_y,
):

    audit = auditoria_y.get(
        mes
    )

    if audit is None:
        return False

    estado = limpiar_texto(
        audit.get(
            "estado",
            "",
        )
    )

    filas = y_por_mes.get(
        mes,
        [],
    )

    # Solo DISPONIBLE puede ser checkpoint definitivo.
    if estado in {
        "ERROR_TECNICO",
        "ERROR_ESTRUCTURAL",
    }:

        return False

    if estado == "MES_SIN_DATO":
        return False

    if estado != "DISPONIBLE":
        return False

    cantidad = _entero_seguro(
        audit.get(
            "cantidad_bancos",
            "",
        )
    )

    if (
        cantidad is None
        or
        cantidad <= 0
    ):

        return False

    if len(
        filas
    ) != cantidad:

        return False

    if any(
        limpiar_texto(
            fila.get(
                "mes",
                "",
            )
        )
        !=
        mes
        for fila
        in filas
    ):

        return False

    claves = [
        (
            limpiar_texto(
                fila.get(
                    "mes",
                    "",
                )
            ),
            limpiar_texto(
                fila.get(
                    "banco_original",
                    "",
                )
            ),
        )
        for fila
        in filas
    ]

    if len(
        claves
    ) != len(
        set(
            claves
        )
    ):

        return False

    if any(
        limpiar_texto(
            audit.get(
                campo,
                "",
            )
        )
        !=
        "OK"
        for campo in (
            "control_mn_estado",
            "control_me_estado",
            "control_total_estado",
        )
    ):

        return False

    # --------------------------------------------------------
    # Consistencia metodológica del control agregado Y.
    #
    # Un checkpoint DISPONIBLE solo es definitivo si contiene
    # las evidencias de la regla actual:
    #
    # abs(suma_fsum - total_sbs) <= unidad_visible / 2
    # --------------------------------------------------------

    evidencias_nueva_regla = (
        "control_mn_unidad_visible",
        "control_mn_tolerancia_fuente",
        "control_me_unidad_visible",
        "control_me_tolerancia_fuente",
        "control_total_unidad_visible",
        "control_total_tolerancia_fuente",
    )

    for campo in evidencias_nueva_regla:

        valor = limpiar_texto(
            audit.get(
                campo,
                "",
            )
        )

        if valor == "":
            return False

    # --------------------------------------------------------
    # NUEVO CONTROL DE CONSISTENCIA DEL PARSER INDIVIDUAL Y.
    #
    # Todos los meses Y deben haber sido procesados usando
    # explícitamente la misma regla de representación float:
    #
    #   tolerancia_float = max(math.ulp(total), math.ulp(mn))
    #
    # Si abs(total - mn) <= tolerancia_float, la diferencia
    # efectiva usada para TOTAL_MENOR_QUE_MN y DOLCRED es 0.0.
    #
    # Los checkpoints creados antes de esta regla no tienen
    # esta marca y deben reprocesarse desde el crudo local.
    # --------------------------------------------------------

    if (
        limpiar_texto(
            audit.get(
                "regla_float_total_mn",
                "",
            )
        )
        !=
        "ULP_1"
    ):

        return False

    return True


def checkpoint_x1_completo(
    mes,
    x1_por_mes,
    auditoria_x1,
):

    audit = auditoria_x1.get(
        mes
    )

    if audit is None:
        return False

    estado = limpiar_texto(
        audit.get(
            "estado",
            "",
        )
    )

    filas = x1_por_mes.get(
        mes,
        [],
    )

    if estado in {
        "ERROR_TECNICO",
        "ERROR_ESTRUCTURAL",
    }:

        return False

    if estado == "MES_SIN_DATO":
        return False

    if estado != "DISPONIBLE":
        return False

    cantidad = _entero_seguro(
        audit.get(
            "cantidad_bancos",
            "",
        )
    )

    if (
        cantidad is None
        or
        cantidad <= 0
    ):

        return False

    if len(
        filas
    ) != cantidad:

        return False

    if any(
        limpiar_texto(
            fila.get(
                "mes",
                "",
            )
        )
        !=
        mes
        for fila
        in filas
    ):

        return False

    claves = [
        (
            limpiar_texto(
                fila.get(
                    "mes",
                    "",
                )
            ),
            limpiar_texto(
                fila.get(
                    "banco_original",
                    "",
                )
            ),
        )
        for fila
        in filas
    ]

    if len(
        claves
    ) != len(
        set(
            claves
        )
    ):

        return False

    if (
        limpiar_texto(
            audit.get(
                "control_mn_estado",
                "",
            )
        )
        !=
        "OK"
    ):

        return False

    if (
        limpiar_texto(
            audit.get(
                "control_me_estado",
                "",
            )
        )
        !=
        "OK"
    ):

        return False

    return True


def checkpoint_mes_completo(
    mes,
    y_por_mes,
    x1_por_mes,
    auditoria_y,
    auditoria_x1,
):

    return (
        checkpoint_y_completo(
            mes,
            y_por_mes,
            auditoria_y,
        )
        and
        checkpoint_x1_completo(
            mes,
            x1_por_mes,
            auditoria_x1,
        )
    )


# ============================================================
# 11. HTTP CON UN REINTENTO TÉCNICO
# ============================================================

def get_con_reintento(
    url,
    timeout,
):

    errores = []

    for intento in range(
        1
        +
        MAX_REINTENTOS_HTTP
    ):

        try:

            respuesta = session.get(
                url,
                timeout=timeout,
            )

            respuesta.raise_for_status()

            return respuesta

        except requests.RequestException as error:

            errores.append(
                f"{type(error).__name__}: {error}"
            )

            if intento < MAX_REINTENTOS_HTTP:

                time.sleep(
                    2
                )

                continue

    raise ErrorTecnicoProduccion(
        "Fallo HTTP después del reintento permitido: "
        +
        " || ".join(
            errores
        )
    )


# ============================================================
# 12. MAPA DE ENLACES OFICIALES
# ============================================================

def cargar_enlaces_reporte(codigo):

    url = URL_REPORTE.format(
        codigo=codigo
    )

    respuesta = get_con_reintento(
        url,
        TIMEOUT_HTTP,
    )

    try:

        soup = BeautifulSoup(
            respuesta.text,
            "lxml",
        )

    except Exception as error:

        raise ErrorTecnicoProduccion(
            f"{codigo}: no pudo parsearse "
            "la página HTML del reporte."
        ) from error

    texto_pagina = sin_tildes(
        soup.get_text(
            " ",
            strip=True,
        )
    )

    if codigo == "B-2359":

        evidencias_requeridas = (
            "creditos",
            "modalidad",
            "moneda",
        )

    elif codigo == "B-2318":

        evidencias_requeridas = (
            "movimiento",
            "depositos",
        )

    else:

        raise ErrorEstructuralProduccion(
            f"{codigo}: código de reporte "
            "no contemplado por el control "
            "de identidad."
        )

    faltantes_identidad = [
        termino
        for termino
        in evidencias_requeridas
        if termino
        not in texto_pagina
    ]

    if faltantes_identidad:

        raise ErrorEstructuralProduccion(
            f"{codigo}: la página respondió HTTP 200, "
            "pero no contiene evidencia suficiente "
            "del reporte esperado. "
            "Faltan términos de identidad: "
            f"{faltantes_identidad}."
        )

    enlaces = []
    urls_vistas = set()

    for enlace in soup.find_all(
        "a",
        href=True,
    ):

        href = enlace.get(
            "href",
            "",
        ).strip()

        if ".xls" not in href.lower():
            continue

        url_xls = urljoin(
            url,
            href,
        )

        if url_xls in urls_vistas:
            continue

        urls_vistas.add(
            url_xls
        )

        enlaces.append({
            "href_original":
                href,

            "url":
                url_xls,

            "texto":
                enlace.get_text(
                    " ",
                    strip=True,
                ),
        })

    if not enlaces:

        raise ErrorEstructuralProduccion(
            f"{codigo}: la página cargó, "
            "pero no se detectaron enlaces XLS."
        )

    return (
        url,
        enlaces,
    )


def construir_mapa_mensual(
    codigo,
    enlaces,
):

    mapa = {}

    for year, month in MESES_OBJETIVO:

        mes = clave_mes(
            year,
            month,
        )

        esperado = identificador_mes_sbs(
            codigo,
            year,
            month,
        )

        coincidencias = []

        for item in enlaces:

            contenido = (
                item[
                    "url"
                ]
                +
                " "
                +
                item[
                    "href_original"
                ]
            ).lower()

            if esperado in contenido:

                coincidencias.append(
                    item
                )

        if len(
            coincidencias
        ) > 1:

            raise ErrorEstructuralProduccion(
                f"{codigo} {mes}: se encontraron "
                f"{len(coincidencias)} enlaces "
                "para el mismo mes."
            )

        mapa[
            mes
        ] = (
            coincidencias[
                0
            ]
            if coincidencias
            else None
        )

    return mapa


# ============================================================
# 13. DESCARGA CRUDA REANUDABLE
# ============================================================

def nombre_archivo_local(
    variable,
    codigo,
    mes,
):

    return (
        f"{variable}_{codigo}_"
        f"{CODIGO_ESTUDIANTE}_{mes}.xls"
    )


def _detectar_formato_binario_bytes(
    contenido
):

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    cabecera = contenido[
        :8
    ]

    if cabecera == firma_ole:

        return "XLS_OLE"

    if cabecera.startswith(
        firmas_zip
    ):

        return "XLSX_ZIP"

    return None


def descargar_o_reutilizar(
    variable,
    codigo,
    mes,
    enlace,
    carpeta,
):

    if enlace is None:

        raise MesSinDatoOficial(
            f"{codigo} {mes}: la página oficial "
            "cargó correctamente, pero no contiene "
            "un enlace exacto para ese mes."
        )

    destino = (
        carpeta
        /
        nombre_archivo_local(
            variable,
            codigo,
            mes,
        )
    )

    if (
        destino.exists()
        and
        destino.stat().st_size > 0
    ):

        try:

            contenido = destino.read_bytes()

        except OSError as error:

            raise ErrorTecnicoProduccion(
                f"{codigo} {mes}: no pudo leerse "
                "el archivo crudo existente."
            ) from error

        formato_existente = (
            _detectar_formato_binario_bytes(
                contenido
            )
        )

        if formato_existente is None:

            raise ErrorTecnicoProduccion(
                f"{codigo} {mes}: el archivo crudo "
                "existente tiene una firma binaria "
                "inválida; no se reutiliza como "
                "OLE ni ZIP/XLSX."
            )

        return {
            "url":
                enlace[
                    "url"
                ],

            "archivo_origen":
                Path(
                    urlparse(
                        enlace[
                            "url"
                        ]
                    ).path
                ).name,

            "archivo_local":
                str(
                    destino
                ),

            "sha256":
                hashlib.sha256(
                    contenido
                ).hexdigest(),

            "bytes":
                len(
                    contenido
                ),

            "reutilizado":
                True,
        }

    respuesta = get_con_reintento(
        enlace[
            "url"
        ],
        TIMEOUT_HTTP,
    )

    contenido = respuesta.content

    inicio = contenido[
        :300
    ].lower()

    if (
        b"<html" in inicio
        or
        b"<!doctype html" in inicio
    ):

        raise ErrorTecnicoProduccion(
            f"{codigo} {mes}: SBS devolvió HTML "
            "en vez del archivo Excel."
        )

    if len(
        contenido
    ) == 0:

        raise ErrorTecnicoProduccion(
            f"{codigo} {mes}: la descarga "
            "devolvió 0 bytes."
        )

    formato_descargado = (
        _detectar_formato_binario_bytes(
            contenido
        )
    )

    if formato_descargado is None:

        raise ErrorTecnicoProduccion(
            f"{codigo} {mes}: el contenido recién "
            "descargado no tiene firma OLE ni "
            "ZIP/XLSX válida; no se reemplazará "
            "el archivo definitivo."
        )

    temporal = destino.with_suffix(
        destino.suffix
        +
        ".tmp"
    )

    try:

        temporal.write_bytes(
            contenido
        )

        temporal.replace(
            destino
        )

    except OSError as error:

        raise ErrorTecnicoProduccion(
            f"{codigo} {mes}: no pudo guardarse "
            "el archivo crudo."
        ) from error

    time.sleep(
        PAUSA_DESCARGAS
    )

    return {
        "url":
            enlace[
                "url"
            ],

        "archivo_origen":
            Path(
                urlparse(
                    enlace[
                        "url"
                    ]
                ).path
            ).name,

        "archivo_local":
            str(
                destino
            ),

        "sha256":
            hashlib.sha256(
                contenido
            ).hexdigest(),

        "bytes":
            len(
                contenido
            ),

        "reutilizado":
            False,
    }


# ============================================================
# 14. DETECTAR FORMATO REAL OLE / ZIP
# ============================================================

def detectar_formato(ruta):

    with Path(
        ruta
    ).open(
        "rb"
    ) as archivo:

        cabecera = archivo.read(
            8
        )

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:

        return "XLS_OLE"

    if cabecera.startswith(
        firmas_zip
    ):

        return "XLSX_ZIP"

    raise ErrorEstructuralProduccion(
        f"Formato binario no reconocido: "
        f"{ruta}"
    )


# ============================================================
# 15. LECTURA DE EXCEL
# ============================================================

def leer_xls(ruta):

    try:

        libro = xlrd.open_workbook(
            filename=str(
                ruta
            ),
            on_demand=True,
        )

    except Exception as error:

        raise ErrorEstructuralProduccion(
            f"No pudo abrirse como XLS/OLE: "
            f"{ruta}"
        ) from error

    hojas = []

    try:

        for nombre in libro.sheet_names():

            hoja = libro.sheet_by_name(
                nombre
            )

            filas = [
                [
                    hoja.cell_value(
                        i,
                        j,
                    )
                    for j in range(
                        hoja.ncols
                    )
                ]
                for i in range(
                    hoja.nrows
                )
            ]

            hojas.append({
                "nombre":
                    nombre,

                "filas":
                    filas,
            })

    finally:

        libro.release_resources()

    return hojas


def leer_xlsx(ruta):

    try:

        with Path(
            ruta
        ).open(
            "rb"
        ) as archivo_binario:

            libro = load_workbook(
                archivo_binario,
                read_only=True,
                data_only=True,
            )

            hojas = []

            try:

                for nombre in libro.sheetnames:

                    hoja = libro[
                        nombre
                    ]

                    filas = [
                        list(
                            fila
                        )
                        for fila
                        in hoja.iter_rows(
                            values_only=True
                        )
                    ]

                    hojas.append({
                        "nombre":
                            nombre,

                        "filas":
                            filas,
                    })

            finally:

                libro.close()

            return hojas

    except ErrorEstructuralProduccion:

        raise

    except Exception as error:

        raise ErrorEstructuralProduccion(
            f"No pudo abrirse como XLSX/ZIP: "
            f"{ruta}"
        ) from error


def leer_archivo(ruta):

    formato = detectar_formato(
        ruta
    )

    if formato == "XLS_OLE":

        hojas = leer_xls(
            ruta
        )

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(
            ruta
        )

    else:

        raise ErrorEstructuralProduccion(
            f"Formato no soportado: "
            f"{formato}"
        )

    return (
        formato,
        hojas,
    )


# ============================================================
# 16. Y / B-2359 - PARSER VALIDADO
# ============================================================

def es_mn(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith(
            "mn"
        )
        and
        "mil" in texto
    )


def es_me(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith(
            "me"
        )
        and
        "mil" in texto
    )


def es_total(valor):

    return normalizar_texto(
        valor
    ).startswith(
        "total"
    )


def encontrar_fila_encabezados_creditos(
    filas
):

    candidatos = []

    for numero_fila, fila in enumerate(
        filas
    ):

        bloques = []

        for columna in range(
            max(
                0,
                len(
                    fila
                ) - 2,
            )
        ):

            if (
                es_mn(
                    fila[
                        columna
                    ]
                )
                and
                es_me(
                    fila[
                        columna + 1
                    ]
                )
                and
                es_total(
                    fila[
                        columna + 2
                    ]
                )
            ):

                bloques.append(
                    columna
                )

        if bloques:

            candidatos.append({
                "fila":
                    numero_fila,

                "bloques":
                    bloques,
            })

    if not candidatos:

        return None

    return max(
        candidatos,
        key=lambda x:
        len(
            x[
                "bloques"
            ]
        ),
    )


def obtener_nombre_banco_creditos(
    filas,
    fila_bancos,
    columna_inicio,
):

    if (
        fila_bancos < 0
        or
        fila_bancos >= len(
            filas
        )
    ):

        return ""

    fila = filas[
        fila_bancos
    ]

    for desplazamiento in (
        0,
        1,
        2,
    ):

        columna = (
            columna_inicio
            +
            desplazamiento
        )

        if columna >= len(
            fila
        ):

            continue

        valor = fila[
            columna
        ]

        if normalizar_texto(
            valor
        ):

            return str(
                valor
            ).strip()

    return ""


def detectar_estructura_creditos(
    filas
):

    encabezado = (
        encontrar_fila_encabezados_creditos(
            filas
        )
    )

    if encabezado is None:

        return None

    fila_encabezados = encabezado[
        "fila"
    ]

    fila_bancos = (
        fila_encabezados
        -
        1
    )

    bloques = []

    for columna_inicio in encabezado[
        "bloques"
    ]:

        bloques.append({
            "banco_original":
                obtener_nombre_banco_creditos(
                    filas,
                    fila_bancos,
                    columna_inicio,
                ),

            "columna_mn":
                columna_inicio,

            "columna_me":
                columna_inicio + 1,

            "columna_total":
                columna_inicio + 2,
        })

    return {
        "fila_bancos":
            fila_bancos,

        "fila_encabezados":
            fila_encabezados,

        "bloques":
            bloques,
    }


def elegir_hoja_principal_creditos(
    hojas
):

    candidatos = []

    for hoja in hojas:

        estructura = (
            detectar_estructura_creditos(
                hoja[
                    "filas"
                ]
            )
        )

        if estructura is not None:

            candidatos.append({
                "hoja":
                    hoja,

                "estructura":
                    estructura,
            })

    if not candidatos:

        return None

    return max(
        candidatos,
        key=lambda x:
        len(
            x[
                "estructura"
            ][
                "bloques"
            ]
        ),
    )


def buscar_fila_total_creditos(
    filas
):

    for numero_fila, fila in enumerate(
        filas
    ):

        for valor in fila:

            if (
                "total creditos"
                in
                sin_tildes(
                    valor
                )
            ):

                return numero_fila

    return None


def es_agregado_creditos(
    nombre
):

    return (
        "total banca multiple"
        in
        sin_tildes(
            nombre
        )
    )


def extraer_y_creditos(
    ruta
):

    formato, hojas = leer_archivo(
        ruta
    )

    seleccionado = (
        elegir_hoja_principal_creditos(
            hojas
        )
    )

    if seleccionado is None:

        raise ErrorEstructuralProduccion(
            "B-2359: no se encontró "
            "la estructura MN|ME|Total."
        )

    hoja = seleccionado[
        "hoja"
    ]

    estructura = seleccionado[
        "estructura"
    ]

    filas = hoja[
        "filas"
    ]

    fila_total = (
        buscar_fila_total_creditos(
            filas
        )
    )

    if fila_total is None:

        raise ErrorEstructuralProduccion(
            "B-2359: no se encontró "
            "la fila 'Total Créditos:'."
        )

    fila = filas[
        fila_total
    ]

    registros = []

    for bloque in estructura[
        "bloques"
    ]:

        banco = bloque[
            "banco_original"
        ]

        col_mn = bloque[
            "columna_mn"
        ]

        col_me = bloque[
            "columna_me"
        ]

        col_total = bloque[
            "columna_total"
        ]

        mn = (
            convertir_numero(
                fila[
                    col_mn
                ]
            )
            if col_mn < len(
                fila
            )
            else None
        )

        me = (
            convertir_numero(
                fila[
                    col_me
                ]
            )
            if col_me < len(
                fila
            )
            else None
        )

        total = (
            convertir_numero(
                fila[
                    col_total
                ]
            )
            if col_total < len(
                fila
            )
            else None
        )

        problemas = []

        if not banco:

            problemas.append(
                "BANCO_SIN_NOMBRE"
            )

        if mn is None:

            problemas.append(
                "MN_FALTANTE"
            )

        if me is None:

            problemas.append(
                "ME_FALTANTE"
            )

        if total is None:

            problemas.append(
                "TOTAL_FALTANTE"
            )

        if (
            mn is not None
            and
            mn < 0
        ):

            problemas.append(
                "MN_NEGATIVO"
            )

        if (
            me is not None
            and
            me < 0
        ):

            problemas.append(
                "ME_NEGATIVO"
            )

        if (
            total is not None
            and
            total <= 0
        ):

            problemas.append(
                "TOTAL_NO_POSITIVO"
            )

        dolcred = None
        tc_implicito = None

        # ====================================================
        # NORMALIZACIÓN EXCLUSIVAMENTE DE REPRESENTACIÓN FLOAT
        #
        # MN y Total son valores raw leídos directamente de la
        # misma fila del archivo SBS y convertidos a float.
        #
        # En algunos casos, dos valores que económicamente son
        # iguales pueden diferir por una sola unidad de precisión
        # binaria del float (1 ULP).
        #
        # Ejemplo Banco Azteca 2021-09:
        #
        # MN    = 325645.26300000004
        # Total = 325645.263
        #
        # Total - MN = -5.820766091346741e-11
        #
        # Esa diferencia equivale a 1 ULP a ese nivel.
        #
        # Esta normalización NO introduce una tolerancia económica,
        # NO utiliza 0.5, porcentajes ni umbrales arbitrarios.
        #
        # Únicamente elimina diferencias de hasta 1 ULP entre
        # los dos valores raw directos ya leídos del archivo SBS.
        # ====================================================

        diferencia_total_mn = None
        diferencia_total_mn_efectiva = None
        tolerancia_float = None

        if (
            mn is not None
            and
            total is not None
        ):

            diferencia_total_mn = (
                total
                -
                mn
            )

            tolerancia_float = max(
                math.ulp(
                    total
                ),
                math.ulp(
                    mn
                ),
            )

            if (
                abs(
                    diferencia_total_mn
                )
                <=
                tolerancia_float
            ):

                diferencia_total_mn_efectiva = 0.0

            else:

                diferencia_total_mn_efectiva = (
                    diferencia_total_mn
                )

            # TOTAL_MENOR_QUE_MN solo se activa cuando,
            # después de eliminar un posible ruido de hasta
            # 1 ULP, la diferencia continúa siendo negativa.
            if (
                diferencia_total_mn_efectiva
                <
                0
            ):

                problemas.append(
                    "TOTAL_MENOR_QUE_MN"
                )

        # ====================================================
        # DOLCRED
        #
        # Fórmula económica original:
        #
        #   (Total - MN) / Total * 100
        #
        # Se mantiene la misma fórmula. La única normalización
        # previa es la eliminación de ruido binario <= 1 ULP
        # en Total - MN.
        #
        # Por tanto, si Total y MN difieren únicamente por
        # 1 ULP, el numerador efectivo es exactamente 0.0.
        # ====================================================

        if (
            total is not None
            and
            total > 0
            and
            mn is not None
        ):

            dolcred = (
                diferencia_total_mn_efectiva
                /
                total
                *
                100
            )

            if not (
                0
                <=
                dolcred
                <=
                100
            ):

                problemas.append(
                    "DOLCRED_FUERA_RANGO"
                )

        # ====================================================
        # TC IMPLÍCITO
        #
        # Se conserva SIN CAMBIOS.
        # ====================================================

        if (
            mn is not None
            and
            me is not None
            and
            total is not None
            and
            me > 0
        ):

            tc_implicito = (
                (
                    total
                    -
                    mn
                )
                /
                me
            )

            if tc_implicito <= 0:

                problemas.append(
                    "TC_NO_POSITIVO"
                )

        registros.append({
            "banco_original":
                banco,

            "mn_soles_miles":
                mn,

            "me_usd_miles":
                me,

            "total_soles_miles":
                total,

            "dolcred_pct":
                dolcred,

            "tc_implicito":
                tc_implicito,

            "es_agregado":
                es_agregado_creditos(
                    banco
                ),

            "problemas":
                problemas,
        })

    return {
        "formato":
            formato,

        "hoja":
            hoja[
                "nombre"
            ],

        "fila_encabezados":
            estructura[
                "fila_encabezados"
            ],

        "fila_bancos":
            estructura[
                "fila_bancos"
            ],

        "fila_total_creditos":
            fila_total,

        "cantidad_bloques":
            len(
                estructura[
                    "bloques"
                ]
            ),

        "registros":
            registros,
    }


# ============================================================
# 17. Y - PRECISIÓN VISIBLE DEL AGREGADO
# ============================================================

def inferir_decimales_formato_excel(
    formato
):

    if formato is None:
        return None

    texto = str(
        formato
    ).strip()

    if (
        not texto
        or
        texto.lower()
        ==
        "general"
    ):

        return None

    seccion = texto.split(
        ";"
    )[
        0
    ]

    seccion = re.sub(
        r'"[^\"]*"',
        "",
        seccion,
    )

    seccion = re.sub(
        r"\[[^\]]*\]",
        "",
        seccion,
    )

    coincidencia = re.search(
        r"\.([0#?]+)",
        seccion,
    )

    if coincidencia:

        return len(
            coincidencia.group(
                1
            )
        )

    if re.search(
        r"[0#?]",
        seccion,
    ):

        return 0

    return None


def redondear_excel_half_up(
    valor,
    decimales,
):

    if (
        valor is None
        or
        decimales is None
        or
        decimales < 0
    ):

        raise ErrorEstructuralProduccion(
            "No puede aplicarse ROUND_HALF_UP "
            "sin valor y decimales válidos."
        )

    numero = Decimal(
        str(
            valor
        )
    )

    cuantizador = (
        Decimal(
            "1"
        )
        if decimales == 0
        else
        Decimal(
            "1"
        ).scaleb(
            -decimales
        )
    )

    return numero.quantize(
        cuantizador,
        rounding=ROUND_HALF_UP,
    )


def localizar_bloque_agregado_y(
    ruta,
    resultado,
):

    _, hojas = leer_archivo(
        ruta
    )

    seleccionado = (
        elegir_hoja_principal_creditos(
            hojas
        )
    )

    if seleccionado is None:

        raise ErrorEstructuralProduccion(
            "B-2359: no pudo recuperarse "
            "la estructura para inspeccionar "
            "el agregado."
        )

    if (
        seleccionado[
            "hoja"
        ][
            "nombre"
        ]
        !=
        resultado[
            "hoja"
        ]
    ):

        raise ErrorEstructuralProduccion(
            "B-2359: la hoja usada para el formato "
            "no coincide con la del parser."
        )

    agregados = [
        bloque
        for bloque
        in seleccionado[
            "estructura"
        ][
            "bloques"
        ]
        if es_agregado_creditos(
            bloque[
                "banco_original"
            ]
        )
    ]

    if len(
        agregados
    ) != 1:

        raise ErrorEstructuralProduccion(
            "B-2359: no se localizó exactamente "
            "un bloque Total Banca Múltiple."
        )

    return agregados[
        0
    ]


def leer_formato_agregado_y_xls(
    ruta,
    resultado,
    bloque_agregado,
):

    try:

        libro = xlrd.open_workbook(
            filename=str(
                ruta
            ),
            formatting_info=True,
            on_demand=True,
        )

    except Exception as error:

        raise ErrorEstructuralProduccion(
            "B-2359: xlrd no pudo abrir el XLS "
            "con formatting_info=True."
        ) from error

    try:

        hoja = libro.sheet_by_name(
            resultado[
                "hoja"
            ]
        )

        fila = resultado[
            "fila_total_creditos"
        ]

        salida = {}

        for variable, columna in (
            (
                "MN",
                bloque_agregado[
                    "columna_mn"
                ],
            ),
            (
                "ME",
                bloque_agregado[
                    "columna_me"
                ],
            ),
            (
                "TOTAL",
                bloque_agregado[
                    "columna_total"
                ],
            ),
        ):

            celda = hoja.cell(
                fila,
                columna,
            )

            xf_index = hoja.cell_xf_index(
                fila,
                columna,
            )

            if (
                xf_index < 0
                or
                xf_index >= len(
                    libro.xf_list
                )
            ):

                raise ErrorEstructuralProduccion(
                    f"B-2359 {variable}: "
                    "xf_index inválido."
                )

            xf = libro.xf_list[
                xf_index
            ]

            format_key = getattr(
                xf,
                "format_key",
                None,
            )

            if format_key is None:

                raise ErrorEstructuralProduccion(
                    f"B-2359 {variable}: "
                    "format_key no disponible."
                )

            formato_obj = (
                libro.format_map.get(
                    format_key
                )
            )

            if formato_obj is None:

                raise ErrorEstructuralProduccion(
                    f"B-2359 {variable}: "
                    "formato Excel no recuperable."
                )

            format_str = getattr(
                formato_obj,
                "format_str",
                None,
            )

            decimales = (
                inferir_decimales_formato_excel(
                    format_str
                )
            )

            if (
                format_str is None
                or
                decimales is None
            ):

                raise ErrorEstructuralProduccion(
                    f"B-2359 {variable}: no puede "
                    "determinarse la precisión visible; "
                    f"format_str={format_str!r}."
                )

            salida[
                variable
            ] = {
                "valor_raw_celda":
                    celda.value,

                "format_str":
                    format_str,

                "decimales_mostrados":
                    decimales,
            }

        return salida

    finally:

        libro.release_resources()


def leer_formato_agregado_y_xlsx(
    ruta,
    resultado,
    bloque_agregado,
):

    try:

        with Path(
            ruta
        ).open(
            "rb"
        ) as archivo_binario:

            libro = load_workbook(
                archivo_binario,
                read_only=False,
                data_only=True,
            )

            try:

                hoja = libro[
                    resultado[
                        "hoja"
                    ]
                ]

                fila = (
                    resultado[
                        "fila_total_creditos"
                    ]
                    +
                    1
                )

                salida = {}

                for variable, columna_0based in (
                    (
                        "MN",
                        bloque_agregado[
                            "columna_mn"
                        ],
                    ),
                    (
                        "ME",
                        bloque_agregado[
                            "columna_me"
                        ],
                    ),
                    (
                        "TOTAL",
                        bloque_agregado[
                            "columna_total"
                        ],
                    ),
                ):

                    celda = hoja.cell(
                        row=fila,
                        column=(
                            columna_0based
                            +
                            1
                        ),
                    )

                    format_str = (
                        celda.number_format
                    )

                    decimales = (
                        inferir_decimales_formato_excel(
                            format_str
                        )
                    )

                    if (
                        format_str is None
                        or
                        decimales is None
                    ):

                        raise ErrorEstructuralProduccion(
                            f"B-2359 {variable}: no puede "
                            "determinarse la precisión visible; "
                            f"format_str={format_str!r}."
                        )

                    salida[
                        variable
                    ] = {
                        "valor_raw_celda":
                            celda.value,

                        "format_str":
                            format_str,

                        "decimales_mostrados":
                            decimales,
                    }

                return salida

            finally:

                libro.close()

    except ErrorEstructuralProduccion:

        raise

    except Exception as error:

        raise ErrorEstructuralProduccion(
            "B-2359: no pudo inspeccionarse "
            "el formato XLSX del agregado."
        ) from error


def obtener_formatos_agregado_y(
    ruta,
    resultado,
):

    bloque = localizar_bloque_agregado_y(
        ruta,
        resultado,
    )

    if (
        resultado[
            "formato"
        ]
        ==
        "XLS_OLE"
    ):

        return leer_formato_agregado_y_xls(
            ruta,
            resultado,
            bloque,
        )

    if (
        resultado[
            "formato"
        ]
        ==
        "XLSX_ZIP"
    ):

        return leer_formato_agregado_y_xlsx(
            ruta,
            resultado,
            bloque,
        )

    raise ErrorEstructuralProduccion(
        "B-2359: formato no compatible "
        "con inspección de precisión."
    )


def validar_agregado_creditos(
    ruta,
    resultado,
    registros,
):

    agregados = [
        registro
        for registro
        in registros
        if registro[
            "es_agregado"
        ]
    ]

    bancos = [
        registro
        for registro
        in registros
        if not registro[
            "es_agregado"
        ]
    ]

    if len(
        agregados
    ) != 1:

        raise ErrorEstructuralProduccion(
            "B-2359: se esperaba exactamente "
            "un Total Banca Múltiple."
        )

    agregado = agregados[
        0
    ]

    formatos = obtener_formatos_agregado_y(
        ruta,
        resultado,
    )

    controles = []

    for variable, campo in (
        (
            "MN",
            "mn_soles_miles",
        ),
        (
            "ME",
            "me_usd_miles",
        ),
        (
            "TOTAL",
            "total_soles_miles",
        ),
    ):

        valores = [
            registro[
                campo
            ]
            for registro
            in bancos
        ]

        if any(
            valor is None
            for valor
            in valores
        ):

            faltantes = [
                registro[
                    "banco_original"
                ]
                for registro
                in bancos
                if registro[
                    campo
                ] is None
            ]

            raise ErrorEstructuralProduccion(
                f"B-2359 {variable}: valores faltantes "
                "impiden validar agregado: "
                f"{faltantes}"
            )

        total_sbs = agregado[
            campo
        ]

        if total_sbs is None:

            raise ErrorEstructuralProduccion(
                f"B-2359 {variable}: "
                "agregado SBS faltante."
            )

        suma = math.fsum(
            valores
        )

        metadata = formatos[
            variable
        ]

        raw_excel = convertir_numero(
            metadata[
                "valor_raw_celda"
            ]
        )

        if raw_excel is None:

            raise ErrorEstructuralProduccion(
                f"B-2359 {variable}: "
                "valor raw agregado no numérico."
            )

        suma_decimal = Decimal(
            str(
                suma
            )
        )

        total_sbs_decimal = Decimal(
            str(
                total_sbs
            )
        )

        raw_excel_decimal = Decimal(
            str(
                raw_excel
            )
        )

        if (
            raw_excel_decimal
            !=
            total_sbs_decimal
        ):

            raise ErrorEstructuralProduccion(
                f"B-2359 {variable}: valor raw "
                "de celda y valor del parser "
                "no coinciden."
            )

        diferencia_decimal = (
            suma_decimal
            -
            total_sbs_decimal
        )

        diferencia_abs_decimal = abs(
            diferencia_decimal
        )

        diferencia_rel_decimal = (
            diferencia_abs_decimal
            /
            abs(
                total_sbs_decimal
            )
            if total_sbs_decimal != 0
            else None
        )

        decimales = metadata[
            "decimales_mostrados"
        ]

        if (
            decimales is None
            or
            decimales < 0
        ):

            raise ErrorEstructuralProduccion(
                f"B-2359 {variable}: no puede "
                "determinarse la precisión visible."
            )

        # Tolerancia derivada exclusivamente
        # de la precisión visible SBS.
        unidad_visible = (
            Decimal(
                "1"
            ).scaleb(
                -decimales
            )
        )

        tolerancia_fuente = (
            unidad_visible
            /
            Decimal(
                "2"
            )
        )

        # ROUND_HALF_UP se conserva solo
        # como evidencia diagnóstica.
        suma_redondeada = (
            redondear_excel_half_up(
                suma,
                decimales,
            )
        )

        total_redondeado = (
            redondear_excel_half_up(
                total_sbs,
                decimales,
            )
        )

        estado = (
            "OK"
            if (
                diferencia_abs_decimal
                <=
                tolerancia_fuente
            )
            else
            "ERROR"
        )

        control = {
            "variable":
                variable,

            "suma_fsum":
                suma,

            "total_sbs_raw":
                total_sbs,

            "diferencia_absoluta_raw":
                str(
                    diferencia_abs_decimal
                ),

            "diferencia_relativa_raw":
                (
                    ""
                    if diferencia_rel_decimal
                    is None
                    else
                    str(
                        diferencia_rel_decimal
                    )
                ),

            "format_str":
                metadata[
                    "format_str"
                ],

            "decimales_mostrados":
                decimales,

            "unidad_visible":
                str(
                    unidad_visible
                ),

            "tolerancia_fuente":
                str(
                    tolerancia_fuente
                ),

            "suma_redondeada":
                str(
                    suma_redondeada
                ),

            "total_redondeado":
                str(
                    total_redondeado
                ),

            "estado":
                estado,
        }

        controles.append(
            control
        )

    errores = [
        control
        for control
        in controles
        if control[
            "estado"
        ]
        !=
        "OK"
    ]

    if errores:

        raise ErrorEstructuralProduccion(
            "B-2359: control Total Banca Múltiple "
            "excede la tolerancia derivada "
            "de la precisión visible SBS: "
            +
            ", ".join(
                control[
                    "variable"
                ]
                for control
                in errores
            )
        )

    return {
        "bancos":
            bancos,

        "agregado":
            agregado,

        "controles":
            controles,
    }


def validar_y_creditos(
    ruta,
    resultado,
):

    registros = resultado[
        "registros"
    ]

    agregado = validar_agregado_creditos(
        ruta,
        resultado,
        registros,
    )

    bancos = agregado[
        "bancos"
    ]

    problemas_estructurales = []

    for registro in bancos:

        problemas_permitidos = {
            "TOTAL_NO_POSITIVO"
        }

        if (
            registro[
                "total_soles_miles"
            ] is not None
            and
            registro[
                "total_soles_miles"
            ] <= 0
        ):

            problemas_permitidos.add(
                "TOTAL_MENOR_QUE_MN"
            )

        graves = [
            problema
            for problema
            in registro[
                "problemas"
            ]
            if problema
            not in problemas_permitidos
        ]

        if graves:

            problemas_estructurales.append({
                "banco":
                    registro[
                        "banco_original"
                    ],

                "problemas":
                    graves,
            })

    if problemas_estructurales:

        raise ErrorEstructuralProduccion(
            "B-2359: problemas estructurales/"
            "económicos no esperados: "
            f"{problemas_estructurales}"
        )

    validos = [
        registro
        for registro
        in bancos
        if registro[
            "dolcred_pct"
        ] is not None
    ]

    faltantes = [
        registro
        for registro
        in bancos
        if registro[
            "dolcred_pct"
        ] is None
    ]

    return {
        "bancos":
            bancos,

        "cantidad_bancos":
            len(
                bancos
            ),

        "cantidad_dolcred_validos":
            len(
                validos
            ),

        "cantidad_faltantes":
            len(
                faltantes
            ),

        "controles":
            agregado[
                "controles"
            ],
    }


# ============================================================
# 18. X1 / B-2318 - PARSER VALIDADO
# ============================================================

def detectar_encabezado_depositos(
    filas
):

    candidatos = []

    for numero_fila, fila in enumerate(
        filas
    ):

        columna_empresa = None
        columna_saldo_final = None

        for numero_columna, valor in enumerate(
            fila
        ):

            texto = normalizar_comparacion(
                valor
            )

            if (
                texto in {
                    "empresa",
                    "empresas",
                }
                or
                texto.startswith(
                    "empresas "
                )
            ):

                columna_empresa = (
                    numero_columna
                )

            if (
                "saldo final"
                in
                texto
            ):

                columna_saldo_final = (
                    numero_columna
                )

        if (
            columna_empresa is not None
            and
            columna_saldo_final is not None
        ):

            candidatos.append({
                "fila_encabezado":
                    numero_fila,

                "columna_empresa":
                    columna_empresa,

                "columna_saldo_final":
                    columna_saldo_final,
            })

    return (
        candidatos[
            0
        ]
        if candidatos
        else None
    )


def obtener_contexto_superior(
    filas,
    fila_encabezado,
    max_filas=8,
):

    inicio = max(
        0,
        fila_encabezado
        -
        max_filas,
    )

    textos = []

    for fila in filas[
        inicio:fila_encabezado
    ]:

        for valor in fila:

            texto = normalizar_comparacion(
                valor
            )

            if texto:

                textos.append(
                    texto
                )

    return " ".join(
        textos
    )


def clasificar_moneda_depositos(
    filas,
    encabezado,
):

    contexto = obtener_contexto_superior(
        filas,
        encabezado[
            "fila_encabezado"
        ],
    )

    tiene_mn = (
        "moneda nacional"
        in
        contexto
    )

    tiene_me = (
        "moneda extranjera"
        in
        contexto
    )

    if (
        tiene_mn
        and
        not tiene_me
    ):

        return "MN"

    if (
        tiene_me
        and
        not tiene_mn
    ):

        return "ME"

    return None


def detectar_tablas_monetarias(
    hojas
):

    candidatos = []

    for hoja in hojas:

        encabezado = (
            detectar_encabezado_depositos(
                hoja[
                    "filas"
                ]
            )
        )

        if encabezado is None:
            continue

        moneda = clasificar_moneda_depositos(
            hoja[
                "filas"
            ],
            encabezado,
        )

        if moneda is None:
            continue

        candidatos.append({
            "moneda":
                moneda,

            "nombre_hoja":
                hoja[
                    "nombre"
                ],

            "filas":
                hoja[
                    "filas"
                ],

            "encabezado":
                encabezado,
        })

    tablas_mn = [
        item
        for item
        in candidatos
        if item[
            "moneda"
        ]
        ==
        "MN"
    ]

    tablas_me = [
        item
        for item
        in candidatos
        if item[
            "moneda"
        ]
        ==
        "ME"
    ]

    return (
        tablas_mn,
        tablas_me,
    )


def es_agregado_depositos(
    nombre
):

    return (
        "total banca multiple"
        in
        normalizar_comparacion(
            nombre
        )
    )


def extraer_saldos(
    tabla
):

    filas = tabla[
        "filas"
    ]

    encabezado = tabla[
        "encabezado"
    ]

    fila_inicio = (
        encabezado[
            "fila_encabezado"
        ]
        +
        1
    )

    col_empresa = encabezado[
        "columna_empresa"
    ]

    col_saldo = encabezado[
        "columna_saldo_final"
    ]

    registros = []

    for numero_fila in range(
        fila_inicio,
        len(
            filas
        ),
    ):

        fila = filas[
            numero_fila
        ]

        if (
            col_empresa >= len(
                fila
            )
            or
            col_saldo >= len(
                fila
            )
        ):

            continue

        nombre = fila[
            col_empresa
        ]

        saldo = convertir_numero(
            fila[
                col_saldo
            ]
        )

        nombre_texto = (
            str(
                nombre
            ).strip()
            if nombre is not None
            else ""
        )

        if not normalizar_texto(
            nombre_texto
        ):

            continue

        if saldo is None:

            continue

        registros.append({
            "fila_origen":
                numero_fila,

            "banco_original":
                nombre_texto,

            "banco_clave":
                normalizar_comparacion(
                    nombre_texto
                ),

            "saldo_final":
                saldo,

            "es_agregado":
                es_agregado_depositos(
                    nombre_texto
                ),
        })

    return registros


def crear_diccionario_bancos(
    registros
):

    resultado = {}
    duplicados = []

    for registro in registros:

        if registro[
            "es_agregado"
        ]:

            continue

        clave = registro[
            "banco_clave"
        ]

        if clave in resultado:

            duplicados.append(
                clave
            )

        else:

            resultado[
                clave
            ] = registro

    return (
        resultado,
        duplicados,
    )


def emparejar_monedas(
    registros_mn,
    registros_me,
):

    dic_mn, duplicados_mn = (
        crear_diccionario_bancos(
            registros_mn
        )
    )

    dic_me, duplicados_me = (
        crear_diccionario_bancos(
            registros_me
        )
    )

    claves_mn = set(
        dic_mn
    )

    claves_me = set(
        dic_me
    )

    comunes = sorted(
        claves_mn
        &
        claves_me
    )

    solo_mn = sorted(
        claves_mn
        -
        claves_me
    )

    solo_me = sorted(
        claves_me
        -
        claves_mn
    )

    emparejados = []

    for clave in comunes:

        reg_mn = dic_mn[
            clave
        ]

        reg_me = dic_me[
            clave
        ]

        mn = reg_mn[
            "saldo_final"
        ]

        me = reg_me[
            "saldo_final"
        ]

        total = (
            mn
            +
            me
        )

        problemas = []

        if mn < 0:

            problemas.append(
                "MN_NEGATIVO"
            )

        if me < 0:

            problemas.append(
                "ME_NEGATIVO"
            )

        if total <= 0:

            problemas.append(
                "TOTAL_NO_POSITIVO"
            )

            doldep = None

        else:

            doldep = (
                me
                /
                total
                *
                100
            )

            if not (
                0
                <=
                doldep
                <=
                100
            ):

                problemas.append(
                    "DOLDEP_FUERA_RANGO"
                )

        emparejados.append({
            "banco_original":
                reg_mn[
                    "banco_original"
                ],

            "banco_original_me":
                reg_me[
                    "banco_original"
                ],

            "saldo_final_mn":
                mn,

            "saldo_final_me":
                me,

            "doldep_pct":
                doldep,

            "problemas":
                problemas,
        })

    return {
        "emparejados":
            emparejados,

        "solo_mn":
            solo_mn,

        "solo_me":
            solo_me,

        "duplicados_mn":
            duplicados_mn,

        "duplicados_me":
            duplicados_me,
    }


def validar_agregado_depositos(
    registros,
    moneda,
):

    individuales = [
        registro
        for registro
        in registros
        if not registro[
            "es_agregado"
        ]
    ]

    agregados = [
        registro
        for registro
        in registros
        if registro[
            "es_agregado"
        ]
    ]

    if len(
        agregados
    ) != 1:

        raise ErrorEstructuralProduccion(
            f"B-2318 {moneda}: se esperaba "
            "exactamente un TOTAL BANCA MÚLTIPLE."
        )

    suma_bancos = sum(
        registro[
            "saldo_final"
        ]
        for registro
        in individuales
    )

    total_sbs = agregados[
        0
    ][
        "saldo_final"
    ]

    diferencia = (
        suma_bancos
        -
        total_sbs
    )

    diferencia_abs = abs(
        diferencia
    )

    coincide = math.isclose(
        suma_bancos,
        total_sbs,
        rel_tol=REL_TOL_X1,
        abs_tol=ABS_TOL_X1,
    )

    if not coincide:

        raise ErrorEstructuralProduccion(
            f"B-2318 {moneda}: la suma de bancos "
            "no reproduce TOTAL BANCA MÚLTIPLE "
            "dentro del control ya validado."
        )

    return {
        "moneda":
            moneda,

        "cantidad_bancos":
            len(
                individuales
            ),

        "suma_bancos":
            suma_bancos,

        "total_sbs":
            total_sbs,

        "diferencia_absoluta":
            diferencia_abs,

        "estado":
            "OK",
    }


def extraer_x1_depositos(
    ruta
):

    formato, hojas = leer_archivo(
        ruta
    )

    tablas_mn, tablas_me = (
        detectar_tablas_monetarias(
            hojas
        )
    )

    if (
        len(
            tablas_mn
        )
        !=
        1
        or
        len(
            tablas_me
        )
        !=
        1
    ):

        raise ErrorEstructuralProduccion(
            "B-2318: se esperaba exactamente "
            "una tabla MN y una tabla ME. "
            f"MN={len(tablas_mn)}, "
            f"ME={len(tablas_me)}."
        )

    tabla_mn = tablas_mn[
        0
    ]

    tabla_me = tablas_me[
        0
    ]

    registros_mn = extraer_saldos(
        tabla_mn
    )

    registros_me = extraer_saldos(
        tabla_me
    )

    control_mn = validar_agregado_depositos(
        registros_mn,
        "MN",
    )

    control_me = validar_agregado_depositos(
        registros_me,
        "ME",
    )

    emparejamiento = emparejar_monedas(
        registros_mn,
        registros_me,
    )

    return {
        "formato":
            formato,

        "tabla_mn":
            tabla_mn,

        "tabla_me":
            tabla_me,

        "control_mn":
            control_mn,

        "control_me":
            control_me,

        "emparejamiento":
            emparejamiento,
    }


def validar_x1_depositos(
    resultado
):

    emp = resultado[
        "emparejamiento"
    ]

    if emp[
        "duplicados_mn"
    ]:

        raise ErrorEstructuralProduccion(
            "B-2318: bancos duplicados en MN: "
            f"{emp['duplicados_mn']}"
        )

    if emp[
        "duplicados_me"
    ]:

        raise ErrorEstructuralProduccion(
            "B-2318: bancos duplicados en ME: "
            f"{emp['duplicados_me']}"
        )

    if emp[
        "solo_mn"
    ]:

        raise ErrorEstructuralProduccion(
            "B-2318: bancos solo en MN: "
            f"{emp['solo_mn']}"
        )

    if emp[
        "solo_me"
    ]:

        raise ErrorEstructuralProduccion(
            "B-2318: bancos solo en ME: "
            f"{emp['solo_me']}"
        )

    problemas_graves = []

    for registro in emp[
        "emparejados"
    ]:

        graves = [
            problema
            for problema
            in registro[
                "problemas"
            ]
            if problema
            !=
            "TOTAL_NO_POSITIVO"
        ]

        if graves:

            problemas_graves.append({
                "banco":
                    registro[
                        "banco_original"
                    ],

                "problemas":
                    graves,
            })

    if problemas_graves:

        raise ErrorEstructuralProduccion(
            "B-2318: problemas no esperados: "
            f"{problemas_graves}"
        )

    faltantes = [
        registro
        for registro
        in emp[
            "emparejados"
        ]
        if registro[
            "doldep_pct"
        ] is None
    ]

    return {
        "bancos":
            emp[
                "emparejados"
            ],

        "cantidad_bancos":
            len(
                emp[
                    "emparejados"
                ]
            ),

        "cantidad_faltantes":
            len(
                faltantes
            ),

        "control_mn":
            resultado[
                "control_mn"
            ],

        "control_me":
            resultado[
                "control_me"
            ],
    }


# ============================================================
# 19. CONSTRUIR FILAS DE DATOS
# ============================================================

def construir_filas_y(
    mes,
    validacion,
):

    filas = []

    for registro in validacion[
        "bancos"
    ]:

        filas.append({
            "mes":
                mes,

            "banco_original":
                registro[
                    "banco_original"
                ],

            "mn_soles_miles":
                (
                    ""
                    if registro[
                        "mn_soles_miles"
                    ] is None
                    else
                    registro[
                        "mn_soles_miles"
                    ]
                ),

            "me_usd_miles":
                (
                    ""
                    if registro[
                        "me_usd_miles"
                    ] is None
                    else
                    registro[
                        "me_usd_miles"
                    ]
                ),

            "total_soles_miles":
                (
                    ""
                    if registro[
                        "total_soles_miles"
                    ] is None
                    else
                    registro[
                        "total_soles_miles"
                    ]
                ),

            "dolcred_pct":
                (
                    ""
                    if registro[
                        "dolcred_pct"
                    ] is None
                    else
                    registro[
                        "dolcred_pct"
                    ]
                ),

            "problemas":
                "|".join(
                    registro[
                        "problemas"
                    ]
                ),
        })

    return filas


def construir_filas_x1(
    mes,
    validacion,
):

    filas = []

    for registro in validacion[
        "bancos"
    ]:

        filas.append({
            "mes":
                mes,

            "banco_original":
                registro[
                    "banco_original"
                ],

            "banco_original_me":
                registro[
                    "banco_original_me"
                ],

            "saldo_final_mn":
                registro[
                    "saldo_final_mn"
                ],

            "saldo_final_me":
                registro[
                    "saldo_final_me"
                ],

            "doldep_pct":
                (
                    ""
                    if registro[
                        "doldep_pct"
                    ] is None
                    else
                    registro[
                        "doldep_pct"
                    ]
                ),

            "problemas":
                "|".join(
                    registro[
                        "problemas"
                    ]
                ),
        })

    return filas


# ============================================================
# 20. AUDITORÍAS
# ============================================================

def ahora_iso():

    return (
        datetime.now()
        .astimezone()
        .isoformat(
            timespec="seconds"
        )
    )


def _control_por_variable(
    controles,
    variable,
):

    encontrados = [
        control
        for control
        in controles
        if control[
            "variable"
        ]
        ==
        variable
    ]

    if len(
        encontrados
    ) != 1:

        raise ErrorEstructuralProduccion(
            f"Control Y {variable} "
            "no unívoco."
        )

    return encontrados[
        0
    ]


def auditoria_y_disponible(
    mes,
    info_descarga,
    resultado,
    validacion,
):

    c_mn = _control_por_variable(
        validacion[
            "controles"
        ],
        "MN",
    )

    c_me = _control_por_variable(
        validacion[
            "controles"
        ],
        "ME",
    )

    c_total = _control_por_variable(
        validacion[
            "controles"
        ],
        "TOTAL",
    )

    return {
        "mes":
            mes,

        "estado":
            "DISPONIBLE",

        "url_fuente":
            info_descarga[
                "url"
            ],

        "archivo_origen":
            info_descarga[
                "archivo_origen"
            ],

        "archivo_local":
            info_descarga[
                "archivo_local"
            ],

        "sha256":
            info_descarga[
                "sha256"
            ],

        "formato_real":
            resultado[
                "formato"
            ],

        "hoja":
            resultado[
                "hoja"
            ],

        "fila_total_creditos":
            resultado[
                "fila_total_creditos"
            ],

        "cantidad_bancos":
            validacion[
                "cantidad_bancos"
            ],

        "cantidad_dolcred_validos":
            validacion[
                "cantidad_dolcred_validos"
            ],

        "cantidad_faltantes":
            validacion[
                "cantidad_faltantes"
            ],

        # NUEVO:
        # Evidencia explícita de que este mes Y fue calculado con
        # la normalización individual de hasta 1 ULP en Total-MN.
        "regla_float_total_mn":
            "ULP_1",

        "control_mn_estado":
            c_mn[
                "estado"
            ],

        "control_mn_diferencia_absoluta_raw":
            c_mn[
                "diferencia_absoluta_raw"
            ],

        "control_mn_diferencia_relativa_raw":
            c_mn[
                "diferencia_relativa_raw"
            ],

        "control_mn_formato_excel":
            c_mn[
                "format_str"
            ],

        "control_mn_decimales_mostrados":
            c_mn[
                "decimales_mostrados"
            ],

        "control_mn_unidad_visible":
            c_mn[
                "unidad_visible"
            ],

        "control_mn_tolerancia_fuente":
            c_mn[
                "tolerancia_fuente"
            ],

        "control_mn_suma_redondeada_half_up":
            c_mn[
                "suma_redondeada"
            ],

        "control_mn_total_redondeado_half_up":
            c_mn[
                "total_redondeado"
            ],

        "control_me_estado":
            c_me[
                "estado"
            ],

        "control_me_diferencia_absoluta_raw":
            c_me[
                "diferencia_absoluta_raw"
            ],

        "control_me_diferencia_relativa_raw":
            c_me[
                "diferencia_relativa_raw"
            ],

        "control_me_formato_excel":
            c_me[
                "format_str"
            ],

        "control_me_decimales_mostrados":
            c_me[
                "decimales_mostrados"
            ],

        "control_me_unidad_visible":
            c_me[
                "unidad_visible"
            ],

        "control_me_tolerancia_fuente":
            c_me[
                "tolerancia_fuente"
            ],

        "control_me_suma_redondeada_half_up":
            c_me[
                "suma_redondeada"
            ],

        "control_me_total_redondeado_half_up":
            c_me[
                "total_redondeado"
            ],

        "control_total_estado":
            c_total[
                "estado"
            ],

        "control_total_diferencia_absoluta_raw":
            c_total[
                "diferencia_absoluta_raw"
            ],

        "control_total_diferencia_relativa_raw":
            c_total[
                "diferencia_relativa_raw"
            ],

        "control_total_formato_excel":
            c_total[
                "format_str"
            ],

        "control_total_decimales_mostrados":
            c_total[
                "decimales_mostrados"
            ],

        "control_total_unidad_visible":
            c_total[
                "unidad_visible"
            ],

        "control_total_tolerancia_fuente":
            c_total[
                "tolerancia_fuente"
            ],

        "control_total_suma_redondeada_half_up":
            c_total[
                "suma_redondeada"
            ],

        "control_total_total_redondeado_half_up":
            c_total[
                "total_redondeado"
            ],

        "fecha_hora_extraccion":
            ahora_iso(),

        "detalle":
            (
                "B-2359 extraído y validado con "
                "tolerancia de fuente derivada "
                "de la precisión visible SBS; "
                "ROUND_HALF_UP se conserva "
                "como diagnóstico."
            ),
    }


def auditoria_x1_disponible(
    mes,
    info_descarga,
    resultado,
    validacion,
):

    cmn = validacion[
        "control_mn"
    ]

    cme = validacion[
        "control_me"
    ]

    return {
        "mes":
            mes,

        "estado":
            "DISPONIBLE",

        "url_fuente":
            info_descarga[
                "url"
            ],

        "archivo_origen":
            info_descarga[
                "archivo_origen"
            ],

        "archivo_local":
            info_descarga[
                "archivo_local"
            ],

        "sha256":
            info_descarga[
                "sha256"
            ],

        "formato_real":
            resultado[
                "formato"
            ],

        "hoja_mn":
            resultado[
                "tabla_mn"
            ][
                "nombre_hoja"
            ],

        "hoja_me":
            resultado[
                "tabla_me"
            ][
                "nombre_hoja"
            ],

        "cantidad_bancos":
            validacion[
                "cantidad_bancos"
            ],

        "cantidad_faltantes":
            validacion[
                "cantidad_faltantes"
            ],

        "control_mn_estado":
            cmn[
                "estado"
            ],

        "control_mn_suma_bancos":
            cmn[
                "suma_bancos"
            ],

        "control_mn_total_sbs":
            cmn[
                "total_sbs"
            ],

        "control_mn_diferencia_absoluta":
            cmn[
                "diferencia_absoluta"
            ],

        "control_me_estado":
            cme[
                "estado"
            ],

        "control_me_suma_bancos":
            cme[
                "suma_bancos"
            ],

        "control_me_total_sbs":
            cme[
                "total_sbs"
            ],

        "control_me_diferencia_absoluta":
            cme[
                "diferencia_absoluta"
            ],

        "fecha_hora_extraccion":
            ahora_iso(),

        "detalle":
            (
                "B-2318 extraído y validado; "
                "MN/ME emparejados solo "
                "dentro del mismo mes."
            ),
    }


def auditoria_error(
    variable,
    mes,
    estado,
    detalle,
    enlace=None,
):

    base = {
        "mes":
            mes,

        "estado":
            estado,

        "url_fuente":
            (
                enlace[
                    "url"
                ]
                if enlace
                else ""
            ),

        "archivo_origen":
            (
                Path(
                    urlparse(
                        enlace[
                            "url"
                        ]
                    ).path
                ).name
                if enlace
                else ""
            ),

        "archivo_local":
            "",

        "sha256":
            "",

        "formato_real":
            "",

        "fecha_hora_extraccion":
            ahora_iso(),

        "detalle":
            detalle,
    }

    if variable == "Y":

        # La nueva columna regla_float_total_mn queda vacía
        # automáticamente en estados de error o MES_SIN_DATO,
        # igual que las demás evidencias no disponibles.
        return {
            columna:
                base.get(
                    columna,
                    "",
                )
            for columna
            in COLUMNAS_AUDITORIA_Y
        }

    return {
        columna:
            base.get(
                columna,
                "",
            )
        for columna
        in COLUMNAS_AUDITORIA_X1
    }


# ============================================================
# 21. PROCESAR Y DE UN MES
# ============================================================

def procesar_y_mes(
    mes,
    enlace,
):

    info = descargar_o_reutilizar(
        "Y",
        REPORTES[
            "Y"
        ][
            "codigo"
        ],
        mes,
        enlace,
        CARPETA_CRUDOS_Y,
    )

    ruta = Path(
        info[
            "archivo_local"
        ]
    )

    resultado = extraer_y_creditos(
        ruta
    )

    validacion = validar_y_creditos(
        ruta,
        resultado,
    )

    filas = construir_filas_y(
        mes,
        validacion,
    )

    if (
        len(
            filas
        )
        !=
        validacion[
            "cantidad_bancos"
        ]
    ):

        raise ErrorEstructuralProduccion(
            f"Y {mes}: filas construidas "
            "!= cantidad_bancos."
        )

    claves = [
        (
            fila[
                "mes"
            ],
            fila[
                "banco_original"
            ],
        )
        for fila
        in filas
    ]

    if len(
        claves
    ) != len(
        set(
            claves
        )
    ):

        raise ErrorEstructuralProduccion(
            f"Y {mes}: duplicados "
            "mes+banco_original dentro del mes."
        )

    audit = auditoria_y_disponible(
        mes,
        info,
        resultado,
        validacion,
    )

    return (
        filas,
        audit,
    )


# ============================================================
# 22. PROCESAR X1 DE UN MES
# ============================================================

def procesar_x1_mes(
    mes,
    enlace,
):

    info = descargar_o_reutilizar(
        "X1",
        REPORTES[
            "X1"
        ][
            "codigo"
        ],
        mes,
        enlace,
        CARPETA_CRUDOS_X1,
    )

    ruta = Path(
        info[
            "archivo_local"
        ]
    )

    resultado = extraer_x1_depositos(
        ruta
    )

    validacion = validar_x1_depositos(
        resultado
    )

    filas = construir_filas_x1(
        mes,
        validacion,
    )

    if (
        len(
            filas
        )
        !=
        validacion[
            "cantidad_bancos"
        ]
    ):

        raise ErrorEstructuralProduccion(
            f"X1 {mes}: filas construidas "
            "!= cantidad_bancos."
        )

    claves = [
        (
            fila[
                "mes"
            ],
            fila[
                "banco_original"
            ],
        )
        for fila
        in filas
    ]

    if len(
        claves
    ) != len(
        set(
            claves
        )
    ):

        raise ErrorEstructuralProduccion(
            f"X1 {mes}: duplicados "
            "mes+banco_original dentro del mes."
        )

    audit = auditoria_x1_disponible(
        mes,
        info,
        resultado,
        validacion,
    )

    return (
        filas,
        audit,
    )


# ============================================================
# 23. RESUMEN TXT
# ============================================================

def escribir_resumen(
    y_por_mes,
    x1_por_mes,
    auditoria_y,
    auditoria_x1,
    estado_final="EN_PROCESO",
):

    lineas = [
        "EXTRACCIÓN PRODUCCIÓN Y + X1",
        "Estudiante: BRICEÑO LEON CRYSTELL HIDEKI",
        f"Código: {CODIGO_ESTUDIANTE}",
        "Período: 2015-11 a 2025-12",
        f"Meses objetivo: {len(MESES_ESPERADOS)}",
        f"Estado general: {estado_final}",
        "",
        "Y / B-2359",
        (
            "  Meses auditados: "
            f"{len(auditoria_y)}"
        ),
        (
            "  Filas banco-mes: "
            f"{sum(len(v) for v in y_por_mes.values())}"
        ),
        "X1 / B-2318",
        (
            "  Meses auditados: "
            f"{len(auditoria_x1)}"
        ),
        (
            "  Filas banco-mes: "
            f"{sum(len(v) for v in x1_por_mes.values())}"
        ),
        "",
        (
            "mes | Y_estado | Y_bancos | "
            "X1_estado | X1_bancos"
        ),
    ]

    for mes in MESES_ESPERADOS:

        ay = auditoria_y.get(
            mes,
            {},
        )

        ax = auditoria_x1.get(
            mes,
            {},
        )

        lineas.append(
            f"{mes} | "
            f"{ay.get('estado', '')} | "
            f"{ay.get('cantidad_bancos', '')} | "
            f"{ax.get('estado', '')} | "
            f"{ax.get('cantidad_bancos', '')}"
        )

    TXT_RESUMEN.write_text(
        "\n".join(
            lineas
        ),
        encoding="utf-8",
    )


# ============================================================
# 24. VALIDACIÓN FINAL
# ============================================================

def validar_final(
    y_por_mes,
    x1_por_mes,
    auditoria_y,
    auditoria_x1,
):

    if (
        sorted(
            auditoria_y
        )
        !=
        MESES_ESPERADOS
    ):

        faltan = sorted(
            set(
                MESES_ESPERADOS
            )
            -
            set(
                auditoria_y
            )
        )

        extras = sorted(
            set(
                auditoria_y
            )
            -
            set(
                MESES_ESPERADOS
            )
        )

        raise ErrorEstructuralProduccion(
            "Auditoría Y no contiene exactamente "
            "los 122 meses. "
            f"Faltan={faltan}; extras={extras}"
        )

    if (
        sorted(
            auditoria_x1
        )
        !=
        MESES_ESPERADOS
    ):

        faltan = sorted(
            set(
                MESES_ESPERADOS
            )
            -
            set(
                auditoria_x1
            )
        )

        extras = sorted(
            set(
                auditoria_x1
            )
            -
            set(
                MESES_ESPERADOS
            )
        )

        raise ErrorEstructuralProduccion(
            "Auditoría X1 no contiene exactamente "
            "los 122 meses. "
            f"Faltan={faltan}; extras={extras}"
        )

    estados_y = {
        mes:
            auditoria_y[
                mes
            ][
                "estado"
            ]
        for mes
        in MESES_ESPERADOS
    }

    estados_x1 = {
        mes:
            auditoria_x1[
                mes
            ][
                "estado"
            ]
        for mes
        in MESES_ESPERADOS
    }

    no_disponibles_y = [
        mes
        for mes, estado
        in estados_y.items()
        if estado
        !=
        "DISPONIBLE"
    ]

    no_disponibles_x1 = [
        mes
        for mes, estado
        in estados_x1.items()
        if estado
        !=
        "DISPONIBLE"
    ]

    if (
        no_disponibles_y
        or
        no_disponibles_x1
    ):

        estados_problematicos = (
            [
                estados_y[
                    mes
                ]
                for mes
                in no_disponibles_y
            ]
            +
            [
                estados_x1[
                    mes
                ]
                for mes
                in no_disponibles_x1
            ]
        )

        if all(
            estado
            ==
            "MES_SIN_DATO"
            for estado
            in estados_problematicos
        ):

            raise DatosFaltantesProduccion(
                "La extracción recorrió los 122 meses, "
                "pero la fuente oficial no publicó "
                "uno o más archivos mensuales exactos. "
                f"Y={no_disponibles_y}; "
                f"X1={no_disponibles_x1}"
            )

        raise ErrorEstructuralProduccion(
            "La auditoría final contiene estados "
            "incompatibles con una extracción completa. "
            f"Y={[(m, estados_y[m]) for m in no_disponibles_y]}; "
            f"X1={[(m, estados_x1[m]) for m in no_disponibles_x1]}"
        )

    if (
        set(
            y_por_mes
        )
        !=
        set(
            MESES_ESPERADOS
        )
    ):

        raise ErrorEstructuralProduccion(
            "Y no contiene exactamente "
            "los 122 meses en el archivo principal."
        )

    if (
        set(
            x1_por_mes
        )
        !=
        set(
            MESES_ESPERADOS
        )
    ):

        raise ErrorEstructuralProduccion(
            "X1 no contiene exactamente "
            "los 122 meses en el archivo principal."
        )

    filas_y = [
        fila
        for mes
        in MESES_ESPERADOS
        for fila
        in y_por_mes[
            mes
        ]
    ]

    filas_x1 = [
        fila
        for mes
        in MESES_ESPERADOS
        for fila
        in x1_por_mes[
            mes
        ]
    ]

    claves_y = [
        (
            limpiar_texto(
                fila[
                    "mes"
                ]
            ),
            limpiar_texto(
                fila[
                    "banco_original"
                ]
            ),
        )
        for fila
        in filas_y
    ]

    claves_x1 = [
        (
            limpiar_texto(
                fila[
                    "mes"
                ]
            ),
            limpiar_texto(
                fila[
                    "banco_original"
                ]
            ),
        )
        for fila
        in filas_x1
    ]

    if len(
        claves_y
    ) != len(
        set(
            claves_y
        )
    ):

        raise ErrorEstructuralProduccion(
            "Y contiene duplicados "
            "mes+banco_original."
        )

    if len(
        claves_x1
    ) != len(
        set(
            claves_x1
        )
    ):

        raise ErrorEstructuralProduccion(
            "X1 contiene duplicados "
            "mes+banco_original."
        )

    for mes in MESES_ESPERADOS:

        if not checkpoint_y_completo(
            mes,
            y_por_mes,
            auditoria_y,
        ):

            raise ErrorEstructuralProduccion(
                f"Checkpoint final Y inválido "
                f"en {mes}."
            )

        if not checkpoint_x1_completo(
            mes,
            x1_por_mes,
            auditoria_x1,
        ):

            raise ErrorEstructuralProduccion(
                f"Checkpoint final X1 inválido "
                f"en {mes}."
            )

    return {
        "meses":
            122,

        "filas_y":
            len(
                filas_y
            ),

        "filas_x1":
            len(
                filas_x1
            ),

        "duplicados_y":
            0,

        "duplicados_x1":
            0,
    }


# ============================================================
# 25. MAIN
# ============================================================

def main():

    (
        y_por_mes,
        x1_por_mes,
        auditoria_y,
        auditoria_x1,
    ) = cargar_estado()

    try:

        print(
            "=" * 100
        )

        print(
            "EXTRACCIÓN PRODUCCIÓN Y + X1"
        )

        print(
            "=" * 100
        )

        print(
            "Período: 2015-11 a 2025-12"
        )

        print(
            "Meses esperados: 122"
        )

        print(
            "No se homologan nombres."
        )

        print(
            "No se interpolan datos."
        )

        print(
            "No se seleccionan bancos."
        )

        print(
            "Un error técnico/estructural "
            "detiene la ejecución y NO "
            "se convierte en sin dato."
        )

        _, enlaces_y = cargar_enlaces_reporte(
            REPORTES[
                "Y"
            ][
                "codigo"
            ]
        )

        _, enlaces_x1 = cargar_enlaces_reporte(
            REPORTES[
                "X1"
            ][
                "codigo"
            ]
        )

        mapa_y = construir_mapa_mensual(
            REPORTES[
                "Y"
            ][
                "codigo"
            ],
            enlaces_y,
        )

        mapa_x1 = construir_mapa_mensual(
            REPORTES[
                "X1"
            ][
                "codigo"
            ],
            enlaces_x1,
        )

        for indice, mes in enumerate(
            MESES_ESPERADOS,
            start=1,
        ):

            print()

            print(
                f"[{indice:03d}/122] {mes}"
            )

            if (
                REANUDAR
                and
                checkpoint_mes_completo(
                    mes,
                    y_por_mes,
                    x1_por_mes,
                    auditoria_y,
                    auditoria_x1,
                )
            ):

                print(
                    "  [CHECKPOINT] Y y X1 completos. "
                    "Se omite el mes."
                )

                continue

            # ------------------------------------------------
            # Y
            # ------------------------------------------------

            if (
                REANUDAR
                and
                checkpoint_y_completo(
                    mes,
                    y_por_mes,
                    auditoria_y,
                )
            ):

                print(
                    "  Y: checkpoint válido. "
                    "Se conserva."
                )

            else:

                y_por_mes[
                    mes
                ] = []

                try:

                    filas_y, audit_y = (
                        procesar_y_mes(
                            mes,
                            mapa_y[
                                mes
                            ],
                        )
                    )

                    y_por_mes[
                        mes
                    ] = filas_y

                    auditoria_y[
                        mes
                    ] = audit_y

                    print(
                        "  Y: DISPONIBLE | "
                        f"bancos={audit_y['cantidad_bancos']} | "
                        f"faltantes={audit_y['cantidad_faltantes']}"
                    )

                except MesSinDatoOficial as error:

                    y_por_mes[
                        mes
                    ] = []

                    auditoria_y[
                        mes
                    ] = auditoria_error(
                        "Y",
                        mes,
                        "MES_SIN_DATO",
                        str(
                            error
                        ),
                        mapa_y[
                            mes
                        ],
                    )

                    print(
                        f"  Y: MES_SIN_DATO | "
                        f"{error}"
                    )

                except ErrorTecnicoProduccion as error:

                    y_por_mes[
                        mes
                    ] = []

                    auditoria_y[
                        mes
                    ] = auditoria_error(
                        "Y",
                        mes,
                        "ERROR_TECNICO",
                        str(
                            error
                        ),
                        mapa_y[
                            mes
                        ],
                    )

                    guardar_estado(
                        y_por_mes,
                        x1_por_mes,
                        auditoria_y,
                        auditoria_x1,
                    )

                    escribir_resumen(
                        y_por_mes,
                        x1_por_mes,
                        auditoria_y,
                        auditoria_x1,
                        estado_final=(
                            "DETENIDO_ERROR_TECNICO_Y"
                        ),
                    )

                    raise

                except ErrorEstructuralProduccion as error:

                    y_por_mes[
                        mes
                    ] = []

                    auditoria_y[
                        mes
                    ] = auditoria_error(
                        "Y",
                        mes,
                        "ERROR_ESTRUCTURAL",
                        str(
                            error
                        ),
                        mapa_y[
                            mes
                        ],
                    )

                    guardar_estado(
                        y_por_mes,
                        x1_por_mes,
                        auditoria_y,
                        auditoria_x1,
                    )

                    escribir_resumen(
                        y_por_mes,
                        x1_por_mes,
                        auditoria_y,
                        auditoria_x1,
                        estado_final=(
                            "DETENIDO_ERROR_ESTRUCTURAL_Y"
                        ),
                    )

                    raise

                guardar_estado(
                    y_por_mes,
                    x1_por_mes,
                    auditoria_y,
                    auditoria_x1,
                )

                escribir_resumen(
                    y_por_mes,
                    x1_por_mes,
                    auditoria_y,
                    auditoria_x1,
                )

                if not checkpoint_y_completo(
                    mes,
                    y_por_mes,
                    auditoria_y,
                ):

                    if (
                        auditoria_y[
                            mes
                        ][
                            "estado"
                        ]
                        !=
                        "MES_SIN_DATO"
                    ):

                        raise ErrorEstructuralProduccion(
                            f"Y {mes}: checkpoint "
                            "recién guardado no es coherente."
                        )

            # ------------------------------------------------
            # X1
            # ------------------------------------------------

            if (
                REANUDAR
                and
                checkpoint_x1_completo(
                    mes,
                    x1_por_mes,
                    auditoria_x1,
                )
            ):

                print(
                    "  X1: checkpoint válido. "
                    "Se conserva."
                )

            else:

                x1_por_mes[
                    mes
                ] = []

                try:

                    filas_x1, audit_x1 = (
                        procesar_x1_mes(
                            mes,
                            mapa_x1[
                                mes
                            ],
                        )
                    )

                    x1_por_mes[
                        mes
                    ] = filas_x1

                    auditoria_x1[
                        mes
                    ] = audit_x1

                    print(
                        "  X1: DISPONIBLE | "
                        f"bancos={audit_x1['cantidad_bancos']} | "
                        f"faltantes={audit_x1['cantidad_faltantes']}"
                    )

                except MesSinDatoOficial as error:

                    x1_por_mes[
                        mes
                    ] = []

                    auditoria_x1[
                        mes
                    ] = auditoria_error(
                        "X1",
                        mes,
                        "MES_SIN_DATO",
                        str(
                            error
                        ),
                        mapa_x1[
                            mes
                        ],
                    )

                    print(
                        f"  X1: MES_SIN_DATO | "
                        f"{error}"
                    )

                except ErrorTecnicoProduccion as error:

                    x1_por_mes[
                        mes
                    ] = []

                    auditoria_x1[
                        mes
                    ] = auditoria_error(
                        "X1",
                        mes,
                        "ERROR_TECNICO",
                        str(
                            error
                        ),
                        mapa_x1[
                            mes
                        ],
                    )

                    guardar_estado(
                        y_por_mes,
                        x1_por_mes,
                        auditoria_y,
                        auditoria_x1,
                    )

                    escribir_resumen(
                        y_por_mes,
                        x1_por_mes,
                        auditoria_y,
                        auditoria_x1,
                        estado_final=(
                            "DETENIDO_ERROR_TECNICO_X1"
                        ),
                    )

                    raise

                except ErrorEstructuralProduccion as error:

                    x1_por_mes[
                        mes
                    ] = []

                    auditoria_x1[
                        mes
                    ] = auditoria_error(
                        "X1",
                        mes,
                        "ERROR_ESTRUCTURAL",
                        str(
                            error
                        ),
                        mapa_x1[
                            mes
                        ],
                    )

                    guardar_estado(
                        y_por_mes,
                        x1_por_mes,
                        auditoria_y,
                        auditoria_x1,
                    )

                    escribir_resumen(
                        y_por_mes,
                        x1_por_mes,
                        auditoria_y,
                        auditoria_x1,
                        estado_final=(
                            "DETENIDO_ERROR_ESTRUCTURAL_X1"
                        ),
                    )

                    raise

                guardar_estado(
                    y_por_mes,
                    x1_por_mes,
                    auditoria_y,
                    auditoria_x1,
                )

                escribir_resumen(
                    y_por_mes,
                    x1_por_mes,
                    auditoria_y,
                    auditoria_x1,
                )

                if not checkpoint_x1_completo(
                    mes,
                    x1_por_mes,
                    auditoria_x1,
                ):

                    if (
                        auditoria_x1[
                            mes
                        ][
                            "estado"
                        ]
                        !=
                        "MES_SIN_DATO"
                    ):

                        raise ErrorEstructuralProduccion(
                            f"X1 {mes}: checkpoint "
                            "recién guardado no es coherente."
                        )

            guardar_estado(
                y_por_mes,
                x1_por_mes,
                auditoria_y,
                auditoria_x1,
            )

            escribir_resumen(
                y_por_mes,
                x1_por_mes,
                auditoria_y,
                auditoria_x1,
            )

        # ----------------------------------------------------
        # VALIDACIÓN FINAL
        # ----------------------------------------------------

        control_final = validar_final(
            y_por_mes,
            x1_por_mes,
            auditoria_y,
            auditoria_x1,
        )

        guardar_estado(
            y_por_mes,
            x1_por_mes,
            auditoria_y,
            auditoria_x1,
        )

        escribir_resumen(
            y_por_mes,
            x1_por_mes,
            auditoria_y,
            auditoria_x1,
            estado_final=(
                "EXTRACCION_Y_X1_COMPLETA_VALIDADA"
            ),
        )

        print()

        print(
            "=" * 100
        )

        print(
            "VALIDACIÓN FINAL"
        )

        print(
            "=" * 100
        )

        print(
            f"Meses exactos: "
            f"{control_final['meses']}"
        )

        print(
            f"Filas Y: "
            f"{control_final['filas_y']}"
        )

        print(
            f"Filas X1: "
            f"{control_final['filas_x1']}"
        )

        print(
            "Duplicados Y mes+banco_original: 0"
        )

        print(
            "Duplicados X1 mes+banco_original: 0"
        )

        print()

        print(
            "ESTADO = "
            "EXTRACCION_Y_X1_COMPLETA_VALIDADA"
        )

    except DatosFaltantesProduccion:

        try:

            guardar_estado(
                y_por_mes,
                x1_por_mes,
                auditoria_y,
                auditoria_x1,
            )

            escribir_resumen(
                y_por_mes,
                x1_por_mes,
                auditoria_y,
                auditoria_x1,
                estado_final=(
                    "COMPLETADO_CON_DATOS_FALTANTES"
                ),
            )

        except Exception:

            pass

        raise

    except Exception:

        try:

            guardar_estado(
                y_por_mes,
                x1_por_mes,
                auditoria_y,
                auditoria_x1,
            )

            escribir_resumen(
                y_por_mes,
                x1_por_mes,
                auditoria_y,
                auditoria_x1,
                estado_final=(
                    "DETENIDO_CON_ERROR"
                ),
            )

        except Exception:

            pass

        raise


if __name__ == "__main__":

    main()