# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Parser piloto robusto - Movimiento de los Depósitos SBS B-2318

from pathlib import Path
import re
import unicodedata

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. CONFIGURACIÓN
# ============================================================

RAIZ = Path(__file__).resolve().parent.parent

CARPETA_PILOTO = (
    RAIZ
    / "datos_crudos"
    / "sbs_piloto"
)

ARCHIVOS_CONTROL = [
    "depositos_2008_06.xls",
    "depositos_2010_12.xls",
    "depositos_2011_12.xls",
    "depositos_2015_06.xls",
    "depositos_2020_06.xls",
    "depositos_2025_12.xls",
]


# ============================================================
# 2. NORMALIZACIÓN DE TEXTO
# ============================================================

def normalizar_texto(valor):
    """
    Normalización para búsqueda y comparación.

    NO modifica el nombre original almacenado.
    """

    if valor is None:
        return ""

    texto = str(valor)

    texto = texto.replace("\n", " ")
    texto = texto.replace("\r", " ")

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


def normalizar_comparacion(valor):
    """
    Normalización un poco más fuerte utilizada
    exclusivamente para emparejar nombres.

    El valor original se conserva aparte.
    """

    texto = normalizar_texto(valor)

    texto = unicodedata.normalize(
        "NFKD",
        texto
    )

    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(caracter)
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip()


# ============================================================
# 3. FECHA DESDE EL NOMBRE DEL ARCHIVO
# ============================================================

def extraer_fecha_archivo(nombre):

    coincidencia = re.search(
        r"depositos_(\d{4})_(\d{2})",
        nombre
    )

    if not coincidencia:
        return None

    year = int(
        coincidencia.group(1)
    )

    month = int(
        coincidencia.group(2)
    )

    return f"{year:04d}-{month:02d}"


# ============================================================
# 4. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato(ruta):

    with open(ruta, "rb") as archivo:
        cabecera = archivo.read(8)

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:
        return "XLS_OLE"

    if cabecera.startswith(
        firmas_zip
    ):
        return "XLSX_ZIP"

    return "DESCONOCIDO"


# ============================================================
# 5. LEER XLS/OLE
# ============================================================

def leer_xls(ruta):

    libro = xlrd.open_workbook(
        filename=str(ruta),
        on_demand=True
    )

    hojas = []

    for nombre in libro.sheet_names():

        hoja = libro.sheet_by_name(
            nombre
        )

        filas = []

        for i in range(
            hoja.nrows
        ):

            fila = [
                hoja.cell_value(i, j)
                for j in range(
                    hoja.ncols
                )
            ]

            filas.append(
                fila
            )

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.release_resources()

    return hojas


# ============================================================
# 6. LEER XLSX/ZIP
# ============================================================

def leer_xlsx(ruta):

    with open(
        ruta,
        "rb"
    ) as archivo_binario:

        libro = load_workbook(
            archivo_binario,
            read_only=True,
            data_only=True
        )

        hojas = []

        for nombre in libro.sheetnames:

            hoja = libro[
                nombre
            ]

            filas = [
                list(fila)
                for fila in hoja.iter_rows(
                    values_only=True
                )
            ]

            hojas.append({
                "nombre": nombre,
                "filas": filas,
            })

        libro.close()

    return hojas


# ============================================================
# 7. LEER ARCHIVO
# ============================================================

def leer_archivo(ruta):

    formato = detectar_formato(
        ruta
    )

    if formato == "XLS_OLE":

        hojas = leer_xls(
            ruta
        )

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(
            ruta
        )

    else:

        raise ValueError(
            f"Formato no reconocido: {ruta.name}"
        )

    return formato, hojas


# ============================================================
# 8. CONVERTIR NÚMERO
# ============================================================

def convertir_numero(valor):

    if valor is None:
        return None

    if isinstance(
        valor,
        (int, float)
    ):
        return float(valor)

    texto = str(valor).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "nd",
        "n/a",
    }:
        return None

    try:

        return float(
            texto.replace(",", "")
        )

    except ValueError:

        return None


# ============================================================
# 9. DETECTAR FILA Y COLUMNAS DE LA TABLA
# ============================================================

def detectar_encabezado_tabla(filas):
    """
    Busca una fila que contenga simultáneamente:

        Empresas
        Saldo Final

    No importa en qué fila ni en qué columnas estén.
    """

    candidatos = []

    for numero_fila, fila in enumerate(
        filas
    ):

        columna_empresa = None
        columna_saldo_final = None

        for numero_columna, valor in enumerate(
            fila
        ):

            texto = normalizar_comparacion(
                valor
            )

            if (
                texto in {
                    "empresa",
                    "empresas",
                }
                or texto.startswith(
                    "empresas "
                )
            ):

                columna_empresa = (
                    numero_columna
                )

            if (
                "saldo final"
                in texto
            ):

                columna_saldo_final = (
                    numero_columna
                )

        if (
            columna_empresa is not None
            and columna_saldo_final is not None
        ):

            candidatos.append({
                "fila_encabezado": numero_fila,
                "columna_empresa": columna_empresa,
                "columna_saldo_final": columna_saldo_final,
            })

    if not candidatos:
        return None

    # Si hubiera más de un candidato en una hoja,
    # tomamos el primero y luego validaremos sus datos.
    return candidatos[0]


# ============================================================
# 10. OBTENER TEXTO SUPERIOR DE UNA TABLA
# ============================================================

def obtener_contexto_superior(
    filas,
    fila_encabezado,
    max_filas=8
):
    """
    Reúne texto de las filas anteriores al encabezado.
    Sirve para identificar si la tabla corresponde
    a MN o ME por su contenido.
    """

    inicio = max(
        0,
        fila_encabezado - max_filas
    )

    textos = []

    for fila in filas[
        inicio:fila_encabezado
    ]:

        for valor in fila:

            texto = normalizar_comparacion(
                valor
            )

            if texto:
                textos.append(
                    texto
                )

    return " ".join(
        textos
    )


# ============================================================
# 11. CLASIFICAR TABLA COMO MN O ME
# ============================================================

def clasificar_moneda(
    filas,
    encabezado
):

    contexto = obtener_contexto_superior(
        filas,
        encabezado["fila_encabezado"]
    )

    tiene_mn = (
        "moneda nacional"
        in contexto
    )

    tiene_me = (
        "moneda extranjera"
        in contexto
    )

    if (
        tiene_mn
        and not tiene_me
    ):
        return "MN"

    if (
        tiene_me
        and not tiene_mn
    ):
        return "ME"

    return None


# ============================================================
# 12. DETECTAR TABLAS MN Y ME
# ============================================================

def detectar_tablas_monetarias(
    hojas
):

    candidatos = []

    for hoja in hojas:

        encabezado = (
            detectar_encabezado_tabla(
                hoja["filas"]
            )
        )

        if encabezado is None:
            continue

        moneda = clasificar_moneda(
            hoja["filas"],
            encabezado
        )

        if moneda is None:
            continue

        candidatos.append({
            "moneda": moneda,
            "nombre_hoja": hoja["nombre"],
            "filas": hoja["filas"],
            "encabezado": encabezado,
        })

    tablas_mn = [
        x
        for x in candidatos
        if x["moneda"] == "MN"
    ]

    tablas_me = [
        x
        for x in candidatos
        if x["moneda"] == "ME"
    ]

    return tablas_mn, tablas_me


# ============================================================
# 13. DETECTAR AGREGADO
# ============================================================

def es_agregado(nombre):

    texto = normalizar_comparacion(
        nombre
    )

    return (
        "total banca multiple"
        in texto
    )


# ============================================================
# 14. EXTRAER SALDOS DE UNA TABLA
# ============================================================

def extraer_saldos(tabla):
    """
    Extrae banco_original + Saldo Final.

    Se detiene naturalmente ignorando filas que
    no tengan simultáneamente nombre y saldo numérico.
    """

    filas = tabla[
        "filas"
    ]

    encabezado = tabla[
        "encabezado"
    ]

    fila_inicio = (
        encabezado["fila_encabezado"]
        + 1
    )

    col_empresa = encabezado[
        "columna_empresa"
    ]

    col_saldo = encabezado[
        "columna_saldo_final"
    ]

    registros = []

    for numero_fila in range(
        fila_inicio,
        len(filas)
    ):

        fila = filas[
            numero_fila
        ]

        if (
            col_empresa >= len(fila)
            or col_saldo >= len(fila)
        ):
            continue

        nombre = fila[
            col_empresa
        ]

        saldo = convertir_numero(
            fila[col_saldo]
        )

        nombre_texto = (
            str(nombre).strip()
            if nombre is not None
            else ""
        )

        if not normalizar_texto(
            nombre_texto
        ):
            continue

        if saldo is None:
            continue

        agregado = es_agregado(
            nombre_texto
        )

        registros.append({
            "fila_origen": numero_fila,
            "banco_original": nombre_texto,
            "banco_clave": normalizar_comparacion(
                nombre_texto
            ),
            "saldo_final": saldo,
            "es_agregado": agregado,
        })

    return registros


# ============================================================
# 15. CREAR DICCIONARIO POR BANCO
# ============================================================

def crear_diccionario_bancos(
    registros
):

    resultado = {}

    duplicados = []

    for registro in registros:

        if registro[
            "es_agregado"
        ]:
            continue

        clave = registro[
            "banco_clave"
        ]

        if clave in resultado:

            duplicados.append(
                clave
            )

        else:

            resultado[
                clave
            ] = registro

    return resultado, duplicados


# ============================================================
# 16. EMPAREJAR MN Y ME
# ============================================================

def emparejar_monedas(
    registros_mn,
    registros_me
):

    dic_mn, duplicados_mn = (
        crear_diccionario_bancos(
            registros_mn
        )
    )

    dic_me, duplicados_me = (
        crear_diccionario_bancos(
            registros_me
        )
    )

    claves_mn = set(
        dic_mn.keys()
    )

    claves_me = set(
        dic_me.keys()
    )

    comunes = sorted(
        claves_mn
        & claves_me
    )

    solo_mn = sorted(
        claves_mn
        - claves_me
    )

    solo_me = sorted(
        claves_me
        - claves_mn
    )

    emparejados = []

    for clave in comunes:

        reg_mn = dic_mn[
            clave
        ]

        reg_me = dic_me[
            clave
        ]

        mn = reg_mn[
            "saldo_final"
        ]

        me = reg_me[
            "saldo_final"
        ]

        total_depositos = (
            mn + me
        )

        problemas = []

        if mn < 0:

            problemas.append(
                "MN_NEGATIVO"
            )

        if me < 0:

            problemas.append(
                "ME_NEGATIVO"
            )

        if total_depositos <= 0:

            problemas.append(
                "TOTAL_NO_POSITIVO"
            )

            doldep = None

        else:

            doldep = (
                me
                / total_depositos
                * 100
            )

            if not (
                0 <= doldep <= 100
            ):

                problemas.append(
                    "DOLDEP_FUERA_RANGO"
                )

        emparejados.append({
            "banco_clave": clave,

            "banco_original_mn":
                reg_mn["banco_original"],

            "banco_original_me":
                reg_me["banco_original"],

            "saldo_final_mn":
                mn,

            "saldo_final_me":
                me,

            "total_depositos":
                total_depositos,

            "doldep_pct":
                doldep,

            "problemas":
                problemas,
        })

    return {
        "emparejados": emparejados,
        "solo_mn": solo_mn,
        "solo_me": solo_me,
        "duplicados_mn": duplicados_mn,
        "duplicados_me": duplicados_me,
    }


# ============================================================
# 17. ANALIZAR ARCHIVO
# ============================================================

def analizar_archivo(
    ruta
):

    fecha = extraer_fecha_archivo(
        ruta.name
    )

    formato, hojas = leer_archivo(
        ruta
    )

    tablas_mn, tablas_me = (
        detectar_tablas_monetarias(
            hojas
        )
    )

    print()
    print("=" * 105)

    print(
        "ARCHIVO:",
        ruta.name
    )

    print("=" * 105)

    print(
        "Fecha:",
        fecha
    )

    print(
        "Formato real:",
        formato
    )

    print(
        "Hojas del archivo:",
        [
            hoja["nombre"]
            for hoja in hojas
        ]
    )

    print()
    print(
        "Tablas MN detectadas:",
        len(tablas_mn)
    )

    for tabla in tablas_mn:

        print(
            "  ->",
            repr(
                tabla["nombre_hoja"]
            ),
            "| encabezado fila",
            tabla["encabezado"][
                "fila_encabezado"
            ],
            "| Empresas col.",
            tabla["encabezado"][
                "columna_empresa"
            ],
            "| Saldo Final col.",
            tabla["encabezado"][
                "columna_saldo_final"
            ]
        )

    print()
    print(
        "Tablas ME detectadas:",
        len(tablas_me)
    )

    for tabla in tablas_me:

        print(
            "  ->",
            repr(
                tabla["nombre_hoja"]
            ),
            "| encabezado fila",
            tabla["encabezado"][
                "fila_encabezado"
            ],
            "| Empresas col.",
            tabla["encabezado"][
                "columna_empresa"
            ],
            "| Saldo Final col.",
            tabla["encabezado"][
                "columna_saldo_final"
            ]
        )

    # --------------------------------------------------------
    # Para este reporte esperamos exactamente
    # una tabla MN y una tabla ME.
    # --------------------------------------------------------

    if (
        len(tablas_mn) != 1
        or len(tablas_me) != 1
    ):

        print()
        print(
            "[ERROR] No se detectó exactamente "
            "una tabla MN y una tabla ME."
        )

        return None

    tabla_mn = tablas_mn[0]
    tabla_me = tablas_me[0]

    registros_mn = extraer_saldos(
        tabla_mn
    )

    registros_me = extraer_saldos(
        tabla_me
    )

    print()
    print(
        "Registros extraídos MN:",
        len(registros_mn)
    )

    print(
        "Registros extraídos ME:",
        len(registros_me)
    )

    # --------------------------------------------------------
    # Mostrar agregado detectado
    # --------------------------------------------------------

    agregados_mn = [
        r
        for r in registros_mn
        if r["es_agregado"]
    ]

    agregados_me = [
        r
        for r in registros_me
        if r["es_agregado"]
    ]

    print()
    print(
        "Agregados MN detectados:",
        [
            r["banco_original"]
            for r in agregados_mn
        ]
    )

    print(
        "Agregados ME detectados:",
        [
            r["banco_original"]
            for r in agregados_me
        ]
    )

    resultado = emparejar_monedas(
        registros_mn,
        registros_me
    )

    print()
    print(
        "Bancos emparejados:",
        len(
            resultado["emparejados"]
        )
    )

    print(
        "Solo en MN:",
        resultado["solo_mn"]
    )

    print(
        "Solo en ME:",
        resultado["solo_me"]
    )

    print(
        "Duplicados MN:",
        resultado["duplicados_mn"]
    )

    print(
        "Duplicados ME:",
        resultado["duplicados_me"]
    )

    print()
    print(
        "RESULTADOS BANCO POR BANCO"
    )

    print("-" * 105)

    for registro in resultado[
        "emparejados"
    ]:

        print()
        print(
            "Banco MN:",
            repr(
                registro[
                    "banco_original_mn"
                ]
            )
        )

        print(
            "Banco ME:",
            repr(
                registro[
                    "banco_original_me"
                ]
            )
        )

        print(
            "  Saldo Final MN:",
            registro[
                "saldo_final_mn"
            ]
        )

        print(
            "  Saldo Final ME:",
            registro[
                "saldo_final_me"
            ]
        )

        print(
            "  Total depósitos:",
            registro[
                "total_depositos"
            ]
        )

        print(
            "  DOLDEP (%):",
            registro[
                "doldep_pct"
            ]
        )

        print(
            "  Problemas:",
            registro[
                "problemas"
            ]
        )

    return {
        "archivo": ruta.name,
        "fecha": fecha,
        "hoja_mn": tabla_mn[
            "nombre_hoja"
        ],
        "hoja_me": tabla_me[
            "nombre_hoja"
        ],
        "registros_mn": registros_mn,
        "registros_me": registros_me,
        "resultado": resultado,
    }


# ============================================================
# 18. FUNCIÓN PRINCIPAL
# ============================================================

def main():

    print()
    print("=" * 105)

    print(
        "PARSER PILOTO DE DEPÓSITOS SBS B-2318"
    )

    print(
        "SEIS PUNTOS DE CONTROL 2008-2025"
    )

    print("=" * 105)

    print()
    print(
        "No se descargarán archivos."
    )

    print(
        "No se modificarán archivos crudos."
    )

    print(
        "DOLDEP se calcula únicamente "
        "como validación piloto."
    )

    resultados = []

    for nombre in ARCHIVOS_CONTROL:

        ruta = (
            CARPETA_PILOTO
            / nombre
        )

        if not ruta.exists():

            print()
            print(
                "[ERROR] No existe:",
                ruta
            )

            continue

        resultado = analizar_archivo(
            ruta
        )

        if resultado is not None:

            resultados.append(
                resultado
            )

    # ========================================================
    # RESUMEN
    # ========================================================

    print()
    print("=" * 105)

    print(
        "RESUMEN FINAL DE DEPÓSITOS"
    )

    print("=" * 105)

    for item in resultados:

        resultado = item[
            "resultado"
        ]

        emparejados = resultado[
            "emparejados"
        ]

        problemas = [
            r
            for r in emparejados
            if r["problemas"]
        ]

        doldep_validos = [
            r["doldep_pct"]
            for r in emparejados
            if (
                r["doldep_pct"]
                is not None
                and not r["problemas"]
            )
        ]

        print()
        print(
            item["archivo"]
        )

        print(
            "  Hoja MN:",
            repr(
                item["hoja_mn"]
            )
        )

        print(
            "  Hoja ME:",
            repr(
                item["hoja_me"]
            )
        )

        print(
            "  Bancos emparejados:",
            len(emparejados)
        )

        print(
            "  Solo MN:",
            len(
                resultado["solo_mn"]
            )
        )

        print(
            "  Solo ME:",
            len(
                resultado["solo_me"]
            )
        )

        print(
            "  Duplicados MN:",
            len(
                resultado["duplicados_mn"]
            )
        )

        print(
            "  Duplicados ME:",
            len(
                resultado["duplicados_me"]
            )
        )

        print(
            "  Registros con problemas:",
            len(problemas)
        )

        print(
            "  DOLDEP mínimo:",
            (
                min(doldep_validos)
                if doldep_validos
                else None
            )
        )

        print(
            "  DOLDEP máximo:",
            (
                max(doldep_validos)
                if doldep_validos
                else None
            )
        )

    print()
    print("=" * 105)

    print(
        "FIN DE LA PRUEBA"
    )

    print("=" * 105)


# ============================================================
# 19. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":
    main()