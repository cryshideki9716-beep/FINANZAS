# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Fecha de extracción: 2026-09-24

from pathlib import Path
from urllib.parse import urljoin, urlparse
import csv
import hashlib
import math
import re
import time
import unicodedata

import requests
from bs4 import BeautifulSoup
import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. OBJETIVO
# ============================================================

"""
DIAGNÓSTICO ÚNICO:

    Y = Dolarización del crédito (%)
    SBS B-2359
    Mes: 2015-11

Objetivo:
    Diagnosticar las pequeñas diferencias observadas entre:

        suma de bancos individuales
        vs.
        TOTAL BANCA MÚLTIPLE

para:

    - MN
    - ME
    - Total

SIN:

    - cambiar el parser;
    - cambiar la fórmula de DOLCRED;
    - ampliar tolerancias;
    - declarar Y validada;
    - ejecutar X1;
    - homologar nombres;
    - seleccionar bancos;
    - interpolar datos.

Se comparará:

    - sum()
    - math.fsum()
    - Total SBS

y se inspeccionará el formato Excel de cada celda.
"""


# ============================================================
# 2. IDENTIFICACIÓN
# ============================================================

CODIGO_ESTUDIANTE = "2024200485D"

YEAR = 2015
MONTH = 11
MES = "2015-11"

CODIGO_REPORTE = "B-2359"

NOMBRE_REPORTE = (
    "Créditos Directos por Tipo, Modalidad y Moneda"
)


# ============================================================
# 3. TOLERANCIA EXISTENTE
#
# NO se modifica.
# Solo se registra para diagnóstico.
# ============================================================

REL_TOL = 1e-14
ABS_TOL = 1e-6


# ============================================================
# 4. FUENTE SBS
# ============================================================

URL_REPORTE = (
    "https://www.sbs.gob.pe/app/stats_net/stats/"
    "EstadisticaSistemaFinancieroResultados.aspx"
    "?c=B-2359"
)


CODIGO_MES = {
    1: "en",
    2: "fe",
    3: "ma",
    4: "ab",
    5: "my",
    6: "jn",
    7: "jl",
    8: "ag",
    9: "se",
    10: "oc",
    11: "no",
    12: "di",
}


# ============================================================
# 5. RUTAS
# ============================================================

RAIZ = Path(
    __file__
).resolve().parent.parent


CARPETA_CRUDOS = (
    RAIZ
    / "datos_crudos"
    / "prueba_Y_X1_2015_11"
)


CARPETA_DIAGNOSTICOS = (
    RAIZ
    / "salidas"
    / "diagnosticos"
)


CARPETA_CRUDOS.mkdir(
    parents=True,
    exist_ok=True
)


CARPETA_DIAGNOSTICOS.mkdir(
    parents=True,
    exist_ok=True
)


ARCHIVO_CREDITO = (
    CARPETA_CRUDOS
    / (
        f"Y_B-2359_"
        f"{YEAR}_{MONTH:02d}_"
        f"{CODIGO_ESTUDIANTE}.xls"
    )
)


TXT_RESULTADO = (
    CARPETA_DIAGNOSTICOS
    / "diagnostico_redondeo_Y_2015_11.txt"
)


CSV_CELDAS = (
    CARPETA_DIAGNOSTICOS
    / "diagnostico_redondeo_Y_2015_11_celdas.csv"
)


CSV_CONTROLES = (
    CARPETA_DIAGNOSTICOS
    / "diagnostico_redondeo_Y_2015_11_controles.csv"
)


# ============================================================
# 6. SESIÓN HTTP
# ============================================================

session = requests.Session()

session.headers.update({

    "User-Agent": (
        "Mozilla/5.0 "
        "(compatible; TrabajoAcademicoUNCP/1.0; "
        "Finanzas-I; "
        "2024200485D)"
    )

})


# ============================================================
# 7. SALIDA TXT
# ============================================================

lineas_reporte = []


def escribir(texto=""):

    texto = str(
        texto
    )

    print(
        texto
    )

    lineas_reporte.append(
        texto
    )


def titulo(texto):

    escribir()
    escribir("=" * 118)
    escribir(texto)
    escribir("=" * 118)


def subtitulo(texto):

    escribir()
    escribir("-" * 118)
    escribir(texto)
    escribir("-" * 118)


# ============================================================
# 8. UTILIDADES DE TEXTO
# ============================================================

def normalizar_texto(valor):

    if valor is None:
        return ""

    texto = str(
        valor
    )

    texto = texto.replace(
        "\n",
        " "
    )

    texto = texto.replace(
        "\r",
        " "
    )

    texto = texto.replace(
        "\xa0",
        " "
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


def normalizar_comparacion(valor):

    texto = normalizar_texto(
        valor
    )

    texto = unicodedata.normalize(
        "NFKD",
        texto
    )

    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(
            caracter
        )
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip()


def sin_tildes(valor):

    return normalizar_comparacion(
        valor
    )


# ============================================================
# 9. CONVERSIÓN NUMÉRICA
#
# MISMA LÓGICA DEL PARSER VALIDADO.
# ============================================================

def convertir_numero(valor):

    if valor is None:
        return None

    if isinstance(
        valor,
        (int, float)
    ):

        return float(
            valor
        )

    texto = str(
        valor
    ).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "n.d",
        "nd",
        "n/a",
        "na",
    }:

        return None

    try:

        return float(
            texto.replace(
                ",",
                ""
            )
        )

    except ValueError:

        return None


# ============================================================
# 10. LOCALIZAR ARCHIVO SBS
# ============================================================

def consultar_reporte():

    respuesta = session.get(
        URL_REPORTE,
        timeout=45
    )

    respuesta.raise_for_status()

    soup = BeautifulSoup(
        respuesta.text,
        "lxml"
    )

    return (
        soup,
        respuesta.status_code,
    )


def localizar_enlaces_xls(
    soup
):

    enlaces = []

    urls_vistas = set()

    for enlace in soup.find_all(
        "a",
        href=True
    ):

        href = enlace.get(
            "href",
            ""
        ).strip()

        if ".xls" not in href.lower():
            continue

        url_xls = urljoin(
            URL_REPORTE,
            href
        )

        if url_xls in urls_vistas:
            continue

        urls_vistas.add(
            url_xls
        )

        enlaces.append({
            "texto":
                enlace.get_text(
                    " ",
                    strip=True
                ),

            "href_original":
                href,

            "url":
                url_xls,
        })

    return enlaces


def identificador_mes():

    return (
        f"{CODIGO_REPORTE}-"
        f"{CODIGO_MES[MONTH]}"
        f"{YEAR}"
    ).lower()


def localizar_archivo_oficial():

    soup, http = consultar_reporte()

    escribir(
        f"Página oficial: {URL_REPORTE}"
    )

    escribir(
        f"HTTP página: {http}"
    )

    enlaces = localizar_enlaces_xls(
        soup
    )

    esperado = identificador_mes()

    coincidencias = []

    for item in enlaces:

        contenido = (
            item[
                "url"
            ]
            + " "
            + item[
                "href_original"
            ]
        ).lower()

        if esperado in contenido:

            coincidencias.append(
                item
            )

    escribir(
        f"Identificador esperado: {esperado}"
    )

    escribir(
        f"Coincidencias: {len(coincidencias)}"
    )

    for item in coincidencias:

        escribir(
            f"  -> {item['url']}"
        )

    if len(
        coincidencias
    ) != 1:

        raise RuntimeError(
            "No se encontró exactamente un "
            "archivo B-2359 para 2015-11."
        )

    return coincidencias[
        0
    ]


# ============================================================
# 11. DESCARGAR / REUTILIZAR ARCHIVO
# ============================================================

def descargar_archivo(
    url,
    destino
):

    respuesta = session.get(
        url,
        timeout=60
    )

    respuesta.raise_for_status()

    contenido = respuesta.content

    inicio = contenido[
        :300
    ].lower()

    if (
        b"<html" in inicio
        or
        b"<!doctype html" in inicio
    ):

        raise RuntimeError(
            "SBS devolvió HTML en vez del Excel."
        )

    temporal = destino.with_suffix(
        destino.suffix
        + ".tmp"
    )

    with temporal.open(
        "wb"
    ) as archivo:

        archivo.write(
            contenido
        )

    temporal.replace(
        destino
    )

    sha256 = hashlib.sha256(
        contenido
    ).hexdigest()

    return {
        "http":
            respuesta.status_code,

        "bytes":
            len(
                contenido
            ),

        "sha256":
            sha256,

        "nombre_origen":
            Path(
                urlparse(
                    url
                ).path
            ).name,
    }


# ============================================================
# 12. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato(
    ruta
):

    with ruta.open(
        "rb"
    ) as archivo:

        cabecera = archivo.read(
            8
        )

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:

        return "XLS_OLE"

    if cabecera.startswith(
        firmas_zip
    ):

        return "XLSX_ZIP"

    return "DESCONOCIDO"


# ============================================================
# 13. LECTURA NORMAL DEL PARSER
#
# NO se cambia.
# ============================================================

def leer_xls(
    ruta
):

    libro = xlrd.open_workbook(
        filename=str(
            ruta
        ),
        on_demand=True
    )

    hojas = []

    for nombre in libro.sheet_names():

        hoja = libro.sheet_by_name(
            nombre
        )

        filas = []

        for i in range(
            hoja.nrows
        ):

            fila = [
                hoja.cell_value(
                    i,
                    j
                )
                for j in range(
                    hoja.ncols
                )
            ]

            filas.append(
                fila
            )

        hojas.append({
            "nombre":
                nombre,

            "filas":
                filas,
        })

    libro.release_resources()

    return hojas


def leer_xlsx(
    ruta
):

    with ruta.open(
        "rb"
    ) as archivo_binario:

        libro = load_workbook(
            archivo_binario,
            read_only=True,
            data_only=True
        )

        hojas = []

        for nombre in libro.sheetnames:

            hoja = libro[
                nombre
            ]

            filas = [
                list(
                    fila
                )
                for fila in hoja.iter_rows(
                    values_only=True
                )
            ]

            hojas.append({
                "nombre":
                    nombre,

                "filas":
                    filas,
            })

        libro.close()

    return hojas


def leer_archivo(
    ruta
):

    formato = detectar_formato(
        ruta
    )

    if formato == "XLS_OLE":

        hojas = leer_xls(
            ruta
        )

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(
            ruta
        )

    else:

        raise RuntimeError(
            f"Formato desconocido: {formato}"
        )

    return (
        formato,
        hojas,
    )


# ============================================================
# 14. RECONOCER MN / ME / TOTAL
#
# MISMO PARSER.
# ============================================================

def es_mn(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith(
            "mn"
        )
        and
        "mil" in texto
    )


def es_me(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith(
            "me"
        )
        and
        "mil" in texto
    )


def es_total(valor):

    texto = normalizar_texto(
        valor
    )

    return texto.startswith(
        "total"
    )


def encontrar_fila_encabezados_creditos(
    filas
):

    candidatos = []

    for numero_fila, fila in enumerate(
        filas
    ):

        bloques = []

        for columna in range(
            max(
                0,
                len(fila) - 2
            )
        ):

            if (
                es_mn(
                    fila[
                        columna
                    ]
                )
                and
                es_me(
                    fila[
                        columna + 1
                    ]
                )
                and
                es_total(
                    fila[
                        columna + 2
                    ]
                )
            ):

                bloques.append(
                    columna
                )

        if bloques:

            candidatos.append({
                "fila":
                    numero_fila,

                "bloques":
                    bloques,
            })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x:
        len(
            x[
                "bloques"
            ]
        )
    )


def obtener_nombre_banco_creditos(
    filas,
    fila_bancos,
    columna_inicio
):

    if (
        fila_bancos < 0
        or
        fila_bancos >= len(
            filas
        )
    ):

        return ""

    fila = filas[
        fila_bancos
    ]

    for desplazamiento in (
        0,
        1,
        2,
    ):

        columna = (
            columna_inicio
            +
            desplazamiento
        )

        if columna >= len(
            fila
        ):
            continue

        valor = fila[
            columna
        ]

        if normalizar_texto(
            valor
        ):

            return str(
                valor
            ).strip()

    return ""


def detectar_estructura_creditos(
    filas
):

    encabezado = encontrar_fila_encabezados_creditos(
        filas
    )

    if encabezado is None:
        return None

    fila_encabezados = encabezado[
        "fila"
    ]

    fila_bancos = (
        fila_encabezados
        -
        1
    )

    bloques = []

    for columna_inicio in encabezado[
        "bloques"
    ]:

        banco = obtener_nombre_banco_creditos(
            filas,
            fila_bancos,
            columna_inicio
        )

        bloques.append({
            "banco_original":
                banco,

            "columna_mn":
                columna_inicio,

            "columna_me":
                columna_inicio + 1,

            "columna_total":
                columna_inicio + 2,
        })

    return {
        "fila_bancos":
            fila_bancos,

        "fila_encabezados":
            fila_encabezados,

        "bloques":
            bloques,
    }


def elegir_hoja_principal_creditos(
    hojas
):

    candidatos = []

    for hoja in hojas:

        estructura = detectar_estructura_creditos(
            hoja[
                "filas"
            ]
        )

        if estructura is None:
            continue

        candidatos.append({
            "hoja":
                hoja,

            "estructura":
                estructura,
        })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x:
        len(
            x[
                "estructura"
            ][
                "bloques"
            ]
        )
    )


def buscar_fila_total_creditos(
    filas
):

    for numero_fila, fila in enumerate(
        filas
    ):

        for valor in fila:

            texto = sin_tildes(
                valor
            )

            if (
                "total creditos"
                in texto
            ):

                return numero_fila

    return None


def es_agregado_creditos(
    nombre
):

    return (
        "total banca multiple"
        in
        sin_tildes(
            nombre
        )
    )


# ============================================================
# 15. PARSER NORMAL DE Y
#
# FÓRMULA SIN CAMBIOS.
# ============================================================

def extraer_y_creditos(
    ruta
):

    formato, hojas = leer_archivo(
        ruta
    )

    seleccionado = elegir_hoja_principal_creditos(
        hojas
    )

    if seleccionado is None:

        raise RuntimeError(
            "No se encontró estructura MN|ME|Total."
        )

    hoja = seleccionado[
        "hoja"
    ]

    estructura = seleccionado[
        "estructura"
    ]

    filas = hoja[
        "filas"
    ]

    fila_total = buscar_fila_total_creditos(
        filas
    )

    if fila_total is None:

        raise RuntimeError(
            "No se encontró 'Total Créditos:'."
        )

    fila = filas[
        fila_total
    ]

    registros = []

    for bloque in estructura[
        "bloques"
    ]:

        banco = bloque[
            "banco_original"
        ]

        col_mn = bloque[
            "columna_mn"
        ]

        col_me = bloque[
            "columna_me"
        ]

        col_total = bloque[
            "columna_total"
        ]

        mn = (
            convertir_numero(
                fila[
                    col_mn
                ]
            )
            if col_mn < len(
                fila
            )
            else None
        )

        me = (
            convertir_numero(
                fila[
                    col_me
                ]
            )
            if col_me < len(
                fila
            )
            else None
        )

        total = (
            convertir_numero(
                fila[
                    col_total
                ]
            )
            if col_total < len(
                fila
            )
            else None
        )

        problemas = []

        if not banco:

            problemas.append(
                "BANCO_SIN_NOMBRE"
            )

        if mn is None:

            problemas.append(
                "MN_FALTANTE"
            )

        if me is None:

            problemas.append(
                "ME_FALTANTE"
            )

        if total is None:

            problemas.append(
                "TOTAL_FALTANTE"
            )

        if (
            mn is not None
            and
            mn < 0
        ):

            problemas.append(
                "MN_NEGATIVO"
            )

        if (
            me is not None
            and
            me < 0
        ):

            problemas.append(
                "ME_NEGATIVO"
            )

        if (
            total is not None
            and
            total <= 0
        ):

            problemas.append(
                "TOTAL_NO_POSITIVO"
            )

        componente_me_soles = None
        dolcred = None
        tc_implicito = None

        if (
            mn is not None
            and
            total is not None
        ):

            if total < mn:

                problemas.append(
                    "TOTAL_MENOR_QUE_MN"
                )

            componente_me_soles = (
                total
                -
                mn
            )

        # ====================================================
        # FÓRMULA VALIDADA - SIN CAMBIOS
        # ====================================================

        if (
            total is not None
            and
            total > 0
            and
            mn is not None
        ):

            dolcred = (
                (
                    total
                    -
                    mn
                )
                /
                total
                *
                100
            )

            if not (
                0
                <=
                dolcred
                <=
                100
            ):

                problemas.append(
                    "DOLCRED_FUERA_RANGO"
                )

        if (
            mn is not None
            and
            me is not None
            and
            total is not None
            and
            me > 0
        ):

            tc_implicito = (
                (
                    total
                    -
                    mn
                )
                /
                me
            )

        registros.append({
            "banco_original":
                banco,

            "mn_soles_miles":
                mn,

            "me_usd_miles":
                me,

            "total_soles_miles":
                total,

            "dolcred_pct":
                dolcred,

            "tc_implicito":
                tc_implicito,

            "es_agregado":
                es_agregado_creditos(
                    banco
                ),

            "problemas":
                problemas,

            # Posiciones reales para el diagnóstico.
            "columna_mn":
                col_mn,

            "columna_me":
                col_me,

            "columna_total":
                col_total,
        })

    return {
        "formato":
            formato,

        "hoja":
            hoja[
                "nombre"
            ],

        "fila_encabezados":
            estructura[
                "fila_encabezados"
            ],

        "fila_bancos":
            estructura[
                "fila_bancos"
            ],

        "fila_total_creditos":
            fila_total,

        "cantidad_bloques":
            len(
                estructura[
                    "bloques"
                ]
            ),

        "registros":
            registros,
    }


# ============================================================
# 16. INFERIR DECIMALES DEL FORMATO EXCEL
# ============================================================

def inferir_decimales_formato_excel(
    formato
):
    """
    Intenta inferir cuántos decimales muestra
    el formato Excel.

    Ejemplos:
        #,##0       -> 0
        #,##0.0     -> 1
        #,##0.00    -> 2
        General     -> None

    Si el formato es complejo/indeterminado,
    devuelve None.
    """

    if formato is None:
        return None

    texto = str(
        formato
    ).strip()

    if not texto:
        return None

    if texto.lower() == "general":
        return None

    # Solo primera sección:
    # positivo;negativo;cero;texto
    seccion = texto.split(
        ";"
    )[0]

    # Quitar textos entre comillas.
    seccion = re.sub(
        r'"[^"]*"',
        "",
        seccion
    )

    # Quitar condiciones/colores [Red], [>=0], etc.
    seccion = re.sub(
        r"\[[^\]]*\]",
        "",
        seccion
    )

    coincidencia = re.search(
        r"\.([0#?]+)",
        seccion
    )

    if coincidencia:

        return len(
            coincidencia.group(1)
        )

    # Si contiene placeholder numérico
    # pero ninguna parte decimal.
    if re.search(
        r"[0#?]",
        seccion
    ):

        return 0

    return None


# ============================================================
# 17. METADATOS DE CELDA XLS/OLE CON XLRD
# ============================================================

def leer_metadatos_celdas_xls(
    ruta,
    nombre_hoja,
    fila,
    registros
):
    """
    Reabre el mismo XLS exclusivamente para
    inspeccionar formato Excel.

    NO sustituye ni modifica el parser.
    """

    libro = xlrd.open_workbook(
        filename=str(
            ruta
        ),
        formatting_info=True,
        on_demand=True
    )

    hoja = libro.sheet_by_name(
        nombre_hoja
    )

    salida = []

    for registro in registros:

        for variable, columna in (
            (
                "MN",
                registro[
                    "columna_mn"
                ]
            ),
            (
                "ME",
                registro[
                    "columna_me"
                ]
            ),
            (
                "TOTAL",
                registro[
                    "columna_total"
                ]
            ),
        ):

            celda = hoja.cell(
                fila,
                columna
            )

            try:

                xf_index = hoja.cell_xf_index(
                    fila,
                    columna
                )

            except Exception:

                xf_index = None

            format_key = None
            format_str = None

            if (
                xf_index is not None
                and
                0 <= xf_index < len(
                    libro.xf_list
                )
            ):

                xf = libro.xf_list[
                    xf_index
                ]

                format_key = getattr(
                    xf,
                    "format_key",
                    None
                )

                if format_key is not None:

                    formato_obj = libro.format_map.get(
                        format_key
                    )

                    if formato_obj is not None:

                        format_str = getattr(
                            formato_obj,
                            "format_str",
                            None
                        )

            salida.append({
                "banco_original":
                    registro[
                        "banco_original"
                    ],

                "es_agregado":
                    registro[
                        "es_agregado"
                    ],

                "variable":
                    variable,

                "fila_excel_0based":
                    fila,

                "columna_excel_0based":
                    columna,

                "valor_parser":
                    (
                        registro[
                            "mn_soles_miles"
                        ]
                        if variable == "MN"
                        else
                        registro[
                            "me_usd_miles"
                        ]
                        if variable == "ME"
                        else
                        registro[
                            "total_soles_miles"
                        ]
                    ),

                "valor_raw":
                    celda.value,

                "valor_raw_repr":
                    repr(
                        celda.value
                    ),

                "ctype_xlrd":
                    celda.ctype,

                "xf_index":
                    xf_index,

                "format_key":
                    format_key,

                "format_str":
                    format_str,

                "decimales_formato":
                    inferir_decimales_formato_excel(
                        format_str
                    ),
            })

    libro.release_resources()

    return salida


# ============================================================
# 18. METADATOS DE CELDA XLSX/ZIP
# ============================================================

def leer_metadatos_celdas_xlsx(
    ruta,
    nombre_hoja,
    fila,
    registros
):

    with ruta.open(
        "rb"
    ) as archivo_binario:

        libro = load_workbook(
            archivo_binario,
            read_only=False,
            data_only=True
        )

        hoja = libro[
            nombre_hoja
        ]

        salida = []

        for registro in registros:

            for variable, columna in (
                (
                    "MN",
                    registro[
                        "columna_mn"
                    ]
                ),
                (
                    "ME",
                    registro[
                        "columna_me"
                    ]
                ),
                (
                    "TOTAL",
                    registro[
                        "columna_total"
                    ]
                ),
            ):

                celda = hoja.cell(
                    row=fila + 1,
                    column=columna + 1
                )

                formato = celda.number_format

                salida.append({
                    "banco_original":
                        registro[
                            "banco_original"
                        ],

                    "es_agregado":
                        registro[
                            "es_agregado"
                        ],

                    "variable":
                        variable,

                    "fila_excel_0based":
                        fila,

                    "columna_excel_0based":
                        columna,

                    "valor_parser":
                        (
                            registro[
                                "mn_soles_miles"
                            ]
                            if variable == "MN"
                            else
                            registro[
                                "me_usd_miles"
                            ]
                            if variable == "ME"
                            else
                            registro[
                                "total_soles_miles"
                            ]
                        ),

                    "valor_raw":
                        celda.value,

                    "valor_raw_repr":
                        repr(
                            celda.value
                        ),

                    "ctype_xlrd":
                        "",

                    "xf_index":
                        celda.style_id,

                    "format_key":
                        "",

                    "format_str":
                        formato,

                    "decimales_formato":
                        inferir_decimales_formato_excel(
                            formato
                        ),
                })

        libro.close()

    return salida


# ============================================================
# 19. LEER METADATOS SEGÚN FORMATO REAL
# ============================================================

def leer_metadatos_excel(
    ruta,
    resultado
):

    formato = resultado[
        "formato"
    ]

    if formato == "XLS_OLE":

        return leer_metadatos_celdas_xls(
            ruta,
            resultado[
                "hoja"
            ],
            resultado[
                "fila_total_creditos"
            ],
            resultado[
                "registros"
            ]
        )

    if formato == "XLSX_ZIP":

        return leer_metadatos_celdas_xlsx(
            ruta,
            resultado[
                "hoja"
            ],
            resultado[
                "fila_total_creditos"
            ],
            resultado[
                "registros"
            ]
        )

    raise RuntimeError(
        "No puede inspeccionarse formato "
        "de un archivo desconocido."
    )


# ============================================================
# 20. DIAGNÓSTICO DE SUMAS
# ============================================================

def diagnosticar_sumas(
    registros
):

    agregados = [
        registro
        for registro in registros
        if registro[
            "es_agregado"
        ]
    ]

    bancos = [
        registro
        for registro in registros
        if not registro[
            "es_agregado"
        ]
    ]

    if len(
        agregados
    ) != 1:

        raise RuntimeError(
            "Se esperaba exactamente un "
            "TOTAL BANCA MÚLTIPLE."
        )

    agregado = agregados[
        0
    ]

    controles = []

    configuracion = [
        (
            "MN",
            "mn_soles_miles"
        ),
        (
            "ME",
            "me_usd_miles"
        ),
        (
            "TOTAL",
            "total_soles_miles"
        ),
    ]

    for nombre_variable, campo in configuracion:

        valores = [
            registro[
                campo
            ]
            for registro in bancos
        ]

        if any(
            valor is None
            for valor in valores
        ):

            raise RuntimeError(
                f"{nombre_variable}: existe "
                "al menos un banco sin valor."
            )

        total_sbs = agregado[
            campo
        ]

        if total_sbs is None:

            raise RuntimeError(
                f"{nombre_variable}: agregado "
                "SBS sin valor."
            )

        suma_sum = sum(
            valores
        )

        suma_fsum = math.fsum(
            valores
        )

        diferencia_sum = (
            suma_sum
            -
            total_sbs
        )

        diferencia_fsum = (
            suma_fsum
            -
            total_sbs
        )

        diferencia_entre_sumas = (
            suma_sum
            -
            suma_fsum
        )

        controles.append({
            "variable":
                nombre_variable,

            "cantidad_bancos":
                len(
                    bancos
                ),

            "suma_sum":
                suma_sum,

            "suma_fsum":
                suma_fsum,

            "total_sbs":
                total_sbs,

            "diferencia_sum":
                diferencia_sum,

            "diferencia_abs_sum":
                abs(
                    diferencia_sum
                ),

            "diferencia_fsum":
                diferencia_fsum,

            "diferencia_abs_fsum":
                abs(
                    diferencia_fsum
                ),

            "diferencia_sum_vs_fsum":
                diferencia_entre_sumas,

            "diferencia_abs_sum_vs_fsum":
                abs(
                    diferencia_entre_sumas
                ),

            # Solo se informa.
            # NO valida ni cambia tolerancia.
            "isclose_sum_tolerancia_actual":
                math.isclose(
                    suma_sum,
                    total_sbs,
                    rel_tol=REL_TOL,
                    abs_tol=ABS_TOL
                ),

            "isclose_fsum_tolerancia_actual":
                math.isclose(
                    suma_fsum,
                    total_sbs,
                    rel_tol=REL_TOL,
                    abs_tol=ABS_TOL
                ),
        })

    return (
        bancos,
        agregado,
        controles,
    )


# ============================================================
# 21. GUARDAR CSV DE CELDAS
# ============================================================

def guardar_csv_celdas(
    metadatos
):

    columnas = [
        "banco_original",
        "es_agregado",
        "variable",
        "fila_excel_0based",
        "columna_excel_0based",
        "valor_parser",
        "valor_raw",
        "valor_raw_repr",
        "ctype_xlrd",
        "xf_index",
        "format_key",
        "format_str",
        "decimales_formato",
    ]

    with CSV_CELDAS.open(
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as archivo:

        escritor = csv.DictWriter(
            archivo,
            fieldnames=columnas
        )

        escritor.writeheader()

        for fila in metadatos:

            escritor.writerow(
                fila
            )


# ============================================================
# 22. GUARDAR CSV DE CONTROLES
# ============================================================

def guardar_csv_controles(
    controles
):

    columnas = [
        "variable",
        "cantidad_bancos",
        "suma_sum",
        "suma_fsum",
        "total_sbs",
        "diferencia_sum",
        "diferencia_abs_sum",
        "diferencia_fsum",
        "diferencia_abs_fsum",
        "diferencia_sum_vs_fsum",
        "diferencia_abs_sum_vs_fsum",
        "isclose_sum_tolerancia_actual",
        "isclose_fsum_tolerancia_actual",
    ]

    with CSV_CONTROLES.open(
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as archivo:

        escritor = csv.DictWriter(
            archivo,
            fieldnames=columnas
        )

        escritor.writeheader()

        for fila in controles:

            escritor.writerow(
                fila
            )


# ============================================================
# 23. MOSTRAR FORMATOS POR BANCO
# ============================================================

def mostrar_celdas(
    registros,
    metadatos
):

    # Índice para acceder a metadatos:
    # (banco, variable)
    indice = {}

    for item in metadatos:

        indice[
            (
                item[
                    "banco_original"
                ],
                item[
                    "variable"
                ]
            )
        ] = item

    subtitulo(
        "VALORES INDIVIDUALES Y FORMATO EXCEL"
    )

    for registro in registros:

        tipo = (
            "TOTAL BANCA MÚLTIPLE"
            if registro[
                "es_agregado"
            ]
            else
            "BANCO"
        )

        escribir()
        escribir(
            f"{tipo}: "
            f"{registro['banco_original']}"
        )

        for variable, campo in (
            (
                "MN",
                "mn_soles_miles"
            ),
            (
                "ME",
                "me_usd_miles"
            ),
            (
                "TOTAL",
                "total_soles_miles"
            ),
        ):

            meta = indice.get(
                (
                    registro[
                        "banco_original"
                    ],
                    variable
                )
            )

            escribir(
                f"  {variable}:"
            )

            escribir(
                f"    valor parser = "
                f"{registro[campo]!r}"
            )

            if meta is None:

                escribir(
                    "    metadatos Excel = "
                    "<NO RECUPERADOS>"
                )

                continue

            escribir(
                f"    valor raw Excel = "
                f"{meta['valor_raw_repr']}"
            )

            escribir(
                f"    fila 0-based = "
                f"{meta['fila_excel_0based']}"
            )

            escribir(
                f"    columna 0-based = "
                f"{meta['columna_excel_0based']}"
            )

            escribir(
                f"    xf/style = "
                f"{meta['xf_index']}"
            )

            escribir(
                f"    format_key = "
                f"{meta['format_key']}"
            )

            escribir(
                f"    formato Excel = "
                f"{meta['format_str']!r}"
            )

            escribir(
                f"    decimales inferidos "
                f"del formato = "
                f"{meta['decimales_formato']}"
            )


# ============================================================
# 24. MOSTRAR CONTROLES DE SUMA
# ============================================================

def mostrar_controles(
    controles
):

    subtitulo(
        "DIAGNÓSTICO SUMA DE BANCOS VS TOTAL BANCA MÚLTIPLE"
    )

    for control in controles:

        escribir()
        escribir(
            f"VARIABLE: {control['variable']}"
        )

        escribir(
            f"  Cantidad de bancos individuales = "
            f"{control['cantidad_bancos']}"
        )

        escribir(
            f"  Suma con sum() = "
            f"{control['suma_sum']!r}"
        )

        escribir(
            f"  Suma con math.fsum() = "
            f"{control['suma_fsum']!r}"
        )

        escribir(
            f"  Total SBS = "
            f"{control['total_sbs']!r}"
        )

        escribir(
            f"  Diferencia sum() - SBS = "
            f"{control['diferencia_sum']!r}"
        )

        escribir(
            f"  Diferencia absoluta sum() = "
            f"{control['diferencia_abs_sum']!r}"
        )

        escribir(
            f"  Diferencia fsum() - SBS = "
            f"{control['diferencia_fsum']!r}"
        )

        escribir(
            f"  Diferencia absoluta fsum() = "
            f"{control['diferencia_abs_fsum']!r}"
        )

        escribir(
            f"  sum() - math.fsum() = "
            f"{control['diferencia_sum_vs_fsum']!r}"
        )

        escribir(
            f"  Diferencia absoluta "
            f"sum() vs fsum() = "
            f"{control['diferencia_abs_sum_vs_fsum']!r}"
        )

        escribir(
            f"  math.isclose(sum, SBS) "
            f"con tolerancia ACTUAL = "
            f"{control['isclose_sum_tolerancia_actual']}"
        )

        escribir(
            f"  math.isclose(fsum, SBS) "
            f"con tolerancia ACTUAL = "
            f"{control['isclose_fsum_tolerancia_actual']}"
        )


# ============================================================
# 25. RESUMEN DE FORMATOS
# ============================================================

def mostrar_resumen_formatos(
    metadatos
):

    subtitulo(
        "RESUMEN DE FORMATOS NUMÉRICOS OBSERVADOS"
    )

    for variable in (
        "MN",
        "ME",
        "TOTAL",
    ):

        filas = [
            item
            for item in metadatos
            if item[
                "variable"
            ]
            == variable
        ]

        formatos = sorted(
            {
                str(
                    item[
                        "format_str"
                    ]
                )
                for item in filas
            }
        )

        decimales = sorted(
            {
                str(
                    item[
                        "decimales_formato"
                    ]
                )
                for item in filas
            }
        )

        escribir()
        escribir(
            f"{variable}:"
        )

        escribir(
            f"  Formatos Excel distintos = "
            f"{formatos}"
        )

        escribir(
            f"  Decimales inferidos distintos = "
            f"{decimales}"
        )


# ============================================================
# 26. MAIN
# ============================================================

def main():

    try:

        titulo(
            "DIAGNÓSTICO DE PRECISIÓN / REDONDEO "
            "Y - B-2359 - 2015-11"
        )

        escribir(
            "Este programa NO valida Y."
        )

        escribir(
            "Este programa NO cambia tolerancias."
        )

        escribir(
            "Este programa NO ejecuta X1."
        )

        escribir(
            "Este programa NO cambia "
            "el parser ni DOLCRED."
        )

        escribir()

        escribir(
            f"Tolerancia actual preservada:"
        )

        escribir(
            f"  rel_tol = {REL_TOL}"
        )

        escribir(
            f"  abs_tol = {ABS_TOL}"
        )

        # ====================================================
        # A. LOCALIZAR Y DESCARGAR MISMO ARCHIVO
        # ====================================================

        titulo(
            "ARCHIVO OFICIAL B-2359"
        )

        enlace = localizar_archivo_oficial()

        info_descarga = descargar_archivo(
            enlace[
                "url"
            ],
            ARCHIVO_CREDITO
        )

        escribir(
            f"Archivo origen SBS: "
            f"{info_descarga['nombre_origen']}"
        )

        escribir(
            f"Archivo local: "
            f"{ARCHIVO_CREDITO}"
        )

        escribir(
            f"HTTP: "
            f"{info_descarga['http']}"
        )

        escribir(
            f"Bytes: "
            f"{info_descarga['bytes']}"
        )

        escribir(
            f"SHA-256: "
            f"{info_descarga['sha256']}"
        )

        formato = detectar_formato(
            ARCHIVO_CREDITO
        )

        escribir(
            f"Formato real: {formato}"
        )

        time.sleep(
            1.2
        )

        # ====================================================
        # B. MISMO PARSER Y
        # ====================================================

        titulo(
            "PARSER Y - SIN CAMBIOS"
        )

        resultado = extraer_y_creditos(
            ARCHIVO_CREDITO
        )

        escribir(
            f"Hoja: "
            f"{resultado['hoja']!r}"
        )

        escribir(
            f"Fila encabezados MN|ME|Total: "
            f"{resultado['fila_encabezados']}"
        )

        escribir(
            f"Fila bancos: "
            f"{resultado['fila_bancos']}"
        )

        escribir(
            f"Fila Total Créditos: "
            f"{resultado['fila_total_creditos']}"
        )

        escribir(
            f"Bloques detectados: "
            f"{resultado['cantidad_bloques']}"
        )

        # ====================================================
        # C. SUMAS
        # ====================================================

        bancos, agregado, controles = (
            diagnosticar_sumas(
                resultado[
                    "registros"
                ]
            )
        )

        escribir(
            f"Cantidad bancos individuales: "
            f"{len(bancos)}"
        )

        escribir(
            f"Agregado encontrado: "
            f"{agregado['banco_original']!r}"
        )

        mostrar_controles(
            controles
        )

        # ====================================================
        # D. FORMATOS EXCEL
        # ====================================================

        titulo(
            "METADATOS DE LAS CELDAS EXCEL"
        )

        metadatos = leer_metadatos_excel(
            ARCHIVO_CREDITO,
            resultado
        )

        mostrar_resumen_formatos(
            metadatos
        )

        mostrar_celdas(
            resultado[
                "registros"
            ],
            metadatos
        )

        # ====================================================
        # E. CSV AUXILIARES
        # ====================================================

        guardar_csv_celdas(
            metadatos
        )

        guardar_csv_controles(
            controles
        )

        # ====================================================
        # F. CLASIFICACIÓN
        # ====================================================

        titulo(
            "CLASIFICACIÓN DEL DIAGNÓSTICO"
        )

        escribir(
            "ESTADO = "
            "DIAGNOSTICO_PRECISION_Y_2015_11_COMPLETADO"
        )

        escribir()

        escribir(
            "IMPORTANTE:"
        )

        escribir(
            "Este estado NO significa que Y "
            "haya sido validada."
        )

        escribir(
            "No se modificó la tolerancia."
        )

        escribir(
            "No se tomó ninguna decisión "
            "sobre aceptar las diferencias."
        )

        escribir(
            "El siguiente paso será interpretar "
            "si las diferencias provienen de "
            "acumulación float, del formato/"
            "redondeo de las celdas SBS o de "
            "otra característica del archivo."
        )

    except Exception as error:

        titulo(
            "ERROR DEL DIAGNÓSTICO"
        )

        escribir(
            f"Tipo = "
            f"{type(error).__name__}"
        )

        escribir(
            f"Detalle = "
            f"{repr(error)}"
        )

        escribir(
            "ESTADO = "
            "DIAGNOSTICO_PRECISION_Y_2015_11_ERROR"
        )

        raise

    finally:

        TXT_RESULTADO.write_text(
            "\n".join(
                lineas_reporte
            ),
            encoding="utf-8"
        )

        print()
        print("=" * 118)
        print("DIAGNÓSTICO FINALIZADO")
        print("=" * 118)

        print()
        print(
            "TXT:"
        )

        print(
            TXT_RESULTADO
        )

        print()
        print(
            "CSV de controles:"
        )

        print(
            CSV_CONTROLES
        )

        print()
        print(
            "CSV de celdas/formato:"
        )

        print(
            CSV_CELDAS
        )


# ============================================================
# 27. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":

    main()