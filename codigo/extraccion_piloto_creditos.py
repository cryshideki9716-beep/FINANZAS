# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Extracción piloto de filas y valores - Créditos SBS

from pathlib import Path
import re

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. CONFIGURACIÓN
# ============================================================

RAIZ = Path(__file__).resolve().parent.parent

CARPETA_PILOTO = (
    RAIZ
    / "datos_crudos"
    / "sbs_piloto"
)

ARCHIVOS_CONTROL = [
    "creditos_2008_06.xls",
    "creditos_2010_12.xls",
    "creditos_2011_12.xls",
    "creditos_2015_06.xls",
    "creditos_2020_06.xls",
    "creditos_2025_12.xls",
]


# ============================================================
# 2. NORMALIZACIÓN DE TEXTO
# ============================================================

def normalizar_texto(valor):

    if valor is None:
        return ""

    texto = str(valor)

    texto = texto.replace("\n", " ")
    texto = texto.replace("\r", " ")

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


# ============================================================
# 3. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato(ruta):

    with open(ruta, "rb") as archivo:
        cabecera = archivo.read(8)

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:
        return "XLS_OLE"

    if cabecera.startswith(firmas_zip):
        return "XLSX_ZIP"

    return "DESCONOCIDO"


# ============================================================
# 4. LEER XLS/OLE
# ============================================================

def leer_xls(ruta):

    libro = xlrd.open_workbook(
        filename=str(ruta),
        on_demand=True
    )

    hojas = []

    for nombre in libro.sheet_names():

        hoja = libro.sheet_by_name(nombre)

        filas = []

        for i in range(hoja.nrows):

            fila = [
                hoja.cell_value(i, j)
                for j in range(hoja.ncols)
            ]

            filas.append(fila)

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.release_resources()

    return hojas


# ============================================================
# 5. LEER XLSX/ZIP
# ============================================================

def leer_xlsx(ruta):

    archivo_binario = open(
        ruta,
        "rb"
    )

    libro = load_workbook(
        archivo_binario,
        read_only=True,
        data_only=True
    )

    hojas = []

    for nombre in libro.sheetnames:

        hoja = libro[nombre]

        filas = []

        for fila in hoja.iter_rows(
            values_only=True
        ):

            filas.append(
                list(fila)
            )

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.close()
    archivo_binario.close()

    return hojas


# ============================================================
# 6. LEER ARCHIVO
# ============================================================

def leer_archivo(ruta):

    formato = detectar_formato(ruta)

    if formato == "XLS_OLE":

        hojas = leer_xls(ruta)

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(ruta)

    else:

        raise ValueError(
            f"Formato no reconocido: {ruta.name}"
        )

    return formato, hojas


# ============================================================
# 7. RECONOCER ENCABEZADOS
# ============================================================

def es_mn(valor):

    texto = normalizar_texto(valor)

    return (
        texto.startswith("mn")
        and "mil" in texto
    )


def es_me(valor):

    texto = normalizar_texto(valor)

    return (
        texto.startswith("me")
        and "mil" in texto
    )


def es_total(valor):

    texto = normalizar_texto(valor)

    return texto.startswith("total")


# ============================================================
# 8. ENCONTRAR FILA MN-ME-TOTAL
# ============================================================

def encontrar_fila_encabezados(filas):

    candidatos = []

    for numero_fila, fila in enumerate(filas):

        bloques = []

        for columna in range(
            len(fila) - 2
        ):

            if (
                es_mn(fila[columna])
                and es_me(fila[columna + 1])
                and es_total(fila[columna + 2])
            ):

                bloques.append(columna)

        if bloques:

            candidatos.append({
                "fila": numero_fila,
                "bloques": bloques,
            })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x: len(x["bloques"])
    )


# ============================================================
# 9. OBTENER NOMBRE DEL BANCO
# ============================================================

def obtener_nombre_banco(
    filas,
    fila_bancos,
    columna_inicio
):

    if fila_bancos < 0:
        return ""

    fila = filas[fila_bancos]

    for desplazamiento in (0, 1, 2):

        columna = (
            columna_inicio
            + desplazamiento
        )

        if columna >= len(fila):
            continue

        valor = fila[columna]

        if normalizar_texto(valor):

            return str(valor).strip()

    return ""


# ============================================================
# 10. DETECTAR ESTRUCTURA DE LA HOJA
# ============================================================

def detectar_estructura(filas):

    encabezado = encontrar_fila_encabezados(
        filas
    )

    if encabezado is None:
        return None

    fila_encabezados = encabezado["fila"]

    fila_bancos = (
        fila_encabezados - 1
    )

    bloques = []

    for columna_inicio in encabezado[
        "bloques"
    ]:

        banco = obtener_nombre_banco(
            filas,
            fila_bancos,
            columna_inicio
        )

        bloques.append({
            "banco_original": banco,
            "columna_mn": columna_inicio,
            "columna_me": columna_inicio + 1,
            "columna_total": columna_inicio + 2,
        })

    return {
        "fila_bancos": fila_bancos,
        "fila_encabezados": fila_encabezados,
        "bloques": bloques,
    }


# ============================================================
# 11. ELEGIR HOJA PRINCIPAL
# ============================================================

def elegir_hoja_principal(hojas):

    candidatos = []

    for hoja in hojas:

        estructura = detectar_estructura(
            hoja["filas"]
        )

        if estructura is not None:

            candidatos.append({
                "hoja": hoja,
                "estructura": estructura,
            })

    if not candidatos:

        return None

    return max(
        candidatos,
        key=lambda x: len(
            x["estructura"]["bloques"]
        )
    )


# ============================================================
# 12. IDENTIFICAR COLUMNA DE CONCEPTOS
# ============================================================

def detectar_columna_conceptos(
    filas,
    estructura
):
    """
    Las columnas anteriores al primer bloque bancario
    son candidatas a contener los conceptos/tipos de crédito.

    Elegimos la que tenga mayor cantidad de texto no vacío
    debajo de los encabezados.
    """

    primera_columna_banco = min(
        bloque["columna_mn"]
        for bloque in estructura["bloques"]
    )

    inicio_datos = (
        estructura["fila_encabezados"]
        + 1
    )

    candidatos = []

    for columna in range(
        primera_columna_banco
    ):

        cantidad_textos = 0

        for fila in filas[
            inicio_datos:
        ]:

            if columna >= len(fila):
                continue

            valor = fila[columna]

            texto = normalizar_texto(
                valor
            )

            if texto:
                cantidad_textos += 1

        candidatos.append(
            (
                columna,
                cantidad_textos
            )
        )

    if not candidatos:
        return None

    columna, cantidad = max(
        candidatos,
        key=lambda x: x[1]
    )

    if cantidad == 0:
        return None

    return columna


# ============================================================
# 13. CONVERTIR VALOR NUMÉRICO
# ============================================================

def convertir_numero(valor):
    """
    Conserva números como números.

    Celdas vacías o textos como '-' quedan como None.

    En esta etapa NO imputamos y NO convertimos
    valores faltantes en cero.
    """

    if valor is None:
        return None

    if isinstance(
        valor,
        (int, float)
    ):

        return float(valor)

    texto = str(valor).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "nd",
        "n/a",
    }:
        return None

    # Intento conservador para números escritos como texto.
    try:

        return float(
            texto.replace(",", "")
        )

    except ValueError:

        return None


# ============================================================
# 14. DETERMINAR SI UNA FILA TIENE DATOS
# ============================================================

def fila_tiene_datos_bancarios(
    fila,
    bloques
):

    for bloque in bloques:

        columnas = [
            bloque["columna_mn"],
            bloque["columna_me"],
            bloque["columna_total"],
        ]

        for columna in columnas:

            if columna >= len(fila):
                continue

            valor = convertir_numero(
                fila[columna]
            )

            if valor is not None:
                return True

    return False


# ============================================================
# 15. EXTRAER FILAS EN FORMATO LARGO
# ============================================================

def extraer_registros(
    filas,
    estructura,
    columna_conceptos,
    archivo
):

    registros = []

    inicio_datos = (
        estructura["fila_encabezados"]
        + 1
    )

    for numero_fila in range(
        inicio_datos,
        len(filas)
    ):

        fila = filas[numero_fila]

        if columna_conceptos >= len(fila):
            continue

        concepto_original = fila[
            columna_conceptos
        ]

        concepto = normalizar_texto(
            concepto_original
        )

        # No queremos filas completamente irrelevantes.
        if not concepto:
            continue

        if not fila_tiene_datos_bancarios(
            fila,
            estructura["bloques"]
        ):
            continue

        for bloque in estructura[
            "bloques"
        ]:

            banco = bloque[
                "banco_original"
            ]

            col_mn = bloque[
                "columna_mn"
            ]

            col_me = bloque[
                "columna_me"
            ]

            col_total = bloque[
                "columna_total"
            ]

            mn = (
                convertir_numero(
                    fila[col_mn]
                )
                if col_mn < len(fila)
                else None
            )

            me = (
                convertir_numero(
                    fila[col_me]
                )
                if col_me < len(fila)
                else None
            )

            total = (
                convertir_numero(
                    fila[col_total]
                )
                if col_total < len(fila)
                else None
            )

            # No generamos un registro si las tres
            # medidas están ausentes.
            if (
                mn is None
                and me is None
                and total is None
            ):
                continue

            registros.append({
                "archivo": archivo,
                "fila_origen": numero_fila,
                "concepto_original": str(
                    concepto_original
                ).strip(),
                "banco_original": banco,
                "mn_soles_miles": mn,
                "me_usd_miles": me,
                "total_soles_miles": total,
            })

    return registros


# ============================================================
# 16. MOSTRAR CONCEPTOS ENCONTRADOS
# ============================================================

def mostrar_conceptos(
    filas,
    estructura,
    columna_conceptos
):

    inicio_datos = (
        estructura["fila_encabezados"]
        + 1
    )

    print()
    print("CONCEPTOS/FILAS CON DATOS")
    print("-" * 90)

    cantidad = 0

    for numero_fila in range(
        inicio_datos,
        len(filas)
    ):

        fila = filas[
            numero_fila
        ]

        if columna_conceptos >= len(fila):
            continue

        concepto = fila[
            columna_conceptos
        ]

        if not normalizar_texto(
            concepto
        ):
            continue

        if not fila_tiene_datos_bancarios(
            fila,
            estructura["bloques"]
        ):
            continue

        cantidad += 1

        print(
            f"Fila {numero_fila}:",
            repr(
                str(concepto).strip()
            )
        )

    print()
    print(
        "Cantidad de filas/conceptos "
        "con datos:",
        cantidad
    )


# ============================================================
# 17. MOSTRAR MUESTRA DE REGISTROS
# ============================================================

def mostrar_muestra_registros(
    registros,
    limite=15
):

    print()
    print("MUESTRA DE VALORES EXTRAÍDOS")
    print("-" * 90)

    for registro in registros[
        :limite
    ]:

        print()
        print(
            "Fila origen:",
            registro["fila_origen"]
        )

        print(
            "Concepto:",
            repr(
                registro["concepto_original"]
            )
        )

        print(
            "Banco:",
            repr(
                registro["banco_original"]
            )
        )

        print(
            "MN (S/ miles):",
            registro["mn_soles_miles"]
        )

        print(
            "ME (US$ miles):",
            registro["me_usd_miles"]
        )

        print(
            "Total (S/ miles):",
            registro["total_soles_miles"]
        )


# ============================================================
# 18. ANALIZAR UN ARCHIVO
# ============================================================

def analizar_archivo(
    ruta
):

    print()
    print("=" * 90)

    print(
        "ARCHIVO:",
        ruta.name
    )

    print("=" * 90)

    formato, hojas = leer_archivo(
        ruta
    )

    seleccionado = elegir_hoja_principal(
        hojas
    )

    if seleccionado is None:

        print(
            "[ERROR] No se encontró "
            "la tabla principal."
        )

        return None

    hoja = seleccionado[
        "hoja"
    ]

    estructura = seleccionado[
        "estructura"
    ]

    filas = hoja[
        "filas"
    ]

    columna_conceptos = (
        detectar_columna_conceptos(
            filas,
            estructura
        )
    )

    print()
    print(
        "Formato:",
        formato
    )

    print(
        "Hoja seleccionada:",
        hoja["nombre"]
    )

    print(
        "Fila de bancos:",
        estructura["fila_bancos"]
    )

    print(
        "Fila MN/ME/Total:",
        estructura["fila_encabezados"]
    )

    print(
        "Bloques bancarios:",
        len(
            estructura["bloques"]
        )
    )

    print(
        "Columna de conceptos detectada:",
        columna_conceptos
    )

    if columna_conceptos is None:

        print()
        print(
            "[ERROR] No se pudo detectar "
            "la columna de conceptos."
        )

        return None

    mostrar_conceptos(
        filas,
        estructura,
        columna_conceptos
    )

    registros = extraer_registros(
        filas,
        estructura,
        columna_conceptos,
        ruta.name
    )

    print()
    print(
        "Registros banco-concepto extraídos:",
        len(registros)
    )

    mostrar_muestra_registros(
        registros,
        limite=15
    )

    return {
        "archivo": ruta.name,
        "formato": formato,
        "hoja": hoja["nombre"],
        "columna_conceptos": columna_conceptos,
        "bloques": len(
            estructura["bloques"]
        ),
        "registros": registros,
    }


# ============================================================
# 19. FUNCIÓN PRINCIPAL
# ============================================================

def main():

    print()
    print("=" * 90)

    print(
        "EXTRACCIÓN PILOTO DE CRÉDITOS SBS"
    )

    print(
        "SEIS PUNTOS DE CONTROL 2008-2025"
    )

    print("=" * 90)

    print()
    print(
        "Esta prueba trabaja exclusivamente "
        "con archivos locales."
    )

    print(
        "NO descarga información de SBS."
    )

    resultados = []

    for nombre in ARCHIVOS_CONTROL:

        ruta = (
            CARPETA_PILOTO
            / nombre
        )

        if not ruta.exists():

            print()
            print(
                "[ERROR] Archivo inexistente:",
                ruta
            )

            continue

        resultado = analizar_archivo(
            ruta
        )

        if resultado is not None:

            resultados.append(
                resultado
            )

    print()
    print("=" * 90)

    print(
        "RESUMEN GENERAL"
    )

    print("=" * 90)

    for resultado in resultados:

        print()

        print(
            resultado["archivo"]
        )

        print(
            "  Hoja:",
            resultado["hoja"]
        )

        print(
            "  Bloques bancarios:",
            resultado["bloques"]
        )

        print(
            "  Columna conceptos:",
            resultado["columna_conceptos"]
        )

        print(
            "  Registros extraídos:",
            len(
                resultado["registros"]
            )
        )

    print()
    print("=" * 90)

    print(
        "FIN DE LA PRUEBA"
    )

    print("=" * 90)

    print()
    print(
        "Los archivos crudos no fueron modificados."
    )

    print(
        "Todavía NO se calculó la dolarización."
    )

    print(
        "Todavía NO se descargaron los 211 meses."
    )


# ============================================================
# 20. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":
    main()