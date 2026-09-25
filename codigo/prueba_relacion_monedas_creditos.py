# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Prueba de relación MN - ME - Total - Tipo de Cambio en créditos SBS

from pathlib import Path
import re
import statistics

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. CONFIGURACIÓN
# ============================================================

RAIZ = Path(__file__).resolve().parent.parent

CARPETA_PILOTO = (
    RAIZ
    / "datos_crudos"
    / "sbs_piloto"
)

ARCHIVOS_CONTROL = [
    "creditos_2008_06.xls",
    "creditos_2010_12.xls",
    "creditos_2011_12.xls",
    "creditos_2015_06.xls",
    "creditos_2020_06.xls",
    "creditos_2025_12.xls",
]


# ============================================================
# 2. NORMALIZACIÓN
# ============================================================

def normalizar_texto(valor):

    if valor is None:
        return ""

    texto = str(valor)

    texto = texto.replace("\n", " ")
    texto = texto.replace("\r", " ")

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


# ============================================================
# 3. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato(ruta):

    with open(ruta, "rb") as archivo:
        cabecera = archivo.read(8)

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:
        return "XLS_OLE"

    if cabecera.startswith(firmas_zip):
        return "XLSX_ZIP"

    return "DESCONOCIDO"


# ============================================================
# 4. LEER XLS
# ============================================================

def leer_xls(ruta):

    libro = xlrd.open_workbook(
        filename=str(ruta),
        on_demand=True
    )

    hojas = []

    for nombre in libro.sheet_names():

        hoja = libro.sheet_by_name(nombre)

        filas = []

        for i in range(hoja.nrows):

            fila = [
                hoja.cell_value(i, j)
                for j in range(hoja.ncols)
            ]

            filas.append(fila)

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.release_resources()

    return hojas


# ============================================================
# 5. LEER XLSX
# ============================================================

def leer_xlsx(ruta):

    archivo_binario = open(
        ruta,
        "rb"
    )

    libro = load_workbook(
        archivo_binario,
        read_only=True,
        data_only=True
    )

    hojas = []

    for nombre in libro.sheetnames:

        hoja = libro[nombre]

        filas = []

        for fila in hoja.iter_rows(
            values_only=True
        ):

            filas.append(
                list(fila)
            )

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.close()
    archivo_binario.close()

    return hojas


# ============================================================
# 6. LEER ARCHIVO
# ============================================================

def leer_archivo(ruta):

    formato = detectar_formato(ruta)

    if formato == "XLS_OLE":

        hojas = leer_xls(ruta)

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(ruta)

    else:

        raise ValueError(
            f"Formato no reconocido: {ruta.name}"
        )

    return formato, hojas


# ============================================================
# 7. RECONOCER MN / ME / TOTAL
# ============================================================

def es_mn(valor):

    texto = normalizar_texto(valor)

    return (
        texto.startswith("mn")
        and "mil" in texto
    )


def es_me(valor):

    texto = normalizar_texto(valor)

    return (
        texto.startswith("me")
        and "mil" in texto
    )


def es_total(valor):

    texto = normalizar_texto(valor)

    return texto.startswith("total")


# ============================================================
# 8. ENCONTRAR FILA DE ENCABEZADOS
# ============================================================

def encontrar_fila_encabezados(filas):

    candidatos = []

    for numero_fila, fila in enumerate(filas):

        bloques = []

        for columna in range(
            len(fila) - 2
        ):

            if (
                es_mn(fila[columna])
                and es_me(fila[columna + 1])
                and es_total(fila[columna + 2])
            ):

                bloques.append(columna)

        if bloques:

            candidatos.append({
                "fila": numero_fila,
                "bloques": bloques,
            })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x: len(x["bloques"])
    )


# ============================================================
# 9. NOMBRE DEL BANCO
# ============================================================

def obtener_nombre_banco(
    filas,
    fila_bancos,
    columna_inicio
):

    fila = filas[fila_bancos]

    for desplazamiento in (0, 1, 2):

        columna = (
            columna_inicio
            + desplazamiento
        )

        if columna >= len(fila):
            continue

        valor = fila[columna]

        if normalizar_texto(valor):

            return str(valor).strip()

    return ""


# ============================================================
# 10. DETECTAR ESTRUCTURA
# ============================================================

def detectar_estructura(filas):

    encabezado = encontrar_fila_encabezados(
        filas
    )

    if encabezado is None:
        return None

    fila_encabezados = encabezado["fila"]
    fila_bancos = fila_encabezados - 1

    bloques = []

    for columna_inicio in encabezado["bloques"]:

        banco = obtener_nombre_banco(
            filas,
            fila_bancos,
            columna_inicio
        )

        bloques.append({
            "banco_original": banco,
            "columna_mn": columna_inicio,
            "columna_me": columna_inicio + 1,
            "columna_total": columna_inicio + 2,
        })

    return {
        "fila_bancos": fila_bancos,
        "fila_encabezados": fila_encabezados,
        "bloques": bloques,
    }


# ============================================================
# 11. ELEGIR HOJA PRINCIPAL
# ============================================================

def elegir_hoja_principal(hojas):

    candidatos = []

    for hoja in hojas:

        estructura = detectar_estructura(
            hoja["filas"]
        )

        if estructura is not None:

            candidatos.append({
                "hoja": hoja,
                "estructura": estructura,
            })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x: len(
            x["estructura"]["bloques"]
        )
    )


# ============================================================
# 12. CONVERTIR NÚMEROS
# ============================================================

def convertir_numero(valor):

    if valor is None:
        return None

    if isinstance(
        valor,
        (int, float)
    ):

        return float(valor)

    texto = str(valor).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "nd",
        "n/a",
    }:
        return None

    try:

        return float(
            texto.replace(",", "")
        )

    except ValueError:

        return None


# ============================================================
# 13. BUSCAR UNA FILA POR CONCEPTO
# ============================================================

def buscar_fila_concepto(
    filas,
    texto_buscado
):

    buscado = normalizar_texto(
        texto_buscado
    )

    for numero_fila, fila in enumerate(
        filas
    ):

        for valor in fila:

            texto = normalizar_texto(
                valor
            )

            if buscado in texto:

                return numero_fila

    return None


# ============================================================
# 14. BUSCAR TIPO DE CAMBIO PUBLICADO
# ============================================================

def buscar_tipo_cambio_publicado(
    filas
):
    """
    Busca la fila 'Tipo de Cambio:'.

    Como todavía no sabemos exactamente en qué columna
    queda el valor en todas las versiones del reporte,
    inspeccionamos las celdas de esa misma fila y
    recuperamos los números encontrados.

    No asumimos todavía cuál de ellos es el TC correcto.
    """

    fila_tc = buscar_fila_concepto(
        filas,
        "tipo de cambio"
    )

    if fila_tc is None:

        return {
            "fila": None,
            "numeros": [],
            "contenido": None,
        }

    fila = filas[
        fila_tc
    ]

    numeros = []

    for columna, valor in enumerate(
        fila
    ):

        numero = convertir_numero(
            valor
        )

        if numero is not None:

            numeros.append({
                "columna": columna,
                "valor": numero,
            })

    return {
        "fila": fila_tc,
        "numeros": numeros,
        "contenido": fila,
    }


# ============================================================
# 15. ANALIZAR TOTAL CRÉDITOS
# ============================================================

def analizar_total_creditos(
    filas,
    estructura
):

    fila_total = buscar_fila_concepto(
        filas,
        "total créditos"
    )

    if fila_total is None:

        # Variante sin tilde, por seguridad
        fila_total = buscar_fila_concepto(
            filas,
            "total creditos"
        )

    if fila_total is None:

        return None

    fila = filas[
        fila_total
    ]

    resultados = []

    for bloque in estructura[
        "bloques"
    ]:

        banco = bloque[
            "banco_original"
        ]

        col_mn = bloque[
            "columna_mn"
        ]

        col_me = bloque[
            "columna_me"
        ]

        col_total = bloque[
            "columna_total"
        ]

        mn = (
            convertir_numero(
                fila[col_mn]
            )
            if col_mn < len(fila)
            else None
        )

        me = (
            convertir_numero(
                fila[col_me]
            )
            if col_me < len(fila)
            else None
        )

        total = (
            convertir_numero(
                fila[col_total]
            )
            if col_total < len(fila)
            else None
        )

        tc_implicito = None
        total_reconstruido = None
        diferencia = None

        if (
            mn is not None
            and me is not None
            and total is not None
            and me != 0
        ):

            tc_implicito = (
                total - mn
            ) / me

            total_reconstruido = (
                mn
                + me * tc_implicito
            )

            diferencia = (
                total
                - total_reconstruido
            )

        resultados.append({
            "banco": banco,
            "mn": mn,
            "me": me,
            "total": total,
            "tc_implicito": tc_implicito,
            "total_reconstruido": total_reconstruido,
            "diferencia": diferencia,
        })

    return {
        "fila_total": fila_total,
        "resultados": resultados,
    }


# ============================================================
# 16. ANALIZAR ARCHIVO
# ============================================================

def analizar_archivo(ruta):

    print()
    print("=" * 100)

    print(
        "ARCHIVO:",
        ruta.name
    )

    print("=" * 100)

    formato, hojas = leer_archivo(
        ruta
    )

    seleccionado = elegir_hoja_principal(
        hojas
    )

    if seleccionado is None:

        print(
            "[ERROR] No se encontró "
            "la tabla principal."
        )

        return None

    hoja = seleccionado[
        "hoja"
    ]

    estructura = seleccionado[
        "estructura"
    ]

    filas = hoja[
        "filas"
    ]

    print()
    print(
        "Formato:",
        formato
    )

    print(
        "Hoja:",
        hoja["nombre"]
    )

    # --------------------------------------------------------
    # Tipo de cambio publicado
    # --------------------------------------------------------

    tc_publicado = (
        buscar_tipo_cambio_publicado(
            filas
        )
    )

    print()
    print(
        "TIPO DE CAMBIO PUBLICADO EN EL ARCHIVO"
    )

    print("-" * 100)

    if tc_publicado["fila"] is None:

        print(
            "No se encontró una fila "
            "'Tipo de Cambio:'."
        )

    else:

        print(
            "Fila:",
            tc_publicado["fila"]
        )

        print(
            "Valores numéricos encontrados "
            "en esa fila:"
        )

        for item in tc_publicado[
            "numeros"
        ]:

            print(
                "  Columna",
                item["columna"],
                "->",
                item["valor"]
            )

    # --------------------------------------------------------
    # Total créditos
    # --------------------------------------------------------

    analisis = analizar_total_creditos(
        filas,
        estructura
    )

    if analisis is None:

        print()
        print(
            "[ERROR] No se encontró "
            "'Total Créditos:'."
        )

        return None

    print()
    print(
        "FILA TOTAL CRÉDITOS:",
        analisis["fila_total"]
    )

    print()
    print(
        "RELACIÓN MN - ME - TOTAL"
    )

    print("-" * 100)

    tc_validos = []

    for item in analisis[
        "resultados"
    ]:

        print()
        print(
            "Banco:",
            repr(item["banco"])
        )

        print(
            "  MN (S/ miles):",
            item["mn"]
        )

        print(
            "  ME (US$ miles):",
            item["me"]
        )

        print(
            "  Total (S/ miles):",
            item["total"]
        )

        print(
            "  TC implícito:",
            item["tc_implicito"]
        )

        if item[
            "tc_implicito"
        ] is not None:

            tc_validos.append(
                item["tc_implicito"]
            )

    # --------------------------------------------------------
    # Resumen del TC implícito
    # --------------------------------------------------------

    print()
    print(
        "RESUMEN DEL TC IMPLÍCITO"
    )

    print("-" * 100)

    if tc_validos:

        promedio = statistics.mean(
            tc_validos
        )

        mediana = statistics.median(
            tc_validos
        )

        minimo = min(
            tc_validos
        )

        maximo = max(
            tc_validos
        )

        rango = (
            maximo - minimo
        )

        print(
            "Bancos utilizables:",
            len(tc_validos)
        )

        print(
            "TC implícito promedio:",
            promedio
        )

        print(
            "TC implícito mediano:",
            mediana
        )

        print(
            "TC implícito mínimo:",
            minimo
        )

        print(
            "TC implícito máximo:",
            maximo
        )

        print(
            "Rango máximo-mínimo:",
            rango
        )

    else:

        promedio = None

        print(
            "No fue posible calcular "
            "TC implícito."
        )

    return {
        "archivo": ruta.name,
        "hoja": hoja["nombre"],
        "fila_total": analisis[
            "fila_total"
        ],
        "tc_publicado": tc_publicado,
        "tc_implicitos": tc_validos,
        "tc_promedio": promedio,
    }


# ============================================================
# 17. FUNCIÓN PRINCIPAL
# ============================================================

def main():

    print()
    print("=" * 100)

    print(
        "PRUEBA MN - ME - TOTAL - TIPO DE CAMBIO"
    )

    print(
        "CRÉDITOS SBS | PILOTO 2008-2025"
    )

    print("=" * 100)

    print()
    print(
        "NO se descargarán archivos."
    )

    print(
        "NO se modificarán los archivos crudos."
    )

    print(
        "NO se calculará todavía DOLCRED."
    )

    resultados = []

    for nombre in ARCHIVOS_CONTROL:

        ruta = (
            CARPETA_PILOTO
            / nombre
        )

        if not ruta.exists():

            print()
            print(
                "[ERROR] No existe:",
                ruta
            )

            continue

        resultado = analizar_archivo(
            ruta
        )

        if resultado is not None:

            resultados.append(
                resultado
            )

    # ========================================================
    # RESUMEN FINAL
    # ========================================================

    print()
    print("=" * 100)

    print(
        "RESUMEN COMPARATIVO DE LOS SEIS PILOTOS"
    )

    print("=" * 100)

    for resultado in resultados:

        print()
        print(
            resultado["archivo"]
        )

        print(
            "  Fila Total Créditos:",
            resultado["fila_total"]
        )

        print(
            "  TC implícito promedio:",
            resultado["tc_promedio"]
        )

        tc_archivo = resultado[
            "tc_publicado"
        ]

        if tc_archivo["fila"] is None:

            print(
                "  Fila Tipo de Cambio: "
                "NO PRESENTE"
            )

        else:

            print(
                "  Fila Tipo de Cambio:",
                tc_archivo["fila"]
            )

            print(
                "  Valores encontrados:",
                [
                    x["valor"]
                    for x in tc_archivo["numeros"]
                ]
            )

    print()
    print("=" * 100)

    print(
        "FIN DE LA PRUEBA"
    )

    print("=" * 100)

    print()
    print(
        "Los seis archivos originales "
        "permanecen sin modificar."
    )


# ============================================================
# 18. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":
    main()