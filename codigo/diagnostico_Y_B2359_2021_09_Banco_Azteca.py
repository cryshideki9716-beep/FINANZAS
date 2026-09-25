# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Fecha de diagnóstico: 2026-09-24

from pathlib import Path
import csv
import hashlib
import math
import re
import unicodedata

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. OBJETIVO
# ============================================================

"""
DIAGNÓSTICO SEPARADO

Y = Dolarización del crédito (%)
SBS B-2359
Mes = 2021-09
Banco objetivo = Banco Azteca

Archivo local:

datos_crudos/Y_B2359_2024200485D/
Y_B-2359_2024200485D_2021-09.xls

ESTE SCRIPT:

- NO descarga nada.
- NO modifica el script de producción.
- NO modifica el parser.
- NO modifica DOLCRED.
- NO modifica tolerancias.
- NO ejecuta X1.
- NO homologa nombres.
- NO interpola.
- NO selecciona bancos.
- NO corrige valores SBS.

OBJETIVOS:

1. Reproducir exactamente el parser B-2359 utilizado
   en producción.

2. Localizar exactamente el bloque de Banco Azteca.

3. Mostrar:
   - nombre exacto SBS;
   - columnas MN / ME / Total;
   - valores raw;
   - valores convertidos;
   - DOLCRED;
   - TC implícito;
   - formatos Excel;
   - decimales visibles;
   - encabezados raw;
   - fila Total Créditos.

4. Verificar:
   - existencia de celdas combinadas;
   - posible desplazamiento de columnas;
   - continuidad MN | ME | Total;
   - bloque anterior;
   - bloque siguiente;
   - comparación con Total Banca Múltiple;
   - correspondencia parser/raw.

5. Mostrar las filas cercanas a "Total Créditos"
   con las tres celdas de Banco Azteca para revisar
   si el valor Total observado pertenece realmente
   a esa fila/concepto.

Este diagnóstico NO declara Y validada.
"""


# ============================================================
# 2. CONFIGURACIÓN
# ============================================================

CODIGO_ESTUDIANTE = "2024200485D"

MES = "2021-09"

BANCO_BUSCADO = "Banco Azteca"

# Cantidad de filas anteriores y posteriores a Total Créditos
# que se mostrarán para Banco Azteca.
RADIO_FILAS_CERCANAS = 10


# ============================================================
# 3. RUTAS
# ============================================================

RAIZ = Path(
    __file__
).resolve().parent.parent


ARCHIVO_LOCAL = (
    RAIZ
    / "datos_crudos"
    / f"Y_B2359_{CODIGO_ESTUDIANTE}"
    / f"Y_B-2359_{CODIGO_ESTUDIANTE}_{MES}.xls"
)


CARPETA_DIAGNOSTICOS = (
    RAIZ
    / "salidas"
    / "diagnosticos"
)

CARPETA_DIAGNOSTICOS.mkdir(
    parents=True,
    exist_ok=True
)


BASE_SALIDA = (
    f"diagnostico_Y_B2359_"
    f"{MES.replace('-', '_')}_Banco_Azteca"
)


TXT_RESULTADO = (
    CARPETA_DIAGNOSTICOS
    / f"{BASE_SALIDA}.txt"
)


CSV_FILAS_CERCANAS = (
    CARPETA_DIAGNOSTICOS
    / f"{BASE_SALIDA}_filas_cercanas.csv"
)


CSV_BLOQUES = (
    CARPETA_DIAGNOSTICOS
    / f"{BASE_SALIDA}_bloques.csv"
)


# ============================================================
# 4. EXCEPCIÓN DEL DIAGNÓSTICO
# ============================================================

class ErrorDiagnostico(Exception):
    pass


# ============================================================
# 5. SALIDA TXT
# ============================================================

lineas_reporte = []


def escribir(texto=""):

    texto = str(
        texto
    )

    print(
        texto
    )

    lineas_reporte.append(
        texto
    )


def titulo(texto):

    escribir()
    escribir("=" * 124)
    escribir(texto)
    escribir("=" * 124)


def subtitulo(texto):

    escribir()
    escribir("-" * 124)
    escribir(texto)
    escribir("-" * 124)


# ============================================================
# 6. UTILIDADES DE TEXTO
# MISMAS BASES DEL PARSER DE PRODUCCIÓN
# ============================================================

def normalizar_texto(valor):

    if valor is None:
        return ""

    texto = str(
        valor
    )

    texto = texto.replace(
        "\n",
        " "
    )

    texto = texto.replace(
        "\r",
        " "
    )

    texto = texto.replace(
        "\xa0",
        " "
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


def normalizar_comparacion(valor):

    texto = normalizar_texto(
        valor
    )

    texto = unicodedata.normalize(
        "NFKD",
        texto
    )

    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(
            caracter
        )
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip()


def sin_tildes(valor):

    return normalizar_comparacion(
        valor
    )


# ============================================================
# 7. CONVERSIÓN NUMÉRICA
# MISMO CRITERIO DEL PARSER DE PRODUCCIÓN
# ============================================================

def convertir_numero(valor):

    if valor is None:
        return None

    if isinstance(
        valor,
        (int, float)
    ):

        return float(
            valor
        )

    texto = str(
        valor
    ).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "n.d",
        "nd",
        "n/a",
        "na",
    }:

        return None

    try:

        return float(
            texto.replace(
                ",",
                ""
            )
        )

    except ValueError:

        return None


# ============================================================
# 8. HASH SHA-256
# ============================================================

def sha256_archivo(ruta):

    objeto = hashlib.sha256()

    with Path(
        ruta
    ).open(
        "rb"
    ) as archivo:

        while True:

            bloque = archivo.read(
                1024 * 1024
            )

            if not bloque:
                break

            objeto.update(
                bloque
            )

    return objeto.hexdigest()


# ============================================================
# 9. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato(ruta):

    with Path(
        ruta
    ).open(
        "rb"
    ) as archivo:

        cabecera = archivo.read(
            8
        )

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:

        return "XLS_OLE"

    if cabecera.startswith(
        firmas_zip
    ):

        return "XLSX_ZIP"

    raise ErrorDiagnostico(
        f"Formato binario no reconocido: {ruta}"
    )


# ============================================================
# 10. LECTURA NORMAL DEL PARSER - XLS
# ============================================================

def leer_xls(ruta):

    try:

        libro = xlrd.open_workbook(
            filename=str(
                ruta
            ),
            on_demand=True
        )

    except Exception as error:

        raise ErrorDiagnostico(
            f"No pudo abrirse como XLS/OLE: {ruta}"
        ) from error

    hojas = []

    try:

        for nombre in libro.sheet_names():

            hoja = libro.sheet_by_name(
                nombre
            )

            filas = []

            for i in range(
                hoja.nrows
            ):

                fila = [
                    hoja.cell_value(
                        i,
                        j
                    )
                    for j in range(
                        hoja.ncols
                    )
                ]

                filas.append(
                    fila
                )

            hojas.append({
                "nombre":
                    nombre,

                "filas":
                    filas,
            })

    finally:

        libro.release_resources()

    return hojas


# ============================================================
# 11. LECTURA NORMAL DEL PARSER - XLSX
# ============================================================

def leer_xlsx(ruta):

    try:

        with Path(
            ruta
        ).open(
            "rb"
        ) as archivo_binario:

            libro = load_workbook(
                archivo_binario,
                read_only=True,
                data_only=True
            )

            hojas = []

            try:

                for nombre in libro.sheetnames:

                    hoja = libro[
                        nombre
                    ]

                    filas = [
                        list(
                            fila
                        )
                        for fila
                        in hoja.iter_rows(
                            values_only=True
                        )
                    ]

                    hojas.append({
                        "nombre":
                            nombre,

                        "filas":
                            filas,
                    })

            finally:

                libro.close()

            return hojas

    except Exception as error:

        raise ErrorDiagnostico(
            f"No pudo abrirse como XLSX/ZIP: {ruta}"
        ) from error


# ============================================================
# 12. LEER ARCHIVO SEGÚN FIRMA
# ============================================================

def leer_archivo(ruta):

    formato = detectar_formato(
        ruta
    )

    if formato == "XLS_OLE":

        hojas = leer_xls(
            ruta
        )

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(
            ruta
        )

    else:

        raise ErrorDiagnostico(
            f"Formato no soportado: {formato}"
        )

    return (
        formato,
        hojas,
    )


# ============================================================
# 13. RECONOCER MN / ME / TOTAL
# MISMA LÓGICA DE PRODUCCIÓN
# ============================================================

def es_mn(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith(
            "mn"
        )
        and
        "mil" in texto
    )


def es_me(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith(
            "me"
        )
        and
        "mil" in texto
    )


def es_total(valor):

    return normalizar_texto(
        valor
    ).startswith(
        "total"
    )


# ============================================================
# 14. FILA DE ENCABEZADOS B-2359
# ============================================================

def encontrar_fila_encabezados_creditos(
    filas
):

    candidatos = []

    for numero_fila, fila in enumerate(
        filas
    ):

        bloques = []

        for columna in range(
            max(
                0,
                len(
                    fila
                ) - 2
            )
        ):

            if (
                es_mn(
                    fila[
                        columna
                    ]
                )
                and
                es_me(
                    fila[
                        columna + 1
                    ]
                )
                and
                es_total(
                    fila[
                        columna + 2
                    ]
                )
            ):

                bloques.append(
                    columna
                )

        if bloques:

            candidatos.append({
                "fila":
                    numero_fila,

                "bloques":
                    bloques,
            })

    if not candidatos:

        return None

    return max(
        candidatos,
        key=lambda x:
        len(
            x[
                "bloques"
            ]
        )
    )


# ============================================================
# 15. NOMBRE DE BANCO
# MISMA LÓGICA DE PRODUCCIÓN
# ============================================================

def obtener_nombre_banco_creditos(
    filas,
    fila_bancos,
    columna_inicio
):

    if (
        fila_bancos < 0
        or
        fila_bancos >= len(
            filas
        )
    ):

        return ""

    fila = filas[
        fila_bancos
    ]

    for desplazamiento in (
        0,
        1,
        2,
    ):

        columna = (
            columna_inicio
            +
            desplazamiento
        )

        if columna >= len(
            fila
        ):

            continue

        valor = fila[
            columna
        ]

        if normalizar_texto(
            valor
        ):

            return str(
                valor
            ).strip()

    return ""


# ============================================================
# 16. DETECTAR ESTRUCTURA
# MISMO PARSER
# ============================================================

def detectar_estructura_creditos(
    filas
):

    encabezado = encontrar_fila_encabezados_creditos(
        filas
    )

    if encabezado is None:

        return None

    fila_encabezados = encabezado[
        "fila"
    ]

    fila_bancos = (
        fila_encabezados
        -
        1
    )

    bloques = []

    for columna_inicio in encabezado[
        "bloques"
    ]:

        bloques.append({

            "banco_original":
                obtener_nombre_banco_creditos(
                    filas,
                    fila_bancos,
                    columna_inicio
                ),

            "columna_mn":
                columna_inicio,

            "columna_me":
                columna_inicio + 1,

            "columna_total":
                columna_inicio + 2,
        })

    return {

        "fila_bancos":
            fila_bancos,

        "fila_encabezados":
            fila_encabezados,

        "bloques":
            bloques,
    }


# ============================================================
# 17. ELEGIR HOJA PRINCIPAL
# MISMA LÓGICA DE PRODUCCIÓN
# ============================================================

def elegir_hoja_principal_creditos(
    hojas
):

    candidatos = []

    for hoja in hojas:

        estructura = detectar_estructura_creditos(
            hoja[
                "filas"
            ]
        )

        if estructura is None:

            continue

        candidatos.append({

            "hoja":
                hoja,

            "estructura":
                estructura,
        })

    if not candidatos:

        return None

    return max(
        candidatos,
        key=lambda x:
        len(
            x[
                "estructura"
            ][
                "bloques"
            ]
        )
    )


# ============================================================
# 18. BUSCAR FILA TOTAL CRÉDITOS
# ============================================================

def buscar_fila_total_creditos(
    filas
):

    for numero_fila, fila in enumerate(
        filas
    ):

        for valor in fila:

            if (
                "total creditos"
                in
                sin_tildes(
                    valor
                )
            ):

                return numero_fila

    return None


# ============================================================
# 19. AGREGADO
# ============================================================

def es_agregado_creditos(
    nombre
):

    return (
        "total banca multiple"
        in
        sin_tildes(
            nombre
        )
    )


# ============================================================
# 20. PARSER Y
# MISMA FÓRMULA DOLCRED
# ============================================================

def extraer_y_creditos(
    ruta
):

    formato, hojas = leer_archivo(
        ruta
    )

    seleccionado = elegir_hoja_principal_creditos(
        hojas
    )

    if seleccionado is None:

        raise ErrorDiagnostico(
            "B-2359: no se encontró "
            "estructura MN|ME|Total."
        )

    hoja = seleccionado[
        "hoja"
    ]

    estructura = seleccionado[
        "estructura"
    ]

    filas = hoja[
        "filas"
    ]

    fila_total = buscar_fila_total_creditos(
        filas
    )

    if fila_total is None:

        raise ErrorDiagnostico(
            "B-2359: no se encontró "
            "la fila Total Créditos."
        )

    fila = filas[
        fila_total
    ]

    registros = []

    for bloque in estructura[
        "bloques"
    ]:

        banco = bloque[
            "banco_original"
        ]

        col_mn = bloque[
            "columna_mn"
        ]

        col_me = bloque[
            "columna_me"
        ]

        col_total = bloque[
            "columna_total"
        ]

        mn = (
            convertir_numero(
                fila[
                    col_mn
                ]
            )
            if col_mn < len(
                fila
            )
            else None
        )

        me = (
            convertir_numero(
                fila[
                    col_me
                ]
            )
            if col_me < len(
                fila
            )
            else None
        )

        total = (
            convertir_numero(
                fila[
                    col_total
                ]
            )
            if col_total < len(
                fila
            )
            else None
        )

        problemas = []

        if not banco:

            problemas.append(
                "BANCO_SIN_NOMBRE"
            )

        if mn is None:

            problemas.append(
                "MN_FALTANTE"
            )

        if me is None:

            problemas.append(
                "ME_FALTANTE"
            )

        if total is None:

            problemas.append(
                "TOTAL_FALTANTE"
            )

        if (
            mn is not None
            and
            mn < 0
        ):

            problemas.append(
                "MN_NEGATIVO"
            )

        if (
            me is not None
            and
            me < 0
        ):

            problemas.append(
                "ME_NEGATIVO"
            )

        if (
            total is not None
            and
            total <= 0
        ):

            problemas.append(
                "TOTAL_NO_POSITIVO"
            )

        if (
            mn is not None
            and
            total is not None
            and
            total < mn
        ):

            problemas.append(
                "TOTAL_MENOR_QUE_MN"
            )

        dolcred = None
        tc_implicito = None

        # ====================================================
        # FÓRMULA DE PRODUCCIÓN - SIN CAMBIOS
        # ====================================================

        if (
            total is not None
            and
            total > 0
            and
            mn is not None
        ):

            dolcred = (
                (
                    total
                    -
                    mn
                )
                /
                total
                *
                100
            )

            if not (
                0
                <=
                dolcred
                <=
                100
            ):

                problemas.append(
                    "DOLCRED_FUERA_RANGO"
                )

        # ====================================================
        # TC IMPLÍCITO - SIN CAMBIOS
        # ====================================================

        if (
            mn is not None
            and
            me is not None
            and
            total is not None
            and
            me > 0
        ):

            tc_implicito = (
                (
                    total
                    -
                    mn
                )
                /
                me
            )

            if tc_implicito <= 0:

                problemas.append(
                    "TC_NO_POSITIVO"
                )

        registros.append({

            "banco_original":
                banco,

            "mn_soles_miles":
                mn,

            "me_usd_miles":
                me,

            "total_soles_miles":
                total,

            "dolcred_pct":
                dolcred,

            "tc_implicito":
                tc_implicito,

            "es_agregado":
                es_agregado_creditos(
                    banco
                ),

            "problemas":
                problemas,

            # Solo se conservan para diagnóstico.
            "columna_mn":
                col_mn,

            "columna_me":
                col_me,

            "columna_total":
                col_total,
        })

    return {

        "formato":
            formato,

        "hoja":
            hoja[
                "nombre"
            ],

        "fila_bancos":
            estructura[
                "fila_bancos"
            ],

        "fila_encabezados":
            estructura[
                "fila_encabezados"
            ],

        "fila_total_creditos":
            fila_total,

        "cantidad_bloques":
            len(
                estructura[
                    "bloques"
                ]
            ),

        "estructura":
            estructura,

        "filas":
            filas,

        "registros":
            registros,
    }


# ============================================================
# 21. INFERIR DECIMALES DEL FORMATO
# ============================================================

def inferir_decimales_formato_excel(
    formato
):

    if formato is None:

        return None

    texto = str(
        formato
    ).strip()

    if not texto:

        return None

    if texto.lower() == "general":

        return None

    seccion = texto.split(
        ";"
    )[0]

    seccion = re.sub(
        r'"[^"]*"',
        "",
        seccion
    )

    seccion = re.sub(
        r"\[[^\]]*\]",
        "",
        seccion
    )

    coincidencia = re.search(
        r"\.([0#?]+)",
        seccion
    )

    if coincidencia:

        return len(
            coincidencia.group(1)
        )

    if re.search(
        r"[0#?]",
        seccion
    ):

        return 0

    return None


# ============================================================
# 22. LOCALIZAR BANCO AZTECA
# ============================================================

def localizar_banco_objetivo(
    resultado
):

    objetivo = normalizar_comparacion(
        BANCO_BUSCADO
    )

    coincidencias = [
        registro
        for registro in resultado[
            "registros"
        ]
        if normalizar_comparacion(
            registro[
                "banco_original"
            ]
        )
        ==
        objetivo
    ]

    # Si no existe igualdad exacta normalizada,
    # mostrar coincidencias parciales para diagnóstico.
    if not coincidencias:

        parciales = [
            registro
            for registro in resultado[
                "registros"
            ]
            if (
                "azteca"
                in
                normalizar_comparacion(
                    registro[
                        "banco_original"
                    ]
                )
            )
        ]

        raise ErrorDiagnostico(
            "No se encontró exactamente Banco Azteca. "
            f"Coincidencias parciales: "
            f"{[r['banco_original'] for r in parciales]}"
        )

    if len(
        coincidencias
    ) != 1:

        raise ErrorDiagnostico(
            "Se esperaba exactamente un bloque "
            "de Banco Azteca y se encontraron "
            f"{len(coincidencias)}."
        )

    return coincidencias[
        0
    ]


# ============================================================
# 23. LOCALIZAR ÍNDICE DEL BLOQUE
# ============================================================

def localizar_indice_bloque(
    resultado,
    registro_objetivo
):

    bloques = resultado[
        "estructura"
    ][
        "bloques"
    ]

    indices = [
        indice
        for indice, bloque in enumerate(
            bloques
        )
        if (
            bloque[
                "columna_mn"
            ]
            ==
            registro_objetivo[
                "columna_mn"
            ]
            and
            bloque[
                "columna_me"
            ]
            ==
            registro_objetivo[
                "columna_me"
            ]
            and
            bloque[
                "columna_total"
            ]
            ==
            registro_objetivo[
                "columna_total"
            ]
        )
    ]

    if len(
        indices
    ) != 1:

        raise ErrorDiagnostico(
            "No pudo determinarse de forma unívoca "
            "el índice del bloque Banco Azteca."
        )

    return indices[
        0
    ]


# ============================================================
# 24. CONVERTIR NÚMERO DE COLUMNA A LETRA EXCEL
# ============================================================

def columna_excel(
    indice_0based
):

    numero = (
        indice_0based
        +
        1
    )

    letras = ""

    while numero:

        numero, resto = divmod(
            numero - 1,
            26
        )

        letras = (
            chr(
                65
                +
                resto
            )
            +
            letras
        )

    return letras


# ============================================================
# 25. METADATOS RAW XLS/OLE
# ============================================================

def inspeccionar_celda_xls(
    libro,
    hoja,
    fila,
    columna
):

    celda = hoja.cell(
        fila,
        columna
    )

    xf_index = hoja.cell_xf_index(
        fila,
        columna
    )

    format_key = None
    format_str = None

    if (
        0 <= xf_index
        <
        len(
            libro.xf_list
        )
    ):

        xf = libro.xf_list[
            xf_index
        ]

        format_key = getattr(
            xf,
            "format_key",
            None
        )

        if format_key is not None:

            formato_obj = libro.format_map.get(
                format_key
            )

            if formato_obj is not None:

                format_str = getattr(
                    formato_obj,
                    "format_str",
                    None
                )

    return {

        "valor_raw":
            celda.value,

        "valor_raw_repr":
            repr(
                celda.value
            ),

        "ctype":
            celda.ctype,

        "xf_index":
            xf_index,

        "format_key":
            format_key,

        "format_str":
            format_str,

        "decimales":
            inferir_decimales_formato_excel(
                format_str
            ),
    }


# ============================================================
# 26. METADATOS RAW XLSX
# ============================================================

def inspeccionar_celda_xlsx(
    hoja,
    fila_0based,
    columna_0based
):

    celda = hoja.cell(
        row=fila_0based + 1,
        column=columna_0based + 1
    )

    format_str = (
        celda.number_format
    )

    return {

        "valor_raw":
            celda.value,

        "valor_raw_repr":
            repr(
                celda.value
            ),

        "ctype":
            str(
                celda.data_type
            ),

        "xf_index":
            celda.style_id,

        "format_key":
            "",

        "format_str":
            format_str,

        "decimales":
            inferir_decimales_formato_excel(
                format_str
            ),
    }


# ============================================================
# 27. OBTENER RANGOS COMBINADOS - XLS
# ============================================================

def rangos_combinados_xls(
    hoja
):

    resultados = []

    for (
        rlo,
        rhi,
        clo,
        chi
    ) in hoja.merged_cells:

        resultados.append({

            "fila_inicio":
                rlo,

            "fila_fin_exclusiva":
                rhi,

            "col_inicio":
                clo,

            "col_fin_exclusiva":
                chi,

            "descripcion":
                (
                    f"filas {rlo}:{rhi - 1}, "
                    f"columnas {clo}:{chi - 1}"
                ),
        })

    return resultados


# ============================================================
# 28. OBTENER RANGOS COMBINADOS - XLSX
# ============================================================

def rangos_combinados_xlsx(
    hoja
):

    resultados = []

    for rango in hoja.merged_cells.ranges:

        resultados.append({

            "fila_inicio":
                rango.min_row - 1,

            "fila_fin_exclusiva":
                rango.max_row,

            "col_inicio":
                rango.min_col - 1,

            "col_fin_exclusiva":
                rango.max_col,

            "descripcion":
                str(
                    rango
                ),
        })

    return resultados


# ============================================================
# 29. INTERSECCIÓN CON RANGO COMBINADO
# ============================================================

def rango_intersecta_celda(
    rango,
    fila,
    columna
):

    return (
        rango[
            "fila_inicio"
        ]
        <=
        fila
        <
        rango[
            "fila_fin_exclusiva"
        ]
        and
        rango[
            "col_inicio"
        ]
        <=
        columna
        <
        rango[
            "col_fin_exclusiva"
        ]
    )


def rangos_intersectan_bloque(
    rangos,
    fila_inicio,
    fila_fin,
    col_inicio,
    col_fin
):

    encontrados = []

    for rango in rangos:

        interseccion_filas = not (
            rango[
                "fila_fin_exclusiva"
            ]
            <=
            fila_inicio
            or
            rango[
                "fila_inicio"
            ]
            >
            fila_fin
        )

        interseccion_columnas = not (
            rango[
                "col_fin_exclusiva"
            ]
            <=
            col_inicio
            or
            rango[
                "col_inicio"
            ]
            >
            col_fin
        )

        if (
            interseccion_filas
            and
            interseccion_columnas
        ):

            encontrados.append(
                rango
            )

    return encontrados


# ============================================================
# 30. INSPECCIÓN RAW COMPLETA SEGÚN FORMATO
# ============================================================

def inspeccionar_archivo_raw(
    ruta,
    resultado
):

    formato = resultado[
        "formato"
    ]

    fila_bancos = resultado[
        "fila_bancos"
    ]

    fila_encabezados = resultado[
        "fila_encabezados"
    ]

    fila_total = resultado[
        "fila_total_creditos"
    ]

    bloques = resultado[
        "estructura"
    ][
        "bloques"
    ]

    salida_bloques = []


    # ========================================================
    # XLS/OLE
    # ========================================================

    if formato == "XLS_OLE":

        try:

            libro = xlrd.open_workbook(
                filename=str(
                    ruta
                ),
                formatting_info=True,
                on_demand=True
            )

        except Exception as error:

            raise ErrorDiagnostico(
                "No pudo abrirse el XLS con "
                "formatting_info=True."
            ) from error

        try:

            hoja = libro.sheet_by_name(
                resultado[
                    "hoja"
                ]
            )

            rangos = rangos_combinados_xls(
                hoja
            )

            for indice, bloque in enumerate(
                bloques
            ):

                item = {

                    "indice_bloque":
                        indice,

                    "banco_original":
                        bloque[
                            "banco_original"
                        ],

                    "es_agregado":
                        es_agregado_creditos(
                            bloque[
                                "banco_original"
                            ]
                        ),

                    "columna_mn":
                        bloque[
                            "columna_mn"
                        ],

                    "columna_me":
                        bloque[
                            "columna_me"
                        ],

                    "columna_total":
                        bloque[
                            "columna_total"
                        ],

                    "columna_mn_excel":
                        columna_excel(
                            bloque[
                                "columna_mn"
                            ]
                        ),

                    "columna_me_excel":
                        columna_excel(
                            bloque[
                                "columna_me"
                            ]
                        ),

                    "columna_total_excel":
                        columna_excel(
                            bloque[
                                "columna_total"
                            ]
                        ),

                    "nombre_raw_col_mn":
                        hoja.cell_value(
                            fila_bancos,
                            bloque[
                                "columna_mn"
                            ]
                        ),

                    "nombre_raw_col_me":
                        hoja.cell_value(
                            fila_bancos,
                            bloque[
                                "columna_me"
                            ]
                        ),

                    "nombre_raw_col_total":
                        hoja.cell_value(
                            fila_bancos,
                            bloque[
                                "columna_total"
                            ]
                        ),

                    "encabezado_mn_raw":
                        hoja.cell_value(
                            fila_encabezados,
                            bloque[
                                "columna_mn"
                            ]
                        ),

                    "encabezado_me_raw":
                        hoja.cell_value(
                            fila_encabezados,
                            bloque[
                                "columna_me"
                            ]
                        ),

                    "encabezado_total_raw":
                        hoja.cell_value(
                            fila_encabezados,
                            bloque[
                                "columna_total"
                            ]
                        ),
                }

                for variable, columna in (
                    (
                        "mn",
                        bloque[
                            "columna_mn"
                        ]
                    ),
                    (
                        "me",
                        bloque[
                            "columna_me"
                        ]
                    ),
                    (
                        "total",
                        bloque[
                            "columna_total"
                        ]
                    ),
                ):

                    meta = inspeccionar_celda_xls(
                        libro,
                        hoja,
                        fila_total,
                        columna
                    )

                    for clave, valor in meta.items():

                        item[
                            f"{variable}_{clave}"
                        ] = valor

                rangos_bloque = rangos_intersectan_bloque(
                    rangos,
                    fila_bancos,
                    fila_total,
                    bloque[
                        "columna_mn"
                    ],
                    bloque[
                        "columna_total"
                    ]
                )

                item[
                    "rangos_combinados_intersectan"
                ] = [
                    r[
                        "descripcion"
                    ]
                    for r in rangos_bloque
                ]

                salida_bloques.append(
                    item
                )

            return {

                "bloques":
                    salida_bloques,

                "rangos_combinados":
                    rangos,

                "nrows":
                    hoja.nrows,

                "ncols":
                    hoja.ncols,
            }

        finally:

            libro.release_resources()


    # ========================================================
    # XLSX/ZIP
    # ========================================================

    if formato == "XLSX_ZIP":

        try:

            with Path(
                ruta
            ).open(
                "rb"
            ) as archivo_binario:

                libro = load_workbook(
                    archivo_binario,
                    read_only=False,
                    data_only=True
                )

                try:

                    hoja = libro[
                        resultado[
                            "hoja"
                        ]
                    ]

                    rangos = rangos_combinados_xlsx(
                        hoja
                    )

                    for indice, bloque in enumerate(
                        bloques
                    ):

                        item = {

                            "indice_bloque":
                                indice,

                            "banco_original":
                                bloque[
                                    "banco_original"
                                ],

                            "es_agregado":
                                es_agregado_creditos(
                                    bloque[
                                        "banco_original"
                                    ]
                                ),

                            "columna_mn":
                                bloque[
                                    "columna_mn"
                                ],

                            "columna_me":
                                bloque[
                                    "columna_me"
                                ],

                            "columna_total":
                                bloque[
                                    "columna_total"
                                ],

                            "columna_mn_excel":
                                columna_excel(
                                    bloque[
                                        "columna_mn"
                                    ]
                                ),

                            "columna_me_excel":
                                columna_excel(
                                    bloque[
                                        "columna_me"
                                    ]
                                ),

                            "columna_total_excel":
                                columna_excel(
                                    bloque[
                                        "columna_total"
                                    ]
                                ),

                            "nombre_raw_col_mn":
                                hoja.cell(
                                    row=fila_bancos + 1,
                                    column=(
                                        bloque[
                                            "columna_mn"
                                        ]
                                        +
                                        1
                                    )
                                ).value,

                            "nombre_raw_col_me":
                                hoja.cell(
                                    row=fila_bancos + 1,
                                    column=(
                                        bloque[
                                            "columna_me"
                                        ]
                                        +
                                        1
                                    )
                                ).value,

                            "nombre_raw_col_total":
                                hoja.cell(
                                    row=fila_bancos + 1,
                                    column=(
                                        bloque[
                                            "columna_total"
                                        ]
                                        +
                                        1
                                    )
                                ).value,

                            "encabezado_mn_raw":
                                hoja.cell(
                                    row=fila_encabezados + 1,
                                    column=(
                                        bloque[
                                            "columna_mn"
                                        ]
                                        +
                                        1
                                    )
                                ).value,

                            "encabezado_me_raw":
                                hoja.cell(
                                    row=fila_encabezados + 1,
                                    column=(
                                        bloque[
                                            "columna_me"
                                        ]
                                        +
                                        1
                                    )
                                ).value,

                            "encabezado_total_raw":
                                hoja.cell(
                                    row=fila_encabezados + 1,
                                    column=(
                                        bloque[
                                            "columna_total"
                                        ]
                                        +
                                        1
                                    )
                                ).value,
                        }

                        for variable, columna in (
                            (
                                "mn",
                                bloque[
                                    "columna_mn"
                                ]
                            ),
                            (
                                "me",
                                bloque[
                                    "columna_me"
                                ]
                            ),
                            (
                                "total",
                                bloque[
                                    "columna_total"
                                ]
                            ),
                        ):

                            meta = inspeccionar_celda_xlsx(
                                hoja,
                                fila_total,
                                columna
                            )

                            for clave, valor in meta.items():

                                item[
                                    f"{variable}_{clave}"
                                ] = valor

                        rangos_bloque = rangos_intersectan_bloque(
                            rangos,
                            fila_bancos,
                            fila_total,
                            bloque[
                                "columna_mn"
                            ],
                            bloque[
                                "columna_total"
                            ]
                        )

                        item[
                            "rangos_combinados_intersectan"
                        ] = [
                            r[
                                "descripcion"
                            ]
                            for r in rangos_bloque
                        ]

                        salida_bloques.append(
                            item
                        )

                    return {

                        "bloques":
                            salida_bloques,

                        "rangos_combinados":
                            rangos,

                        "nrows":
                            hoja.max_row,

                        "ncols":
                            hoja.max_column,
                    }

                finally:

                    libro.close()

        except Exception as error:

            raise ErrorDiagnostico(
                "No pudo inspeccionarse XLSX "
                "con estilos y celdas combinadas."
            ) from error

    raise ErrorDiagnostico(
        f"Formato no soportado: {formato}"
    )


# ============================================================
# 31. BUSCAR RAW DE UN BLOQUE POR ÍNDICE
# ============================================================

def obtener_bloque_raw(
    inspeccion,
    indice
):

    coincidencias = [
        bloque
        for bloque in inspeccion[
            "bloques"
        ]
        if bloque[
            "indice_bloque"
        ]
        ==
        indice
    ]

    if len(
        coincidencias
    ) != 1:

        raise ErrorDiagnostico(
            f"No pudo recuperarse bloque raw índice {indice}."
        )

    return coincidencias[
        0
    ]


# ============================================================
# 32. OBTENER REGISTRO DEL PARSER POR ÍNDICE
# ============================================================

def obtener_registro_parser_por_indice(
    resultado,
    indice
):

    registros = resultado[
        "registros"
    ]

    if not (
        0
        <=
        indice
        <
        len(
            registros
        )
    ):

        return None

    return registros[
        indice
    ]


# ============================================================
# 33. DESCRIBIR BLOQUE
# ============================================================

def mostrar_bloque(
    etiqueta,
    raw,
    parser=None
):

    subtitulo(
        etiqueta
    )

    if raw is None:

        escribir(
            "<NO EXISTE>"
        )

        return

    escribir(
        f"Índice bloque = "
        f"{raw['indice_bloque']}"
    )

    escribir(
        f"Nombre exacto SBS = "
        f"{raw['banco_original']!r}"
    )

    escribir(
        "Columnas 0-based MN / ME / Total = "
        f"{raw['columna_mn']} / "
        f"{raw['columna_me']} / "
        f"{raw['columna_total']}"
    )

    escribir(
        "Columnas Excel MN / ME / Total = "
        f"{raw['columna_mn_excel']} / "
        f"{raw['columna_me_excel']} / "
        f"{raw['columna_total_excel']}"
    )

    escribir(
        "Nombre raw en fila bancos = "
        f"{raw['nombre_raw_col_mn']!r} | "
        f"{raw['nombre_raw_col_me']!r} | "
        f"{raw['nombre_raw_col_total']!r}"
    )

    escribir(
        "Encabezados raw = "
        f"{raw['encabezado_mn_raw']!r} | "
        f"{raw['encabezado_me_raw']!r} | "
        f"{raw['encabezado_total_raw']!r}"
    )

    escribir(
        f"MN raw = "
        f"{raw['mn_valor_raw_repr']}"
    )

    escribir(
        f"ME raw = "
        f"{raw['me_valor_raw_repr']}"
    )

    escribir(
        f"Total raw = "
        f"{raw['total_valor_raw_repr']}"
    )

    escribir(
        f"MN convertido raw = "
        f"{convertir_numero(raw['mn_valor_raw'])!r}"
    )

    escribir(
        f"ME convertido raw = "
        f"{convertir_numero(raw['me_valor_raw'])!r}"
    )

    escribir(
        f"Total convertido raw = "
        f"{convertir_numero(raw['total_valor_raw'])!r}"
    )

    escribir(
        f"MN formato = "
        f"{raw['mn_format_str']!r}"
    )

    escribir(
        f"ME formato = "
        f"{raw['me_format_str']!r}"
    )

    escribir(
        f"Total formato = "
        f"{raw['total_format_str']!r}"
    )

    escribir(
        "Decimales visibles MN / ME / Total = "
        f"{raw['mn_decimales']} / "
        f"{raw['me_decimales']} / "
        f"{raw['total_decimales']}"
    )

    escribir(
        "Rangos combinados que intersectan "
        "este bloque desde fila bancos hasta "
        "Total Créditos = "
        f"{raw['rangos_combinados_intersectan']}"
    )

    if parser is not None:

        escribir()
        escribir(
            "Valores producidos por parser:"
        )

        escribir(
            f"  MN = "
            f"{parser['mn_soles_miles']!r}"
        )

        escribir(
            f"  ME = "
            f"{parser['me_usd_miles']!r}"
        )

        escribir(
            f"  Total = "
            f"{parser['total_soles_miles']!r}"
        )

        escribir(
            f"  DOLCRED = "
            f"{parser['dolcred_pct']!r}"
        )

        escribir(
            f"  TC implícito = "
            f"{parser['tc_implicito']!r}"
        )

        escribir(
            f"  Problemas = "
            f"{parser['problemas']}"
        )


# ============================================================
# 34. IDENTIFICAR CONCEPTO DE UNA FILA
# ============================================================

def obtener_concepto_fila(
    fila,
    primera_columna_bancos
):

    textos = []

    limite = min(
        primera_columna_bancos,
        len(
            fila
        )
    )

    for columna in range(
        limite
    ):

        valor = fila[
            columna
        ]

        if normalizar_texto(
            valor
        ):

            textos.append(
                str(
                    valor
                ).strip()
            )

    if textos:

        return " | ".join(
            textos
        )

    # Respaldo: revisar primeras 6 columnas.
    for valor in fila[
        :6
    ]:

        if normalizar_texto(
            valor
        ):

            textos.append(
                str(
                    valor
                ).strip()
            )

    return " | ".join(
        textos
    )


# ============================================================
# 35. FILAS CERCANAS A TOTAL CRÉDITOS
# ============================================================

def construir_filas_cercanas(
    resultado,
    bloque_objetivo
):

    filas = resultado[
        "filas"
    ]

    fila_total = resultado[
        "fila_total_creditos"
    ]

    primera_columna_bancos = min(
        bloque[
            "columna_mn"
        ]
        for bloque in resultado[
            "estructura"
        ][
            "bloques"
        ]
    )

    inicio = max(
        0,
        fila_total
        -
        RADIO_FILAS_CERCANAS
    )

    fin = min(
        len(
            filas
        )
        -
        1,
        fila_total
        +
        RADIO_FILAS_CERCANAS
    )

    salida = []

    for numero_fila in range(
        inicio,
        fin + 1
    ):

        fila = filas[
            numero_fila
        ]

        def valor_seguro(columna):

            if columna < len(
                fila
            ):

                return fila[
                    columna
                ]

            return None

        mn_raw = valor_seguro(
            bloque_objetivo[
                "columna_mn"
            ]
        )

        me_raw = valor_seguro(
            bloque_objetivo[
                "columna_me"
            ]
        )

        total_raw = valor_seguro(
            bloque_objetivo[
                "columna_total"
            ]
        )

        concepto = obtener_concepto_fila(
            fila,
            primera_columna_bancos
        )

        salida.append({

            "fila_0based":
                numero_fila,

            "fila_excel":
                numero_fila + 1,

            "es_fila_total_creditos":
                (
                    numero_fila
                    ==
                    fila_total
                ),

            "concepto_raw":
                concepto,

            "mn_raw":
                mn_raw,

            "mn_raw_repr":
                repr(
                    mn_raw
                ),

            "mn_convertido":
                convertir_numero(
                    mn_raw
                ),

            "me_raw":
                me_raw,

            "me_raw_repr":
                repr(
                    me_raw
                ),

            "me_convertido":
                convertir_numero(
                    me_raw
                ),

            "total_raw":
                total_raw,

            "total_raw_repr":
                repr(
                    total_raw
                ),

            "total_convertido":
                convertir_numero(
                    total_raw
                ),
        })

    return salida


# ============================================================
# 36. MOSTRAR FILAS CERCANAS
# ============================================================

def mostrar_filas_cercanas(
    filas_cercanas
):

    titulo(
        "BANCO AZTECA - FILAS CERCANAS A TOTAL CRÉDITOS"
    )

    for fila in filas_cercanas:

        marcador = (
            " <<< FILA USADA POR EL PARSER"
            if fila[
                "es_fila_total_creditos"
            ]
            else ""
        )

        escribir()

        escribir(
            f"Fila 0-based = "
            f"{fila['fila_0based']} | "
            f"Fila Excel = "
            f"{fila['fila_excel']}"
            f"{marcador}"
        )

        escribir(
            f"  Concepto = "
            f"{fila['concepto_raw']!r}"
        )

        escribir(
            f"  MN raw = "
            f"{fila['mn_raw_repr']} | "
            f"convertido = "
            f"{fila['mn_convertido']!r}"
        )

        escribir(
            f"  ME raw = "
            f"{fila['me_raw_repr']} | "
            f"convertido = "
            f"{fila['me_convertido']!r}"
        )

        escribir(
            f"  Total raw = "
            f"{fila['total_raw_repr']} | "
            f"convertido = "
            f"{fila['total_convertido']!r}"
        )

        mn = fila[
            "mn_convertido"
        ]

        me = fila[
            "me_convertido"
        ]

        total = fila[
            "total_convertido"
        ]

        if (
            mn is not None
            and
            total is not None
        ):

            escribir(
                f"  Total - MN = "
                f"{total - mn!r}"
            )

        if (
            mn is not None
            and
            me is not None
            and
            total is not None
            and
            me > 0
        ):

            escribir(
                f"  TC implícito fila = "
                f"{((total - mn) / me)!r}"
            )


# ============================================================
# 37. COMPARACIÓN PARSER VS RAW BANCO AZTECA
# ============================================================

def comparar_parser_raw(
    registro,
    raw
):

    comparaciones = []

    configuracion = (
        (
            "MN",
            "mn_soles_miles",
            "mn_valor_raw",
        ),
        (
            "ME",
            "me_usd_miles",
            "me_valor_raw",
        ),
        (
            "TOTAL",
            "total_soles_miles",
            "total_valor_raw",
        ),
    )

    for (
        variable,
        campo_parser,
        campo_raw
    ) in configuracion:

        valor_parser = registro[
            campo_parser
        ]

        valor_raw_convertido = convertir_numero(
            raw[
                campo_raw
            ]
        )

        coincide = (
            valor_parser
            ==
            valor_raw_convertido
        )

        comparaciones.append({

            "variable":
                variable,

            "valor_parser":
                valor_parser,

            "valor_raw_convertido":
                valor_raw_convertido,

            "coincide":
                coincide,
        })

    return comparaciones


# ============================================================
# 38. COMPARACIÓN CON TOTAL BANCA MÚLTIPLE
# ============================================================

def comparar_con_agregado(
    azteca,
    agregado
):

    salida = []

    for variable, campo in (
        (
            "MN",
            "mn_soles_miles"
        ),
        (
            "ME",
            "me_usd_miles"
        ),
        (
            "TOTAL",
            "total_soles_miles"
        ),
    ):

        valor_banco = azteca[
            campo
        ]

        valor_agregado = agregado[
            campo
        ]

        porcentaje = None

        if (
            valor_banco is not None
            and
            valor_agregado is not None
            and
            valor_agregado != 0
        ):

            porcentaje = (
                valor_banco
                /
                valor_agregado
                *
                100
            )

        salida.append({

            "variable":
                variable,

            "banco":
                valor_banco,

            "agregado":
                valor_agregado,

            "participacion_pct":
                porcentaje,
        })

    return salida


# ============================================================
# 39. COMPROBAR ALINEACIÓN DE BLOQUES
# ============================================================

def diagnosticar_alineacion(
    resultado,
    indice_objetivo
):

    bloques = resultado[
        "estructura"
    ][
        "bloques"
    ]

    bloque = bloques[
        indice_objetivo
    ]

    alertas = []

    # MN | ME | Total deben ser consecutivos.
    if not (
        bloque[
            "columna_me"
        ]
        ==
        bloque[
            "columna_mn"
        ]
        +
        1
        and
        bloque[
            "columna_total"
        ]
        ==
        bloque[
            "columna_mn"
        ]
        +
        2
    ):

        alertas.append(
            "BLOQUE_AZTECA_NO_ES_TRIPLETA_CONSECUTIVA"
        )

    # Comparar separación respecto del bloque anterior.
    if indice_objetivo > 0:

        anterior = bloques[
            indice_objetivo - 1
        ]

        separacion_anterior = (
            bloque[
                "columna_mn"
            ]
            -
            anterior[
                "columna_total"
            ]
        )

        if separacion_anterior != 1:

            alertas.append(
                "SEPARACION_RESPECTO_BLOQUE_ANTERIOR="
                f"{separacion_anterior}"
            )

    else:

        separacion_anterior = None

    # Comparar separación respecto del bloque siguiente.
    if (
        indice_objetivo
        <
        len(
            bloques
        )
        -
        1
    ):

        siguiente = bloques[
            indice_objetivo + 1
        ]

        separacion_siguiente = (
            siguiente[
                "columna_mn"
            ]
            -
            bloque[
                "columna_total"
            ]
        )

        if separacion_siguiente != 1:

            alertas.append(
                "SEPARACION_RESPECTO_BLOQUE_SIGUIENTE="
                f"{separacion_siguiente}"
            )

    else:

        separacion_siguiente = None

    # Revisar columnas duplicadas/solapadas de todos los bloques.
    columnas = []

    for b in bloques:

        columnas.extend([
            b[
                "columna_mn"
            ],
            b[
                "columna_me"
            ],
            b[
                "columna_total"
            ],
        ])

    repetidas = sorted({
        columna
        for columna in columnas
        if columnas.count(
            columna
        )
        >
        1
    })

    if repetidas:

        alertas.append(
            f"COLUMNAS_SOLAPADAS={repetidas}"
        )

    return {

        "separacion_anterior":
            separacion_anterior,

        "separacion_siguiente":
            separacion_siguiente,

        "columnas_solapadas":
            repetidas,

        "alertas":
            alertas,
    }


# ============================================================
# 40. GUARDAR CSV BLOQUES
# ============================================================

def guardar_csv_bloques(
    bloques
):

    columnas = [
        "indice_bloque",
        "banco_original",
        "es_agregado",
        "columna_mn",
        "columna_me",
        "columna_total",
        "columna_mn_excel",
        "columna_me_excel",
        "columna_total_excel",
        "nombre_raw_col_mn",
        "nombre_raw_col_me",
        "nombre_raw_col_total",
        "encabezado_mn_raw",
        "encabezado_me_raw",
        "encabezado_total_raw",
        "mn_valor_raw",
        "mn_valor_raw_repr",
        "mn_ctype",
        "mn_format_str",
        "mn_decimales",
        "me_valor_raw",
        "me_valor_raw_repr",
        "me_ctype",
        "me_format_str",
        "me_decimales",
        "total_valor_raw",
        "total_valor_raw_repr",
        "total_ctype",
        "total_format_str",
        "total_decimales",
        "rangos_combinados_intersectan",
    ]

    with CSV_BLOQUES.open(
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as archivo:

        escritor = csv.DictWriter(
            archivo,
            fieldnames=columnas,
            extrasaction="ignore"
        )

        escritor.writeheader()

        for fila in bloques:

            copia = dict(
                fila
            )

            copia[
                "rangos_combinados_intersectan"
            ] = " | ".join(
                fila[
                    "rangos_combinados_intersectan"
                ]
            )

            escritor.writerow(
                copia
            )


# ============================================================
# 41. GUARDAR CSV FILAS CERCANAS
# ============================================================

def guardar_csv_filas_cercanas(
    filas
):

    columnas = [
        "fila_0based",
        "fila_excel",
        "es_fila_total_creditos",
        "concepto_raw",
        "mn_raw",
        "mn_raw_repr",
        "mn_convertido",
        "me_raw",
        "me_raw_repr",
        "me_convertido",
        "total_raw",
        "total_raw_repr",
        "total_convertido",
    ]

    with CSV_FILAS_CERCANAS.open(
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as archivo:

        escritor = csv.DictWriter(
            archivo,
            fieldnames=columnas,
            extrasaction="ignore"
        )

        escritor.writeheader()

        for fila in filas:

            escritor.writerow(
                fila
            )


# ============================================================
# 42. MAIN
# ============================================================

def main():

    try:

        titulo(
            "DIAGNÓSTICO LOCAL "
            "Y / B-2359 / 2021-09 / BANCO AZTECA"
        )

        escribir(
            "No se realizará ninguna descarga."
        )

        escribir(
            "No se modifica producción."
        )

        escribir(
            "No se modifica el parser."
        )

        escribir(
            "No se modifica DOLCRED."
        )

        escribir(
            "No se modifica ninguna tolerancia."
        )

        escribir()

        escribir(
            f"Archivo local esperado = "
            f"{ARCHIVO_LOCAL}"
        )

        # ====================================================
        # A. VALIDAR ARCHIVO LOCAL
        # ====================================================

        if not ARCHIVO_LOCAL.exists():

            raise ErrorDiagnostico(
                "No existe el archivo local esperado: "
                f"{ARCHIVO_LOCAL}"
            )

        if (
            ARCHIVO_LOCAL.stat().st_size
            <=
            0
        ):

            raise ErrorDiagnostico(
                "El archivo local existe "
                "pero tiene 0 bytes."
            )

        escribir(
            f"Tamaño = "
            f"{ARCHIVO_LOCAL.stat().st_size} bytes"
        )

        escribir(
            f"SHA-256 = "
            f"{sha256_archivo(ARCHIVO_LOCAL)}"
        )

        formato = detectar_formato(
            ARCHIVO_LOCAL
        )

        escribir(
            f"Formato real = "
            f"{formato}"
        )

        # ====================================================
        # B. EJECUTAR MISMO PARSER Y
        # ====================================================

        resultado = extraer_y_creditos(
            ARCHIVO_LOCAL
        )

        titulo(
            "ESTRUCTURA GENERAL DETECTADA"
        )

        escribir(
            f"Hoja elegida = "
            f"{resultado['hoja']!r}"
        )

        escribir(
            f"Fila de bancos 0-based = "
            f"{resultado['fila_bancos']}"
        )

        escribir(
            f"Fila de bancos Excel = "
            f"{resultado['fila_bancos'] + 1}"
        )

        escribir(
            f"Fila encabezados 0-based = "
            f"{resultado['fila_encabezados']}"
        )

        escribir(
            f"Fila encabezados Excel = "
            f"{resultado['fila_encabezados'] + 1}"
        )

        escribir(
            f"Fila Total Créditos 0-based = "
            f"{resultado['fila_total_creditos']}"
        )

        escribir(
            f"Fila Total Créditos Excel = "
            f"{resultado['fila_total_creditos'] + 1}"
        )

        escribir(
            f"Cantidad total de bloques = "
            f"{resultado['cantidad_bloques']}"
        )

        # ====================================================
        # C. LOCALIZAR BANCO AZTECA
        # ====================================================

        azteca = localizar_banco_objetivo(
            resultado
        )

        indice_azteca = localizar_indice_bloque(
            resultado,
            azteca
        )

        escribir(
            f"Índice Banco Azteca = "
            f"{indice_azteca}"
        )

        # ====================================================
        # D. INSPECCIÓN RAW / FORMATOS / MERGED CELLS
        # ====================================================

        inspeccion = inspeccionar_archivo_raw(
            ARCHIVO_LOCAL,
            resultado
        )

        raw_azteca = obtener_bloque_raw(
            inspeccion,
            indice_azteca
        )

        indice_anterior = (
            indice_azteca
            -
            1
        )

        indice_siguiente = (
            indice_azteca
            +
            1
        )

        raw_anterior = (
            obtener_bloque_raw(
                inspeccion,
                indice_anterior
            )
            if indice_anterior >= 0
            else None
        )

        raw_siguiente = (
            obtener_bloque_raw(
                inspeccion,
                indice_siguiente
            )
            if indice_siguiente
            <
            len(
                inspeccion[
                    "bloques"
                ]
            )
            else None
        )

        parser_anterior = (
            obtener_registro_parser_por_indice(
                resultado,
                indice_anterior
            )
            if indice_anterior >= 0
            else None
        )

        parser_siguiente = (
            obtener_registro_parser_por_indice(
                resultado,
                indice_siguiente
            )
            if indice_siguiente
            <
            len(
                resultado[
                    "registros"
                ]
            )
            else None
        )

        # ====================================================
        # E. BLOQUE AZTECA
        # ====================================================

        titulo(
            "BANCO AZTECA - BLOQUE DETECTADO"
        )

        escribir(
            f"Nombre exacto SBS = "
            f"{azteca['banco_original']!r}"
        )

        escribir(
            "Columnas 0-based MN / ME / Total = "
            f"{azteca['columna_mn']} / "
            f"{azteca['columna_me']} / "
            f"{azteca['columna_total']}"
        )

        escribir(
            "Columnas Excel MN / ME / Total = "
            f"{columna_excel(azteca['columna_mn'])} / "
            f"{columna_excel(azteca['columna_me'])} / "
            f"{columna_excel(azteca['columna_total'])}"
        )

        escribir()

        escribir(
            "ENCABEZADOS RAW:"
        )

        escribir(
            f"  MN = "
            f"{raw_azteca['encabezado_mn_raw']!r}"
        )

        escribir(
            f"  ME = "
            f"{raw_azteca['encabezado_me_raw']!r}"
        )

        escribir(
            f"  Total = "
            f"{raw_azteca['encabezado_total_raw']!r}"
        )

        escribir()

        escribir(
            "VALORES RAW FILA TOTAL CRÉDITOS:"
        )

        escribir(
            f"  MN raw = "
            f"{raw_azteca['mn_valor_raw_repr']}"
        )

        escribir(
            f"  ME raw = "
            f"{raw_azteca['me_valor_raw_repr']}"
        )

        escribir(
            f"  Total raw = "
            f"{raw_azteca['total_valor_raw_repr']}"
        )

        escribir()

        escribir(
            "VALORES CONVERTIDOS POR EL PARSER:"
        )

        escribir(
            f"  MN = "
            f"{azteca['mn_soles_miles']!r}"
        )

        escribir(
            f"  ME = "
            f"{azteca['me_usd_miles']!r}"
        )

        escribir(
            f"  Total = "
            f"{azteca['total_soles_miles']!r}"
        )

        escribir()

        escribir(
            f"DOLCRED calculado = "
            f"{azteca['dolcred_pct']!r}"
        )

        escribir(
            f"TC implícito = "
            f"{azteca['tc_implicito']!r}"
        )

        escribir(
            f"Problemas parser = "
            f"{azteca['problemas']}"
        )

        if (
            azteca[
                "mn_soles_miles"
            ] is not None
            and
            azteca[
                "total_soles_miles"
            ] is not None
        ):

            escribir(
                "Total - MN = "
                f"{(
                    azteca['total_soles_miles']
                    -
                    azteca['mn_soles_miles']
                )!r}"
            )

        escribir()

        escribir(
            "FORMATOS EXCEL:"
        )

        escribir(
            f"  MN formato = "
            f"{raw_azteca['mn_format_str']!r}"
        )

        escribir(
            f"  MN decimales visibles = "
            f"{raw_azteca['mn_decimales']}"
        )

        escribir(
            f"  ME formato = "
            f"{raw_azteca['me_format_str']!r}"
        )

        escribir(
            f"  ME decimales visibles = "
            f"{raw_azteca['me_decimales']}"
        )

        escribir(
            f"  Total formato = "
            f"{raw_azteca['total_format_str']!r}"
        )

        escribir(
            f"  Total decimales visibles = "
            f"{raw_azteca['total_decimales']}"
        )

        escribir()

        escribir(
            "CELDAS COMBINADAS QUE INTERSECTAN "
            "EL BLOQUE BANCO AZTECA:"
        )

        if raw_azteca[
            "rangos_combinados_intersectan"
        ]:

            for rango in raw_azteca[
                "rangos_combinados_intersectan"
            ]:

                escribir(
                    f"  * {rango}"
                )

        else:

            escribir(
                "  Ninguna."
            )

        # ====================================================
        # F. BLOQUE ANTERIOR Y SIGUIENTE
        # ====================================================

        mostrar_bloque(
            "BLOQUE ANTERIOR A BANCO AZTECA",
            raw_anterior,
            parser_anterior
        )

        mostrar_bloque(
            "BLOQUE BANCO AZTECA",
            raw_azteca,
            azteca
        )

        mostrar_bloque(
            "BLOQUE SIGUIENTE A BANCO AZTECA",
            raw_siguiente,
            parser_siguiente
        )

        # ====================================================
        # G. ALINEACIÓN
        # ====================================================

        alineacion = diagnosticar_alineacion(
            resultado,
            indice_azteca
        )

        titulo(
            "CONTROL DE ALINEACIÓN"
        )

        escribir(
            "Separación entre Total del bloque anterior "
            "y MN Banco Azteca = "
            f"{alineacion['separacion_anterior']}"
        )

        escribir(
            "Separación entre Total Banco Azteca "
            "y MN bloque siguiente = "
            f"{alineacion['separacion_siguiente']}"
        )

        escribir(
            f"Columnas solapadas = "
            f"{alineacion['columnas_solapadas']}"
        )

        escribir(
            f"Alertas de alineación = "
            f"{alineacion['alertas']}"
        )

        encabezados_ok = (
            es_mn(
                raw_azteca[
                    "encabezado_mn_raw"
                ]
            )
            and
            es_me(
                raw_azteca[
                    "encabezado_me_raw"
                ]
            )
            and
            es_total(
                raw_azteca[
                    "encabezado_total_raw"
                ]
            )
        )

        escribir(
            "Encabezados forman MN | ME | Total = "
            f"{encabezados_ok}"
        )

        # ====================================================
        # H. PARSER VS RAW
        # ====================================================

        comparaciones = comparar_parser_raw(
            azteca,
            raw_azteca
        )

        titulo(
            "COMPARACIÓN PARSER VS CELDAS RAW"
        )

        for control in comparaciones:

            escribir(
                f"{control['variable']}: "
                f"parser={control['valor_parser']!r} | "
                f"raw convertido="
                f"{control['valor_raw_convertido']!r} | "
                f"coincide={control['coincide']}"
            )

        parser_raw_ok = all(
            control[
                "coincide"
            ]
            for control in comparaciones
        )

        escribir()

        escribir(
            f"TODOS_LOS_VALORES_PARSER_RAW_COINCIDEN = "
            f"{parser_raw_ok}"
        )

        # ====================================================
        # I. TOTAL BANCA MÚLTIPLE
        # ====================================================

        agregados = [
            registro
            for registro in resultado[
                "registros"
            ]
            if registro[
                "es_agregado"
            ]
        ]

        if len(
            agregados
        ) != 1:

            raise ErrorDiagnostico(
                "Se esperaba exactamente un "
                "Total Banca Múltiple."
            )

        agregado = agregados[
            0
        ]

        comparacion_agregado = comparar_con_agregado(
            azteca,
            agregado
        )

        titulo(
            "COMPARACIÓN BANCO AZTECA VS TOTAL BANCA MÚLTIPLE"
        )

        escribir(
            f"Agregado SBS = "
            f"{agregado['banco_original']!r}"
        )

        escribir(
            f"Agregado MN = "
            f"{agregado['mn_soles_miles']!r}"
        )

        escribir(
            f"Agregado ME = "
            f"{agregado['me_usd_miles']!r}"
        )

        escribir(
            f"Agregado Total = "
            f"{agregado['total_soles_miles']!r}"
        )

        escribir(
            f"Agregado DOLCRED = "
            f"{agregado['dolcred_pct']!r}"
        )

        escribir(
            f"Agregado TC implícito = "
            f"{agregado['tc_implicito']!r}"
        )

        escribir()

        for control in comparacion_agregado:

            escribir(
                f"{control['variable']}: "
                f"Banco Azteca="
                f"{control['banco']!r} | "
                f"Total Banca Múltiple="
                f"{control['agregado']!r} | "
                f"participación="
                f"{control['participacion_pct']!r}%"
            )

        # ====================================================
        # J. FILAS CERCANAS
        # ====================================================

        bloque_objetivo = resultado[
            "estructura"
        ][
            "bloques"
        ][
            indice_azteca
        ]

        filas_cercanas = construir_filas_cercanas(
            resultado,
            bloque_objetivo
        )

        mostrar_filas_cercanas(
            filas_cercanas
        )

        # ====================================================
        # K. REVISIÓN DE MERGED CELLS
        # ====================================================

        titulo(
            "REVISIÓN DE CELDAS COMBINADAS"
        )

        escribir(
            f"Cantidad total de rangos combinados "
            f"en hoja = "
            f"{len(inspeccion['rangos_combinados'])}"
        )

        relevantes = rangos_intersectan_bloque(
            inspeccion[
                "rangos_combinados"
            ],
            resultado[
                "fila_bancos"
            ],
            resultado[
                "fila_total_creditos"
            ],
            azteca[
                "columna_mn"
            ],
            azteca[
                "columna_total"
            ]
        )

        if relevantes:

            escribir(
                "Rangos combinados que intersectan "
                "Banco Azteca:"
            )

            for rango in relevantes:

                escribir(
                    f"  * {rango['descripcion']}"
                )

        else:

            escribir(
                "No existen celdas combinadas que "
                "intersecten el bloque Banco Azteca "
                "entre la fila de nombres y "
                "la fila Total Créditos."
            )

        # ====================================================
        # L. CONCLUSIÓN DIAGNÓSTICA AUTOMÁTICA
        # ====================================================

        titulo(
            "CONCLUSIÓN DIAGNÓSTICA AUTOMÁTICA"
        )

        estructura_tripleta_ok = (
            raw_azteca[
                "columna_me"
            ]
            ==
            raw_azteca[
                "columna_mn"
            ]
            +
            1
            and
            raw_azteca[
                "columna_total"
            ]
            ==
            raw_azteca[
                "columna_mn"
            ]
            +
            2
        )

        sin_solapamiento = (
            len(
                alineacion[
                    "columnas_solapadas"
                ]
            )
            ==
            0
        )

        sin_alertas_alineacion = (
            len(
                alineacion[
                    "alertas"
                ]
            )
            ==
            0
        )

        mn_raw = convertir_numero(
            raw_azteca[
                "mn_valor_raw"
            ]
        )

        me_raw = convertir_numero(
            raw_azteca[
                "me_valor_raw"
            ]
        )

        total_raw = convertir_numero(
            raw_azteca[
                "total_valor_raw"
            ]
        )

        fuente_raw_presenta_total_menor_mn = (
            mn_raw is not None
            and
            total_raw is not None
            and
            total_raw < mn_raw
        )

        escribir(
            f"Parser reproduce exactamente raw = "
            f"{parser_raw_ok}"
        )

        escribir(
            f"Tripleta MN|ME|Total consecutiva = "
            f"{estructura_tripleta_ok}"
        )

        escribir(
            f"Encabezados MN|ME|Total correctos = "
            f"{encabezados_ok}"
        )

        escribir(
            f"Sin columnas solapadas = "
            f"{sin_solapamiento}"
        )

        escribir(
            f"Sin alertas de alineación = "
            f"{sin_alertas_alineacion}"
        )

        escribir(
            "Archivo raw SBS presenta "
            "Total < MN para Banco Azteca = "
            f"{fuente_raw_presenta_total_menor_mn}"
        )

        escribir(
            f"MN raw = {mn_raw!r}"
        )

        escribir(
            f"ME raw = {me_raw!r}"
        )

        escribir(
            f"Total raw = {total_raw!r}"
        )

        if (
            parser_raw_ok
            and
            estructura_tripleta_ok
            and
            encabezados_ok
            and
            sin_solapamiento
            and
            fuente_raw_presenta_total_menor_mn
        ):

            escribir()
            escribir(
                "INDICIO PRINCIPAL DEL DIAGNÓSTICO:"
            )

            escribir(
                "Los valores que producen "
                "TOTAL_MENOR_QUE_MN y "
                "DOLCRED_FUERA_RANGO ya están "
                "presentes en las celdas raw "
                "MN/ME/Total que SBS ubica dentro "
                "del bloque Banco Azteca en la "
                "fila Total Créditos."
            )

            escribir(
                "El parser reproduce esas celdas "
                "sin discrepancia."
            )

            escribir(
                "Por tanto, con la evidencia "
                "estructural observada por este script, "
                "el problema NO aparece como una "
                "alteración numérica introducida por "
                "convertir_numero() ni como un "
                "desplazamiento simple de columnas."
            )

            escribir(
                "Las filas cercanas deben revisarse "
                "para determinar si el propio archivo "
                "SBS contiene una particularidad "
                "contable/de presentación para Banco "
                "Azteca en 2021-09."
            )

        else:

            escribir()
            escribir(
                "El diagnóstico encontró una o más "
                "condiciones que impiden atribuir "
                "todavía el problema exclusivamente "
                "al contenido raw SBS."
            )

            escribir(
                "Revisar bloques, merged cells, "
                "alineación y filas cercanas."
            )

        # ====================================================
        # M. GUARDAR CSV
        # ====================================================

        guardar_csv_bloques(
            inspeccion[
                "bloques"
            ]
        )

        guardar_csv_filas_cercanas(
            filas_cercanas
        )

        # ====================================================
        # N. ESTADO
        # ====================================================

        titulo(
            "CLASIFICACIÓN"
        )

        escribir(
            "ESTADO = "
            "DIAGNOSTICO_Y_B2359_2021_09_"
            "BANCO_AZTECA_COMPLETADO"
        )

        escribir()

        escribir(
            "Este estado NO valida ni modifica "
            "Y / 2021-09."
        )

        escribir(
            "No se modificó producción."
        )

        escribir(
            "No se modificó el parser."
        )

        escribir(
            "No se modificó DOLCRED."
        )

        escribir(
            "No se modificó ninguna tolerancia."
        )

    except Exception as error:

        titulo(
            "ERROR DEL DIAGNÓSTICO"
        )

        escribir(
            f"Tipo = "
            f"{type(error).__name__}"
        )

        escribir(
            f"Detalle = "
            f"{repr(error)}"
        )

        escribir(
            "ESTADO = "
            "DIAGNOSTICO_Y_B2359_2021_09_"
            "BANCO_AZTECA_ERROR"
        )

        raise

    finally:

        TXT_RESULTADO.write_text(
            "\n".join(
                lineas_reporte
            ),
            encoding="utf-8"
        )

        print()
        print("=" * 124)
        print("DIAGNÓSTICO FINALIZADO")
        print("=" * 124)

        print()
        print("TXT:")
        print(
            TXT_RESULTADO
        )

        print()
        print("CSV bloques:")
        print(
            CSV_BLOQUES
        )

        print()
        print("CSV filas cercanas:")
        print(
            CSV_FILAS_CERCANAS
        )


# ============================================================
# 43. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":

    main()