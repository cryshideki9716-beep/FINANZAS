# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Validación final del parser de créditos usando seis archivos piloto SBS

from pathlib import Path
import re

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. CONFIGURACIÓN
# ============================================================

RAIZ = Path(__file__).resolve().parent.parent

CARPETA_PILOTO = (
    RAIZ
    / "datos_crudos"
    / "sbs_piloto"
)

ARCHIVOS_CONTROL = [
    "creditos_2008_06.xls",
    "creditos_2010_12.xls",
    "creditos_2011_12.xls",
    "creditos_2015_06.xls",
    "creditos_2020_06.xls",
    "creditos_2025_12.xls",
]


# ============================================================
# 2. EXTRAER FECHA DESDE EL NOMBRE DEL ARCHIVO
# ============================================================

def extraer_fecha_archivo(nombre):

    patron = r"creditos_(\d{4})_(\d{2})"

    coincidencia = re.search(
        patron,
        nombre
    )

    if not coincidencia:

        return None

    year = int(
        coincidencia.group(1)
    )

    month = int(
        coincidencia.group(2)
    )

    return f"{year:04d}-{month:02d}"


# ============================================================
# 3. NORMALIZAR TEXTO
# ============================================================

def normalizar_texto(valor):

    if valor is None:
        return ""

    texto = str(valor)

    texto = texto.replace(
        "\n",
        " "
    )

    texto = texto.replace(
        "\r",
        " "
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


# ============================================================
# 4. CONVERTIR NÚMEROS
# ============================================================

def convertir_numero(valor):

    if valor is None:
        return None

    if isinstance(
        valor,
        (int, float)
    ):

        return float(valor)

    texto = str(valor).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "nd",
        "n/a",
    }:

        return None

    try:

        return float(
            texto.replace(",", "")
        )

    except ValueError:

        return None


# ============================================================
# 5. DETECTAR FORMATO
# ============================================================

def detectar_formato(ruta):

    with open(
        ruta,
        "rb"
    ) as archivo:

        cabecera = archivo.read(
            8
        )

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:
        return "XLS_OLE"

    if cabecera.startswith(
        firmas_zip
    ):
        return "XLSX_ZIP"

    return "DESCONOCIDO"


# ============================================================
# 6. LEER XLS/OLE
# ============================================================

def leer_xls(ruta):

    libro = xlrd.open_workbook(
        filename=str(ruta),
        on_demand=True
    )

    hojas = []

    for nombre in libro.sheet_names():

        hoja = libro.sheet_by_name(
            nombre
        )

        filas = []

        for i in range(
            hoja.nrows
        ):

            fila = [
                hoja.cell_value(i, j)
                for j in range(
                    hoja.ncols
                )
            ]

            filas.append(
                fila
            )

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.release_resources()

    return hojas


# ============================================================
# 7. LEER XLSX/ZIP
# ============================================================

def leer_xlsx(ruta):

    archivo_binario = open(
        ruta,
        "rb"
    )

    libro = load_workbook(
        archivo_binario,
        read_only=True,
        data_only=True
    )

    hojas = []

    for nombre in libro.sheetnames:

        hoja = libro[
            nombre
        ]

        filas = []

        for fila in hoja.iter_rows(
            values_only=True
        ):

            filas.append(
                list(fila)
            )

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.close()
    archivo_binario.close()

    return hojas


# ============================================================
# 8. LEER ARCHIVO
# ============================================================

def leer_archivo(ruta):

    formato = detectar_formato(
        ruta
    )

    if formato == "XLS_OLE":

        hojas = leer_xls(
            ruta
        )

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(
            ruta
        )

    else:

        raise ValueError(
            f"Formato no reconocido: "
            f"{ruta.name}"
        )

    return formato, hojas


# ============================================================
# 9. RECONOCER MN / ME / TOTAL
# ============================================================

def es_mn(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith("mn")
        and "mil" in texto
    )


def es_me(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith("me")
        and "mil" in texto
    )


def es_total(valor):

    texto = normalizar_texto(
        valor
    )

    return texto.startswith(
        "total"
    )


# ============================================================
# 10. ENCONTRAR ENCABEZADOS
# ============================================================

def encontrar_fila_encabezados(
    filas
):

    candidatos = []

    for numero_fila, fila in enumerate(
        filas
    ):

        bloques = []

        for columna in range(
            len(fila) - 2
        ):

            if (
                es_mn(
                    fila[columna]
                )
                and es_me(
                    fila[columna + 1]
                )
                and es_total(
                    fila[columna + 2]
                )
            ):

                bloques.append(
                    columna
                )

        if bloques:

            candidatos.append({
                "fila": numero_fila,
                "bloques": bloques,
            })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x: len(
            x["bloques"]
        )
    )


# ============================================================
# 11. NOMBRE DEL BANCO
# ============================================================

def obtener_nombre_banco(
    filas,
    fila_bancos,
    columna_inicio
):

    fila = filas[
        fila_bancos
    ]

    for desplazamiento in (
        0,
        1,
        2
    ):

        columna = (
            columna_inicio
            + desplazamiento
        )

        if columna >= len(
            fila
        ):
            continue

        valor = fila[
            columna
        ]

        if normalizar_texto(
            valor
        ):

            return str(
                valor
            ).strip()

    return ""


# ============================================================
# 12. DETECTAR ESTRUCTURA
# ============================================================

def detectar_estructura(
    filas
):

    encabezado = (
        encontrar_fila_encabezados(
            filas
        )
    )

    if encabezado is None:
        return None

    fila_encabezados = (
        encabezado["fila"]
    )

    fila_bancos = (
        fila_encabezados - 1
    )

    bloques = []

    for columna_inicio in encabezado[
        "bloques"
    ]:

        banco = obtener_nombre_banco(
            filas,
            fila_bancos,
            columna_inicio
        )

        bloques.append({
            "banco_original": banco,
            "columna_mn": columna_inicio,
            "columna_me": columna_inicio + 1,
            "columna_total": columna_inicio + 2,
        })

    return {
        "fila_bancos": fila_bancos,
        "fila_encabezados": fila_encabezados,
        "bloques": bloques,
    }


# ============================================================
# 13. ELEGIR HOJA PRINCIPAL
# ============================================================

def elegir_hoja_principal(
    hojas
):

    candidatos = []

    for hoja in hojas:

        estructura = (
            detectar_estructura(
                hoja["filas"]
            )
        )

        if estructura is not None:

            candidatos.append({
                "hoja": hoja,
                "estructura": estructura,
            })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x: len(
            x["estructura"][
                "bloques"
            ]
        )
    )


# ============================================================
# 14. BUSCAR FILA TOTAL CRÉDITOS
# ============================================================

def buscar_fila_total_creditos(
    filas
):

    for numero_fila, fila in enumerate(
        filas
    ):

        for valor in fila:

            texto = normalizar_texto(
                valor
            )

            texto_sin_tilde = (
                texto.replace(
                    "é",
                    "e"
                )
            )

            if (
                "total creditos"
                in texto_sin_tilde
            ):

                return numero_fila

    return None


# ============================================================
# 15. EXTRAER Y VALIDAR BANCO-MES
# ============================================================

def validar_archivo(
    ruta
):

    fecha = extraer_fecha_archivo(
        ruta.name
    )

    formato, hojas = leer_archivo(
        ruta
    )

    seleccionado = (
        elegir_hoja_principal(
            hojas
        )
    )

    if seleccionado is None:

        raise ValueError(
            "No se encontró la tabla "
            f"principal en {ruta.name}"
        )

    hoja = seleccionado[
        "hoja"
    ]

    estructura = seleccionado[
        "estructura"
    ]

    filas = hoja[
        "filas"
    ]

    fila_total = (
        buscar_fila_total_creditos(
            filas
        )
    )

    if fila_total is None:

        raise ValueError(
            "'Total Créditos:' no encontrado "
            f"en {ruta.name}"
        )

    fila = filas[
        fila_total
    ]

    registros = []

    for bloque in estructura[
        "bloques"
    ]:

        banco = bloque[
            "banco_original"
        ]

        col_mn = bloque[
            "columna_mn"
        ]

        col_me = bloque[
            "columna_me"
        ]

        col_total = bloque[
            "columna_total"
        ]

        mn = (
            convertir_numero(
                fila[col_mn]
            )
            if col_mn < len(fila)
            else None
        )

        me = (
            convertir_numero(
                fila[col_me]
            )
            if col_me < len(fila)
            else None
        )

        total = (
            convertir_numero(
                fila[col_total]
            )
            if col_total < len(fila)
            else None
        )

        # ----------------------------------------------------
        # Estado inicial
        # ----------------------------------------------------

        problemas = []

        # ----------------------------------------------------
        # Validaciones de existencia
        # ----------------------------------------------------

        if not banco:

            problemas.append(
                "BANCO_SIN_NOMBRE"
            )

        if mn is None:

            problemas.append(
                "MN_FALTANTE"
            )

        if me is None:

            problemas.append(
                "ME_FALTANTE"
            )

        if total is None:

            problemas.append(
                "TOTAL_FALTANTE"
            )

        # ----------------------------------------------------
        # Inicializar variables derivadas
        # ----------------------------------------------------

        componente_me_soles = None
        dolcred = None
        tc_implicito = None

        # ----------------------------------------------------
        # Validaciones económicas
        # ----------------------------------------------------

        if mn is not None:

            if mn < 0:

                problemas.append(
                    "MN_NEGATIVO"
                )

        if me is not None:

            if me < 0:

                problemas.append(
                    "ME_NEGATIVO"
                )

        if total is not None:

            if total <= 0:

                problemas.append(
                    "TOTAL_NO_POSITIVO"
                )

        if (
            mn is not None
            and total is not None
        ):

            if total < mn:

                problemas.append(
                    "TOTAL_MENOR_QUE_MN"
                )

            componente_me_soles = (
                total - mn
            )

        # ----------------------------------------------------
        # DOLCRED
        # ----------------------------------------------------

        if (
            total is not None
            and total > 0
            and mn is not None
        ):

            dolcred = (
                (total - mn)
                / total
                * 100
            )

            if not (
                0 <= dolcred <= 100
            ):

                problemas.append(
                    "DOLCRED_FUERA_RANGO"
                )

        # ----------------------------------------------------
        # TC implícito
        # ----------------------------------------------------

        if (
            mn is not None
            and me is not None
            and total is not None
            and me > 0
        ):

            tc_implicito = (
                (total - mn)
                / me
            )

            if tc_implicito <= 0:

                problemas.append(
                    "TC_NO_POSITIVO"
                )

        # ----------------------------------------------------
        # Clasificación provisional del bloque
        # ----------------------------------------------------

        texto_banco = (
            normalizar_texto(
                banco
            )
        )

        posible_agregado = any(
            palabra in texto_banco
            for palabra in [
                "total",
                "sistema",
                "banca múltiple",
                "banca multiple",
            ]
        )

        registros.append({
            "fecha": fecha,
            "archivo": ruta.name,
            "formato": formato,
            "hoja": hoja["nombre"],
            "fila_total_creditos": fila_total,
            "banco_original": banco,
            "mn_soles_miles": mn,
            "me_usd_miles": me,
            "total_soles_miles": total,
            "componente_me_soles": componente_me_soles,
            "tc_implicito": tc_implicito,
            "dolcred_pct": dolcred,
            "posible_agregado": posible_agregado,
            "problemas": problemas,
        })

    return registros


# ============================================================
# 16. MOSTRAR RESULTADOS
# ============================================================

def mostrar_archivo(
    nombre,
    registros
):

    print()
    print("=" * 105)

    print(
        "ARCHIVO:",
        nombre
    )

    print("=" * 105)

    validos = 0
    problematicos = 0

    for registro in registros:

        print()

        print(
            "Banco:",
            repr(
                registro[
                    "banco_original"
                ]
            )
        )

        print(
            "  MN:",
            registro[
                "mn_soles_miles"
            ]
        )

        print(
            "  ME:",
            registro[
                "me_usd_miles"
            ]
        )

        print(
            "  Total:",
            registro[
                "total_soles_miles"
            ]
        )

        print(
            "  Componente ME en soles:",
            registro[
                "componente_me_soles"
            ]
        )

        print(
            "  TC implícito:",
            registro[
                "tc_implicito"
            ]
        )

        print(
            "  DOLCRED (%):",
            registro[
                "dolcred_pct"
            ]
        )

        print(
            "  ¿Posible agregado?:",
            registro[
                "posible_agregado"
            ]
        )

        print(
            "  Problemas:",
            registro[
                "problemas"
            ]
        )

        if registro[
            "problemas"
        ]:

            problematicos += 1

        else:

            validos += 1

    print()
    print("-" * 105)

    print(
        "Bloques totales:",
        len(registros)
    )

    print(
        "Sin problemas:",
        validos
    )

    print(
        "Con problemas:",
        problematicos
    )


# ============================================================
# 17. FUNCIÓN PRINCIPAL
# ============================================================

def main():

    print()
    print("=" * 105)

    print(
        "VALIDACIÓN FINAL DE CRÉDITOS SBS"
    )

    print(
        "SEIS ARCHIVOS PILOTO"
    )

    print("=" * 105)

    print()
    print(
        "No se descargarán archivos."
    )

    print(
        "No se modificarán datos crudos."
    )

    print(
        "DOLCRED se calcula únicamente "
        "como prueba de consistencia."
    )

    todos = []

    resumen = []

    for nombre in ARCHIVOS_CONTROL:

        ruta = (
            CARPETA_PILOTO
            / nombre
        )

        if not ruta.exists():

            print()
            print(
                "[ERROR] No existe:",
                ruta
            )

            continue

        registros = validar_archivo(
            ruta
        )

        todos.extend(
            registros
        )

        mostrar_archivo(
            nombre,
            registros
        )

        problematicos = [
            r
            for r in registros
            if r["problemas"]
        ]

        agregados = [
            r
            for r in registros
            if r["posible_agregado"]
        ]

        dolcred_validos = [
            r["dolcred_pct"]
            for r in registros
            if r["dolcred_pct"]
            is not None
            and not r["problemas"]
        ]

        resumen.append({
            "archivo": nombre,
            "bloques": len(registros),
            "problemas": len(
                problematicos
            ),
            "agregados": len(
                agregados
            ),
            "dolcred_min": (
                min(dolcred_validos)
                if dolcred_validos
                else None
            ),
            "dolcred_max": (
                max(dolcred_validos)
                if dolcred_validos
                else None
            ),
        })

    # ========================================================
    # RESUMEN GENERAL
    # ========================================================

    print()
    print("=" * 105)

    print(
        "RESUMEN FINAL DE VALIDACIÓN"
    )

    print("=" * 105)

    for item in resumen:

        print()
        print(
            item["archivo"]
        )

        print(
            "  Bloques:",
            item["bloques"]
        )

        print(
            "  Con problemas:",
            item["problemas"]
        )

        print(
            "  Posibles agregados:",
            item["agregados"]
        )

        print(
            "  DOLCRED mínimo:",
            item["dolcred_min"]
        )

        print(
            "  DOLCRED máximo:",
            item["dolcred_max"]
        )

    # ========================================================
    # PROBLEMAS ENCONTRADOS
    # ========================================================

    registros_problema = [
        r
        for r in todos
        if r["problemas"]
    ]

    print()
    print("=" * 105)

    print(
        "REGISTROS CON PROBLEMAS"
    )

    print("=" * 105)

    if not registros_problema:

        print()
        print(
            "[OK] No se detectaron "
            "inconsistencias automáticas."
        )

    else:

        for registro in registros_problema:

            print()
            print(
                registro["fecha"],
                "|",
                repr(
                    registro[
                        "banco_original"
                    ]
                ),
                "|",
                registro["problemas"]
            )

    # ========================================================
    # POSIBLES AGREGADOS
    # ========================================================

    registros_agregados = [
        r
        for r in todos
        if r["posible_agregado"]
    ]

    print()
    print("=" * 105)

    print(
        "POSIBLES AGREGADOS"
    )

    print("=" * 105)

    if not registros_agregados:

        print()
        print(
            "No se detectaron nombres "
            "evidentes de agregados."
        )

    else:

        for registro in registros_agregados:

            print()
            print(
                registro["fecha"],
                "|",
                repr(
                    registro[
                        "banco_original"
                    ]
                )
            )

    print()
    print("=" * 105)

    print(
        "FIN DE LA VALIDACIÓN"
    )

    print("=" * 105)


# ============================================================
# 18. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":
    main()