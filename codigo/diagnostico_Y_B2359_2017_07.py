# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Fecha de diagnóstico: 2026-09-24

from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv
import hashlib
import math
import re
import unicodedata

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. OBJETIVO
# ============================================================

"""
DIAGNÓSTICO SEPARADO Y / SBS B-2359 / 2017-07

Usa exclusivamente el archivo crudo local ya descargado:

    datos_crudos/Y_B2359_2024200485D/
    Y_B-2359_2024200485D_2017-07.xls

NO descarga nada.
NO modifica producción.
NO cambia el parser B-2359.
NO cambia DOLCRED.
NO cambia tolerancias.
NO ejecuta X1.
NO homologa nombres.
NO interpola.
NO selecciona bancos.

Objetivos del diagnóstico:

1. Reproducir exactamente el parser Y de producción.
2. Diagnosticar MN, ME y TOTAL:
   - cantidad de bancos individuales;
   - sum();
   - math.fsum();
   - Total Banca Múltiple SBS raw;
   - diferencia absoluta y relativa;
   - formato Excel del agregado;
   - decimales mostrados;
   - ROUND_HALF_UP de suma y total;
   - coincidencia a precisión visible.
3. Mostrar todos los valores individuales MN:
   - nombre original SBS;
   - valor raw de la celda;
   - valor convertido por el parser;
   - clasificación banco/agregado.
4. Auditar la estructura del mes para detectar:
   - bloques MN|ME|Total;
   - nombres vacíos o duplicados;
   - columnas repetidas/solapadas;
   - celdas no numéricas o faltantes;
   - Total <= 0;
   - celdas numéricas de la fila Total Créditos fuera de los bloques detectados.

Este programa NO declara Y validada.
"""


# ============================================================
# 2. CONFIGURACIÓN
# ============================================================

CODIGO_ESTUDIANTE = "2024200485D"
MES = "2017-07"

RAIZ = Path(__file__).resolve().parent.parent

ARCHIVO_LOCAL = (
    RAIZ
    / "datos_crudos"
    / f"Y_B2359_{CODIGO_ESTUDIANTE}"
    / f"Y_B-2359_{CODIGO_ESTUDIANTE}_{MES}.xls"
)

CARPETA_DIAGNOSTICOS = RAIZ / "salidas" / "diagnosticos"
CARPETA_DIAGNOSTICOS.mkdir(parents=True, exist_ok=True)

TXT_RESULTADO = (
    CARPETA_DIAGNOSTICOS
    / f"diagnostico_Y_B2359_{MES.replace('-', '_')}.txt"
)

CSV_CONTROLES = (
    CARPETA_DIAGNOSTICOS
    / f"diagnostico_Y_B2359_{MES.replace('-', '_')}_controles.csv"
)

CSV_BLOQUES = (
    CARPETA_DIAGNOSTICOS
    / f"diagnostico_Y_B2359_{MES.replace('-', '_')}_bloques.csv"
)


# ============================================================
# 3. EXCEPCIÓN LOCAL DEL DIAGNÓSTICO
# ============================================================

class ErrorDiagnosticoY(Exception):
    pass


# ============================================================
# 4. SALIDA TXT
# ============================================================

lineas_reporte = []


def escribir(texto=""):
    texto = str(texto)
    print(texto)
    lineas_reporte.append(texto)


def titulo(texto):
    escribir()
    escribir("=" * 122)
    escribir(texto)
    escribir("=" * 122)


def subtitulo(texto):
    escribir()
    escribir("-" * 122)
    escribir(texto)
    escribir("-" * 122)


# ============================================================
# 5. UTILIDADES DE TEXTO
# MISMAS BASES DEL SCRIPT DE PRODUCCIÓN
# ============================================================


def normalizar_texto(valor):
    if valor is None:
        return ""

    texto = str(valor)
    texto = texto.replace("\n", " ")
    texto = texto.replace("\r", " ")
    texto = texto.replace("\xa0", " ")
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip().lower()


def normalizar_comparacion(valor):
    texto = normalizar_texto(valor)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(caracter)
    )
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip()


def sin_tildes(valor):
    return normalizar_comparacion(valor)


# ============================================================
# 6. CONVERSIÓN NUMÉRICA
# MISMO CRITERIO DEL PARSER DE PRODUCCIÓN
# ============================================================


def convertir_numero(valor):
    if valor is None:
        return None

    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "n.d",
        "nd",
        "n/a",
        "na",
    }:
        return None

    try:
        return float(texto.replace(",", ""))
    except ValueError:
        return None


# ============================================================
# 7. HASH DEL ARCHIVO LOCAL
# ============================================================


def sha256_archivo(ruta):
    objeto = hashlib.sha256()

    with Path(ruta).open("rb") as archivo:
        while True:
            bloque = archivo.read(1024 * 1024)
            if not bloque:
                break
            objeto.update(bloque)

    return objeto.hexdigest()


# ============================================================
# 8. DETECTAR FORMATO REAL OLE / ZIP
# MISMO CRITERIO DE PRODUCCIÓN
# ============================================================


def detectar_formato(ruta):
    with Path(ruta).open("rb") as archivo:
        cabecera = archivo.read(8)

    firma_ole = bytes.fromhex("D0CF11E0A1B11AE1")
    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:
        return "XLS_OLE"

    if cabecera.startswith(firmas_zip):
        return "XLSX_ZIP"

    raise ErrorDiagnosticoY(
        f"Formato binario no reconocido: {ruta}"
    )


# ============================================================
# 9. LECTURA NORMAL DEL PARSER
# MISMA LÓGICA DE PRODUCCIÓN
# ============================================================


def leer_xls(ruta):
    try:
        libro = xlrd.open_workbook(
            filename=str(ruta),
            on_demand=True,
        )
    except Exception as error:
        raise ErrorDiagnosticoY(
            f"No pudo abrirse como XLS/OLE: {ruta}"
        ) from error

    hojas = []

    try:
        for nombre in libro.sheet_names():
            hoja = libro.sheet_by_name(nombre)

            filas = [
                [
                    hoja.cell_value(i, j)
                    for j in range(hoja.ncols)
                ]
                for i in range(hoja.nrows)
            ]

            hojas.append({
                "nombre": nombre,
                "filas": filas,
            })

    finally:
        libro.release_resources()

    return hojas


def leer_xlsx(ruta):
    try:
        with Path(ruta).open("rb") as archivo_binario:
            libro = load_workbook(
                archivo_binario,
                read_only=True,
                data_only=True,
            )

            hojas = []

            try:
                for nombre in libro.sheetnames:
                    hoja = libro[nombre]

                    filas = [
                        list(fila)
                        for fila in hoja.iter_rows(
                            values_only=True
                        )
                    ]

                    hojas.append({
                        "nombre": nombre,
                        "filas": filas,
                    })

            finally:
                libro.close()

            return hojas

    except ErrorDiagnosticoY:
        raise

    except Exception as error:
        raise ErrorDiagnosticoY(
            f"No pudo abrirse como XLSX/ZIP: {ruta}"
        ) from error


def leer_archivo(ruta):
    formato = detectar_formato(ruta)

    if formato == "XLS_OLE":
        hojas = leer_xls(ruta)

    elif formato == "XLSX_ZIP":
        hojas = leer_xlsx(ruta)

    else:
        raise ErrorDiagnosticoY(
            f"Formato no soportado: {formato}"
        )

    return formato, hojas


# ============================================================
# 10. Y / B-2359 - PARSER VALIDADO
# COPIA LÓGICA DEL PARSER DE PRODUCCIÓN
# ============================================================


def es_mn(valor):
    texto = normalizar_texto(valor)

    return (
        texto.startswith("mn")
        and
        "mil" in texto
    )


def es_me(valor):
    texto = normalizar_texto(valor)

    return (
        texto.startswith("me")
        and
        "mil" in texto
    )


def es_total(valor):
    return normalizar_texto(valor).startswith(
        "total"
    )


def encontrar_fila_encabezados_creditos(filas):
    candidatos = []

    for numero_fila, fila in enumerate(filas):
        bloques = []

        for columna in range(
            max(
                0,
                len(fila) - 2
            )
        ):
            if (
                es_mn(fila[columna])
                and
                es_me(fila[columna + 1])
                and
                es_total(fila[columna + 2])
            ):
                bloques.append(
                    columna
                )

        if bloques:
            candidatos.append({
                "fila": numero_fila,
                "bloques": bloques,
            })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x: len(
            x["bloques"]
        ),
    )


def obtener_nombre_banco_creditos(
    filas,
    fila_bancos,
    columna_inicio,
):
    if (
        fila_bancos < 0
        or
        fila_bancos >= len(filas)
    ):
        return ""

    fila = filas[
        fila_bancos
    ]

    for desplazamiento in (
        0,
        1,
        2,
    ):
        columna = (
            columna_inicio
            +
            desplazamiento
        )

        if columna >= len(fila):
            continue

        valor = fila[
            columna
        ]

        if normalizar_texto(valor):
            return str(
                valor
            ).strip()

    return ""


def detectar_estructura_creditos(filas):
    encabezado = encontrar_fila_encabezados_creditos(
        filas
    )

    if encabezado is None:
        return None

    fila_encabezados = encabezado[
        "fila"
    ]

    fila_bancos = (
        fila_encabezados
        -
        1
    )

    bloques = []

    for columna_inicio in encabezado[
        "bloques"
    ]:
        bloques.append({
            "banco_original":
                obtener_nombre_banco_creditos(
                    filas,
                    fila_bancos,
                    columna_inicio,
                ),

            "columna_mn":
                columna_inicio,

            "columna_me":
                columna_inicio + 1,

            "columna_total":
                columna_inicio + 2,
        })

    return {
        "fila_bancos":
            fila_bancos,

        "fila_encabezados":
            fila_encabezados,

        "bloques":
            bloques,
    }


def elegir_hoja_principal_creditos(hojas):
    candidatos = []

    for hoja in hojas:
        estructura = detectar_estructura_creditos(
            hoja["filas"]
        )

        if estructura is not None:
            candidatos.append({
                "hoja":
                    hoja,

                "estructura":
                    estructura,
            })

    if not candidatos:
        return None

    return max(
        candidatos,
        key=lambda x: len(
            x["estructura"]["bloques"]
        ),
    )


def buscar_fila_total_creditos(filas):
    for numero_fila, fila in enumerate(
        filas
    ):
        for valor in fila:
            if (
                "total creditos"
                in
                sin_tildes(valor)
            ):
                return numero_fila

    return None


def es_agregado_creditos(nombre):
    return (
        "total banca multiple"
        in
        sin_tildes(nombre)
    )


def extraer_y_creditos(ruta):
    formato, hojas = leer_archivo(
        ruta
    )

    seleccionado = elegir_hoja_principal_creditos(
        hojas
    )

    if seleccionado is None:
        raise ErrorDiagnosticoY(
            "B-2359: no se encontró "
            "la estructura MN|ME|Total."
        )

    hoja = seleccionado[
        "hoja"
    ]

    estructura = seleccionado[
        "estructura"
    ]

    filas = hoja[
        "filas"
    ]

    fila_total = buscar_fila_total_creditos(
        filas
    )

    if fila_total is None:
        raise ErrorDiagnosticoY(
            "B-2359: no se encontró "
            "la fila 'Total Créditos:'."
        )

    fila = filas[
        fila_total
    ]

    registros = []

    for bloque in estructura[
        "bloques"
    ]:
        banco = bloque[
            "banco_original"
        ]

        col_mn = bloque[
            "columna_mn"
        ]

        col_me = bloque[
            "columna_me"
        ]

        col_total = bloque[
            "columna_total"
        ]

        mn = (
            convertir_numero(
                fila[col_mn]
            )
            if col_mn < len(fila)
            else None
        )

        me = (
            convertir_numero(
                fila[col_me]
            )
            if col_me < len(fila)
            else None
        )

        total = (
            convertir_numero(
                fila[col_total]
            )
            if col_total < len(fila)
            else None
        )

        problemas = []

        if not banco:
            problemas.append(
                "BANCO_SIN_NOMBRE"
            )

        if mn is None:
            problemas.append(
                "MN_FALTANTE"
            )

        if me is None:
            problemas.append(
                "ME_FALTANTE"
            )

        if total is None:
            problemas.append(
                "TOTAL_FALTANTE"
            )

        if (
            mn is not None
            and
            mn < 0
        ):
            problemas.append(
                "MN_NEGATIVO"
            )

        if (
            me is not None
            and
            me < 0
        ):
            problemas.append(
                "ME_NEGATIVO"
            )

        if (
            total is not None
            and
            total <= 0
        ):
            problemas.append(
                "TOTAL_NO_POSITIVO"
            )

        dolcred = None
        tc_implicito = None

        if (
            mn is not None
            and
            total is not None
            and
            total < mn
        ):
            problemas.append(
                "TOTAL_MENOR_QUE_MN"
            )

        # Fórmula validada. NO se modifica.
        if (
            total is not None
            and
            total > 0
            and
            mn is not None
        ):
            dolcred = (
                (
                    total
                    -
                    mn
                )
                /
                total
                *
                100
            )

            if not (
                0
                <=
                dolcred
                <=
                100
            ):
                problemas.append(
                    "DOLCRED_FUERA_RANGO"
                )

        if (
            mn is not None
            and
            me is not None
            and
            total is not None
            and
            me > 0
        ):
            tc_implicito = (
                (
                    total
                    -
                    mn
                )
                /
                me
            )

            if tc_implicito <= 0:
                problemas.append(
                    "TC_NO_POSITIVO"
                )

        registros.append({
            "banco_original":
                banco,

            "mn_soles_miles":
                mn,

            "me_usd_miles":
                me,

            "total_soles_miles":
                total,

            "dolcred_pct":
                dolcred,

            "tc_implicito":
                tc_implicito,

            "es_agregado":
                es_agregado_creditos(
                    banco
                ),

            "problemas":
                problemas,
        })

    return {
        "formato":
            formato,

        "hoja":
            hoja["nombre"],

        "fila_encabezados":
            estructura[
                "fila_encabezados"
            ],

        "fila_bancos":
            estructura[
                "fila_bancos"
            ],

        "fila_total_creditos":
            fila_total,

        "cantidad_bloques":
            len(
                estructura["bloques"]
            ),

        "registros":
            registros,
    }


# ============================================================
# 11. PRECISIÓN VISIBLE DEL AGREGADO
# MISMA LÓGICA DE PRODUCCIÓN
# ============================================================


def inferir_decimales_formato_excel(
    formato
):
    if formato is None:
        return None

    texto = str(
        formato
    ).strip()

    if (
        not texto
        or
        texto.lower() == "general"
    ):
        return None

    seccion = texto.split(
        ";"
    )[0]

    seccion = re.sub(
        r'"[^\"]*"',
        "",
        seccion,
    )

    seccion = re.sub(
        r"\[[^\]]*\]",
        "",
        seccion,
    )

    coincidencia = re.search(
        r"\.([0#?]+)",
        seccion,
    )

    if coincidencia:
        return len(
            coincidencia.group(1)
        )

    if re.search(
        r"[0#?]",
        seccion,
    ):
        return 0

    return None


def redondear_excel_half_up(
    valor,
    decimales,
):
    if (
        valor is None
        or
        decimales is None
        or
        decimales < 0
    ):
        raise ErrorDiagnosticoY(
            "No puede aplicarse ROUND_HALF_UP "
            "sin valor y decimales válidos."
        )

    numero = Decimal(
        str(valor)
    )

    cuantizador = (
        Decimal("1")
        if decimales == 0
        else
        Decimal("1").scaleb(
            -decimales
        )
    )

    return numero.quantize(
        cuantizador,
        rounding=ROUND_HALF_UP,
    )


# ============================================================
# 12. RECUPERAR ESTRUCTURA SELECCIONADA
# SOLO PARA INSPECCIÓN. NO CAMBIA EL PARSER.
# ============================================================


def recuperar_estructura_seleccionada(
    ruta,
    resultado,
):
    _, hojas = leer_archivo(
        ruta
    )

    seleccionado = elegir_hoja_principal_creditos(
        hojas
    )

    if seleccionado is None:
        raise ErrorDiagnosticoY(
            "No pudo recuperarse "
            "la estructura seleccionada."
        )

    if (
        seleccionado["hoja"]["nombre"]
        !=
        resultado["hoja"]
    ):
        raise ErrorDiagnosticoY(
            "La hoja seleccionada en la inspección "
            "no coincide con la del parser."
        )

    return (
        seleccionado["hoja"],
        seleccionado["estructura"],
    )


# ============================================================
# 13. LEER METADATOS RAW DE TODAS LAS CELDAS DE LOS BLOQUES
# ============================================================


def leer_celdas_bloques_xls(
    ruta,
    resultado,
    estructura,
):
    libro = xlrd.open_workbook(
        filename=str(ruta),
        formatting_info=True,
        on_demand=True,
    )

    try:
        hoja = libro.sheet_by_name(
            resultado["hoja"]
        )

        fila_total = resultado[
            "fila_total_creditos"
        ]

        fila_encabezados = resultado[
            "fila_encabezados"
        ]

        fila_bancos = resultado[
            "fila_bancos"
        ]

        salida = []

        for indice_bloque, bloque in enumerate(
            estructura["bloques"],
            start=1,
        ):
            item = {
                "indice_bloque":
                    indice_bloque,

                "banco_original":
                    bloque[
                        "banco_original"
                    ],

                "es_agregado":
                    es_agregado_creditos(
                        bloque[
                            "banco_original"
                        ]
                    ),

                "columna_mn":
                    bloque[
                        "columna_mn"
                    ],

                "columna_me":
                    bloque[
                        "columna_me"
                    ],

                "columna_total":
                    bloque[
                        "columna_total"
                    ],

                "encabezado_mn_raw":
                    hoja.cell_value(
                        fila_encabezados,
                        bloque["columna_mn"],
                    ),

                "encabezado_me_raw":
                    hoja.cell_value(
                        fila_encabezados,
                        bloque["columna_me"],
                    ),

                "encabezado_total_raw":
                    hoja.cell_value(
                        fila_encabezados,
                        bloque["columna_total"],
                    ),

                "nombre_celda_1_raw":
                    hoja.cell_value(
                        fila_bancos,
                        bloque["columna_mn"],
                    ),

                "nombre_celda_2_raw":
                    hoja.cell_value(
                        fila_bancos,
                        bloque["columna_me"],
                    ),

                "nombre_celda_3_raw":
                    hoja.cell_value(
                        fila_bancos,
                        bloque["columna_total"],
                    ),
            }

            for variable, columna in (
                (
                    "MN",
                    bloque["columna_mn"],
                ),
                (
                    "ME",
                    bloque["columna_me"],
                ),
                (
                    "TOTAL",
                    bloque["columna_total"],
                ),
            ):
                celda = hoja.cell(
                    fila_total,
                    columna,
                )

                xf_index = hoja.cell_xf_index(
                    fila_total,
                    columna,
                )

                format_key = None
                format_str = None

                if (
                    0 <= xf_index
                    <
                    len(libro.xf_list)
                ):
                    xf = libro.xf_list[
                        xf_index
                    ]

                    format_key = getattr(
                        xf,
                        "format_key",
                        None,
                    )

                    if format_key is not None:
                        formato_obj = (
                            libro.format_map.get(
                                format_key
                            )
                        )

                        if formato_obj is not None:
                            format_str = getattr(
                                formato_obj,
                                "format_str",
                                None,
                            )

                item[
                    f"{variable.lower()}_raw"
                ] = celda.value

                item[
                    f"{variable.lower()}_raw_repr"
                ] = repr(
                    celda.value
                )

                item[
                    f"{variable.lower()}_ctype"
                ] = celda.ctype

                item[
                    f"{variable.lower()}_xf_index"
                ] = xf_index

                item[
                    f"{variable.lower()}_format_key"
                ] = format_key

                item[
                    f"{variable.lower()}_format_str"
                ] = format_str

                item[
                    f"{variable.lower()}_decimales"
                ] = inferir_decimales_formato_excel(
                    format_str
                )

            salida.append(
                item
            )

        return salida

    finally:
        libro.release_resources()


def leer_celdas_bloques_xlsx(
    ruta,
    resultado,
    estructura,
):
    with Path(ruta).open(
        "rb"
    ) as archivo_binario:
        libro = load_workbook(
            archivo_binario,
            read_only=False,
            data_only=True,
        )

        try:
            hoja = libro[
                resultado["hoja"]
            ]

            fila_total = (
                resultado[
                    "fila_total_creditos"
                ]
                +
                1
            )

            fila_encabezados = (
                resultado[
                    "fila_encabezados"
                ]
                +
                1
            )

            fila_bancos = (
                resultado[
                    "fila_bancos"
                ]
                +
                1
            )

            salida = []

            for indice_bloque, bloque in enumerate(
                estructura["bloques"],
                start=1,
            ):
                item = {
                    "indice_bloque":
                        indice_bloque,

                    "banco_original":
                        bloque[
                            "banco_original"
                        ],

                    "es_agregado":
                        es_agregado_creditos(
                            bloque[
                                "banco_original"
                            ]
                        ),

                    "columna_mn":
                        bloque[
                            "columna_mn"
                        ],

                    "columna_me":
                        bloque[
                            "columna_me"
                        ],

                    "columna_total":
                        bloque[
                            "columna_total"
                        ],

                    "encabezado_mn_raw":
                        hoja.cell(
                            fila_encabezados,
                            bloque[
                                "columna_mn"
                            ]
                            +
                            1,
                        ).value,

                    "encabezado_me_raw":
                        hoja.cell(
                            fila_encabezados,
                            bloque[
                                "columna_me"
                            ]
                            +
                            1,
                        ).value,

                    "encabezado_total_raw":
                        hoja.cell(
                            fila_encabezados,
                            bloque[
                                "columna_total"
                            ]
                            +
                            1,
                        ).value,

                    "nombre_celda_1_raw":
                        hoja.cell(
                            fila_bancos,
                            bloque[
                                "columna_mn"
                            ]
                            +
                            1,
                        ).value,

                    "nombre_celda_2_raw":
                        hoja.cell(
                            fila_bancos,
                            bloque[
                                "columna_me"
                            ]
                            +
                            1,
                        ).value,

                    "nombre_celda_3_raw":
                        hoja.cell(
                            fila_bancos,
                            bloque[
                                "columna_total"
                            ]
                            +
                            1,
                        ).value,
                }

                for variable, columna_0based in (
                    (
                        "MN",
                        bloque[
                            "columna_mn"
                        ],
                    ),
                    (
                        "ME",
                        bloque[
                            "columna_me"
                        ],
                    ),
                    (
                        "TOTAL",
                        bloque[
                            "columna_total"
                        ],
                    ),
                ):
                    celda = hoja.cell(
                        row=fila_total,
                        column=columna_0based + 1,
                    )

                    format_str = (
                        celda.number_format
                    )

                    item[
                        f"{variable.lower()}_raw"
                    ] = celda.value

                    item[
                        f"{variable.lower()}_raw_repr"
                    ] = repr(
                        celda.value
                    )

                    item[
                        f"{variable.lower()}_ctype"
                    ] = str(
                        celda.data_type
                    )

                    item[
                        f"{variable.lower()}_xf_index"
                    ] = celda.style_id

                    item[
                        f"{variable.lower()}_format_key"
                    ] = ""

                    item[
                        f"{variable.lower()}_format_str"
                    ] = format_str

                    item[
                        f"{variable.lower()}_decimales"
                    ] = inferir_decimales_formato_excel(
                        format_str
                    )

                salida.append(
                    item
                )

            return salida

        finally:
            libro.close()


def leer_celdas_bloques(
    ruta,
    resultado,
    estructura,
):
    if (
        resultado["formato"]
        ==
        "XLS_OLE"
    ):
        return leer_celdas_bloques_xls(
            ruta,
            resultado,
            estructura,
        )

    if (
        resultado["formato"]
        ==
        "XLSX_ZIP"
    ):
        return leer_celdas_bloques_xlsx(
            ruta,
            resultado,
            estructura,
        )

    raise ErrorDiagnosticoY(
        "Formato no compatible "
        "con inspección de celdas."
    )


# ============================================================
# 14. DIAGNÓSTICO DE CONTROLES MN / ME / TOTAL
# ============================================================


def diagnosticar_controles(
    resultado,
    bloques_raw,
):
    registros = resultado[
        "registros"
    ]

    agregados = [
        r
        for r in registros
        if r["es_agregado"]
    ]

    bancos = [
        r
        for r in registros
        if not r["es_agregado"]
    ]

    if len(agregados) != 1:
        raise ErrorDiagnosticoY(
            "Se esperaba exactamente un "
            "Total Banca Múltiple. "
            f"Encontrados={len(agregados)}"
        )

    agregado = agregados[
        0
    ]

    raw_agregados = [
        b
        for b in bloques_raw
        if b["es_agregado"]
    ]

    if len(raw_agregados) != 1:
        raise ErrorDiagnosticoY(
            "La inspección raw no encontró "
            "exactamente un bloque agregado."
        )

    raw_agregado = raw_agregados[
        0
    ]

    controles = []

    for variable, campo in (
        (
            "MN",
            "mn_soles_miles",
        ),
        (
            "ME",
            "me_usd_miles",
        ),
        (
            "TOTAL",
            "total_soles_miles",
        ),
    ):
        valores = [
            r[campo]
            for r in bancos
        ]

        faltantes = [
            r["banco_original"]
            for r in bancos
            if r[campo] is None
        ]

        if faltantes:
            raise ErrorDiagnosticoY(
                f"{variable}: existen "
                "valores faltantes en bancos: "
                f"{faltantes}"
            )

        suma_sum = sum(
            valores
        )

        suma_fsum = math.fsum(
            valores
        )

        total_sbs = agregado[
            campo
        ]

        if total_sbs is None:
            raise ErrorDiagnosticoY(
                f"{variable}: "
                "Total Banca Múltiple "
                "no tiene valor numérico."
            )

        diferencia_sum = (
            suma_sum
            -
            total_sbs
        )

        diferencia_fsum = (
            suma_fsum
            -
            total_sbs
        )

        diferencia_abs_sum = abs(
            diferencia_sum
        )

        diferencia_abs_fsum = abs(
            diferencia_fsum
        )

        if total_sbs != 0:
            diferencia_rel_sum = (
                diferencia_abs_sum
                /
                abs(total_sbs)
            )

            diferencia_rel_fsum = (
                diferencia_abs_fsum
                /
                abs(total_sbs)
            )

        else:
            diferencia_rel_sum = None
            diferencia_rel_fsum = None

        prefijo = variable.lower()

        format_str = raw_agregado[
            f"{prefijo}_format_str"
        ]

        decimales = raw_agregado[
            f"{prefijo}_decimales"
        ]

        raw_celda = raw_agregado[
            f"{prefijo}_raw"
        ]

        if (
            format_str is None
            or
            decimales is None
        ):
            suma_redondeada = None
            total_redondeado = None
            coincide = None

        else:
            suma_redondeada = redondear_excel_half_up(
                suma_fsum,
                decimales,
            )

            total_redondeado = redondear_excel_half_up(
                total_sbs,
                decimales,
            )

            coincide = (
                suma_redondeada
                ==
                total_redondeado
            )

        controles.append({
            "variable":
                variable,

            "cantidad_bancos":
                len(bancos),

            "suma_sum":
                suma_sum,

            "suma_fsum":
                suma_fsum,

            "total_sbs_raw_parser":
                total_sbs,

            "total_sbs_raw_celda":
                raw_celda,

            "total_sbs_raw_celda_repr":
                repr(raw_celda),

            "diferencia_sum":
                diferencia_sum,

            "diferencia_absoluta_sum":
                diferencia_abs_sum,

            "diferencia_relativa_sum":
                diferencia_rel_sum,

            "diferencia_fsum":
                diferencia_fsum,

            "diferencia_absoluta_fsum":
                diferencia_abs_fsum,

            "diferencia_relativa_fsum":
                diferencia_rel_fsum,

            "diferencia_sum_vs_fsum":
                (
                    suma_sum
                    -
                    suma_fsum
                ),

            "formato_excel_agregado":
                format_str,

            "decimales_mostrados":
                decimales,

            "suma_redondeada_half_up":
                (
                    ""
                    if suma_redondeada is None
                    else
                    str(suma_redondeada)
                ),

            "total_redondeado_half_up":
                (
                    ""
                    if total_redondeado is None
                    else
                    str(total_redondeado)
                ),

            "coincide_precision_mostrada":
                coincide,
        })

    return (
        bancos,
        agregado,
        controles,
    )


# ============================================================
# 15. AUDITORÍA ESTRUCTURAL DEL MES
# ============================================================


def auditar_estructura(
    resultado,
    estructura,
    bloques_raw,
    hoja_normal,
):
    alertas = []

    # --------------------------------------------------------
    # A. Correspondencia 1 bloque = 1 registro del parser
    # --------------------------------------------------------

    if (
        len(estructura["bloques"])
        !=
        len(resultado["registros"])
    ):
        alertas.append(
            "CANTIDAD_BLOQUES_DISTINTA_"
            "DE_REGISTROS_PARSER"
        )

    # --------------------------------------------------------
    # B. Nombres vacíos
    # --------------------------------------------------------

    vacios = [
        bloque["indice_bloque"]
        for bloque in bloques_raw
        if not normalizar_texto(
            bloque["banco_original"]
        )
    ]

    if vacios:
        alertas.append(
            f"NOMBRES_VACIOS_EN_BLOQUES={vacios}"
        )

    # --------------------------------------------------------
    # C. Duplicados exactos / normalizados
    # Solo diagnóstico; no se homologa nada.
    # --------------------------------------------------------

    nombres_exactos = [
        b["banco_original"]
        for b in bloques_raw
    ]

    duplicados_exactos = sorted({
        nombre
        for nombre in nombres_exactos
        if (
            nombres_exactos.count(nombre)
            >
            1
        )
    })

    if duplicados_exactos:
        alertas.append(
            "NOMBRES_DUPLICADOS_EXACTOS="
            f"{duplicados_exactos}"
        )

    nombres_norm = [
        normalizar_comparacion(
            nombre
        )
        for nombre in nombres_exactos
    ]

    duplicados_norm = sorted({
        nombre
        for nombre in nombres_norm
        if (
            nombre
            and
            nombres_norm.count(nombre) > 1
        )
    })

    if duplicados_norm:
        alertas.append(
            "NOMBRES_DUPLICADOS_NORMALIZADOS_"
            "SOLO_DIAGNOSTICO="
            f"{duplicados_norm}"
        )

    # --------------------------------------------------------
    # D. Exactamente un agregado
    # --------------------------------------------------------

    agregados = [
        b
        for b in bloques_raw
        if b["es_agregado"]
    ]

    if len(agregados) != 1:
        alertas.append(
            f"CANTIDAD_AGREGADOS={len(agregados)}"
        )

    # --------------------------------------------------------
    # E. Tripletas consecutivas y sin solapamiento
    # --------------------------------------------------------

    todas_columnas = []

    for bloque in estructura[
        "bloques"
    ]:
        cols = (
            bloque["columna_mn"],
            bloque["columna_me"],
            bloque["columna_total"],
        )

        if not (
            cols[1]
            ==
            cols[0] + 1
            and
            cols[2]
            ==
            cols[0] + 2
        ):
            alertas.append(
                "TRIPLETA_NO_CONSECUTIVA="
                f"{bloque['banco_original']}:"
                f"{cols}"
            )

        todas_columnas.extend(
            cols
        )

    columnas_repetidas = sorted({
        c
        for c in todas_columnas
        if (
            todas_columnas.count(c)
            >
            1
        )
    })

    if columnas_repetidas:
        alertas.append(
            "COLUMNAS_SOLAPADAS="
            f"{columnas_repetidas}"
        )

    # --------------------------------------------------------
    # F. Encabezados raw deben seguir MN | ME | Total
    # --------------------------------------------------------

    for bloque in bloques_raw:
        if not es_mn(
            bloque["encabezado_mn_raw"]
        ):
            alertas.append(
                "ENCABEZADO_MN_DISTINTO="
                f"{bloque['banco_original']!r}:"
                f"{bloque['encabezado_mn_raw']!r}"
            )

        if not es_me(
            bloque["encabezado_me_raw"]
        ):
            alertas.append(
                "ENCABEZADO_ME_DISTINTO="
                f"{bloque['banco_original']!r}:"
                f"{bloque['encabezado_me_raw']!r}"
            )

        if not es_total(
            bloque[
                "encabezado_total_raw"
            ]
        ):
            alertas.append(
                "ENCABEZADO_TOTAL_DISTINTO="
                f"{bloque['banco_original']!r}:"
                f"{bloque['encabezado_total_raw']!r}"
            )

    # --------------------------------------------------------
    # G. Valores raw faltantes/no numéricos
    # --------------------------------------------------------

    for bloque in bloques_raw:
        for variable in (
            "mn",
            "me",
            "total",
        ):
            raw = bloque[
                f"{variable}_raw"
            ]

            convertido = convertir_numero(
                raw
            )

            if convertido is None:
                alertas.append(
                    f"CELDA_{variable.upper()}_"
                    "NO_NUMERICA_O_FALTANTE="
                    f"{bloque['banco_original']!r}:"
                    f"{raw!r}"
                )

    # --------------------------------------------------------
    # H. Total <= 0 por banco
    # --------------------------------------------------------

    for registro in resultado[
        "registros"
    ]:
        if registro[
            "es_agregado"
        ]:
            continue

        total = registro[
            "total_soles_miles"
        ]

        if (
            total is not None
            and
            total <= 0
        ):
            alertas.append(
                "BANCO_TOTAL_NO_POSITIVO="
                f"{registro['banco_original']!r}:"
                f"{total!r}"
            )

    # --------------------------------------------------------
    # I. Celdas numéricas en la fila Total Créditos fuera
    #    de columnas MN/ME/Total detectadas.
    # --------------------------------------------------------

    fila_total = (
        hoja_normal["filas"][
            resultado[
                "fila_total_creditos"
            ]
        ]
    )

    columnas_detectadas = set(
        todas_columnas
    )

    numericas_fuera = []

    for indice_columna, raw in enumerate(
        fila_total
    ):
        if indice_columna in columnas_detectadas:
            continue

        numero = convertir_numero(
            raw
        )

        if numero is not None:
            numericas_fuera.append(
                (
                    indice_columna,
                    raw,
                )
            )

    # No se declara automáticamente error:
    # algunas columnas de concepto pueden contener
    # códigos u otros valores. Se registra para revisar.
    if numericas_fuera:
        alertas.append(
            "CELDAS_NUMERICAS_FUERA_DE_BLOQUES="
            f"{numericas_fuera}"
        )

    # --------------------------------------------------------
    # J. Bancos incluidos/excluidos según la regla actual
    # --------------------------------------------------------

    incluidos = [
        r["banco_original"]
        for r in resultado["registros"]
        if not r["es_agregado"]
    ]

    excluidos = [
        r["banco_original"]
        for r in resultado["registros"]
        if r["es_agregado"]
    ]

    return {
        "alertas":
            alertas,

        "incluidos":
            incluidos,

        "excluidos":
            excluidos,

        "numericas_fuera":
            numericas_fuera,

        "duplicados_exactos":
            duplicados_exactos,

        "duplicados_normalizados":
            duplicados_norm,
    }


# ============================================================
# 16. MOSTRAR CONTROLES
# ============================================================


def mostrar_controles(controles):
    titulo(
        "CONTROLES TOTAL BANCA MÚLTIPLE - 2017-07"
    )

    for control in controles:
        subtitulo(
            control["variable"]
        )

        escribir(
            "Cantidad de bancos individuales = "
            f"{control['cantidad_bancos']}"
        )

        escribir(
            f"sum() = "
            f"{control['suma_sum']!r}"
        )

        escribir(
            f"math.fsum() = "
            f"{control['suma_fsum']!r}"
        )

        escribir(
            "Total Banca Múltiple SBS raw "
            "(parser) = "
            f"{control['total_sbs_raw_parser']!r}"
        )

        escribir(
            "Total Banca Múltiple SBS raw "
            "(celda) = "
            f"{control['total_sbs_raw_celda_repr']}"
        )

        escribir(
            "Diferencia absoluta sum() = "
            f"{control['diferencia_absoluta_sum']!r}"
        )

        escribir(
            "Diferencia relativa sum() = "
            f"{control['diferencia_relativa_sum']!r}"
        )

        escribir(
            "Diferencia absoluta math.fsum() = "
            f"{control['diferencia_absoluta_fsum']!r}"
        )

        escribir(
            "Diferencia relativa math.fsum() = "
            f"{control['diferencia_relativa_fsum']!r}"
        )

        escribir(
            "sum() - math.fsum() = "
            f"{control['diferencia_sum_vs_fsum']!r}"
        )

        escribir(
            "Formato Excel agregado = "
            f"{control['formato_excel_agregado']!r}"
        )

        escribir(
            "Decimales mostrados = "
            f"{control['decimales_mostrados']}"
        )

        escribir(
            "Suma ROUND_HALF_UP = "
            f"{control['suma_redondeada_half_up']}"
        )

        escribir(
            "Total SBS ROUND_HALF_UP = "
            f"{control['total_redondeado_half_up']}"
        )

        escribir(
            "¿Coinciden a precisión mostrada? = "
            f"{control['coincide_precision_mostrada']}"
        )


# ============================================================
# 17. MOSTRAR TODOS LOS VALORES INDIVIDUALES MN
# ============================================================


def mostrar_valores_mn(
    bloques_raw
):
    titulo(
        "VALORES INDIVIDUALES MN - "
        "TODOS LOS BLOQUES DETECTADOS"
    )

    for bloque in bloques_raw:
        clasificacion = (
            "AGREGADO_EXCLUIDO_DE_LA_SUMA"
            if bloque["es_agregado"]
            else
            "BANCO_INCLUIDO_EN_LA_SUMA"
        )

        escribir()

        escribir(
            f"Bloque "
            f"{bloque['indice_bloque']:02d} | "
            f"{bloque['banco_original']!r}"
        )

        escribir(
            f"  Clasificación = "
            f"{clasificacion}"
        )

        escribir(
            "  Columnas MN/ME/TOTAL = "
            f"{bloque['columna_mn']}/"
            f"{bloque['columna_me']}/"
            f"{bloque['columna_total']}"
        )

        escribir(
            "  Encabezado MN raw = "
            f"{bloque['encabezado_mn_raw']!r}"
        )

        escribir(
            "  Nombre raw en las 3 columnas = "
            f"{bloque['nombre_celda_1_raw']!r} | "
            f"{bloque['nombre_celda_2_raw']!r} | "
            f"{bloque['nombre_celda_3_raw']!r}"
        )

        escribir(
            f"  MN raw = "
            f"{bloque['mn_raw_repr']}"
        )

        escribir(
            "  MN convertido = "
            f"{convertir_numero(bloque['mn_raw'])!r}"
        )

        escribir(
            f"  MN ctype = "
            f"{bloque['mn_ctype']!r}"
        )

        escribir(
            f"  MN formato = "
            f"{bloque['mn_format_str']!r}"
        )

        escribir(
            "  MN decimales mostrados = "
            f"{bloque['mn_decimales']}"
        )


# ============================================================
# 18. MOSTRAR AUDITORÍA ESTRUCTURAL
# ============================================================


def mostrar_auditoria_estructural(
    resultado,
    estructura,
    bloques_raw,
    auditoria,
):
    titulo(
        "AUDITORÍA ESTRUCTURAL B-2359 / 2017-07"
    )

    escribir(
        f"Hoja seleccionada = "
        f"{resultado['hoja']!r}"
    )

    escribir(
        f"Fila bancos = "
        f"{resultado['fila_bancos']}"
    )

    escribir(
        f"Fila encabezados = "
        f"{resultado['fila_encabezados']}"
    )

    escribir(
        f"Fila Total Créditos = "
        f"{resultado['fila_total_creditos']}"
    )

    escribir(
        f"Bloques detectados = "
        f"{len(estructura['bloques'])}"
    )

    escribir(
        "Registros producidos por parser = "
        f"{len(resultado['registros'])}"
    )

    escribir()
    escribir(
        "Bancos incluidos en las sumas:"
    )

    for nombre in auditoria[
        "incluidos"
    ]:
        escribir(
            f"  + {nombre!r}"
        )

    escribir()
    escribir(
        "Bloques excluidos por ser agregado:"
    )

    for nombre in auditoria[
        "excluidos"
    ]:
        escribir(
            f"  - {nombre!r}"
        )

    escribir()

    escribir(
        "Duplicados exactos = "
        f"{auditoria['duplicados_exactos']}"
    )

    escribir(
        "Duplicados normalizados "
        "(solo diagnóstico) = "
        f"{auditoria['duplicados_normalizados']}"
    )

    escribir()
    escribir(
        "Tripletas detectadas:"
    )

    for bloque in bloques_raw:
        escribir(
            f"  {bloque['indice_bloque']:02d} | "
            f"{bloque['banco_original']!r} | "
            f"cols "
            f"{bloque['columna_mn']},"
            f"{bloque['columna_me']},"
            f"{bloque['columna_total']} | "
            f"headers "
            f"{bloque['encabezado_mn_raw']!r} / "
            f"{bloque['encabezado_me_raw']!r} / "
            f"{bloque['encabezado_total_raw']!r} | "
            f"raw MN={bloque['mn_raw_repr']} | "
            f"ME={bloque['me_raw_repr']} | "
            f"TOTAL={bloque['total_raw_repr']}"
        )

    escribir()

    escribir(
        "Celdas numéricas fuera de las columnas "
        "MN/ME/TOTAL detectadas en la fila "
        "Total Créditos = "
        f"{auditoria['numericas_fuera']}"
    )

    escribir()

    if auditoria[
        "alertas"
    ]:
        escribir(
            "ALERTAS ESTRUCTURALES/DIAGNÓSTICAS:"
        )

        for alerta in auditoria[
            "alertas"
        ]:
            escribir(
                f"  * {alerta}"
            )

    else:
        escribir(
            "ALERTAS ESTRUCTURALES/DIAGNÓSTICAS: "
            "ninguna"
        )


# ============================================================
# 19. GUARDAR CSV DE CONTROLES
# ============================================================


def guardar_csv_controles(
    controles
):
    columnas = [
        "variable",
        "cantidad_bancos",
        "suma_sum",
        "suma_fsum",
        "total_sbs_raw_parser",
        "total_sbs_raw_celda",
        "total_sbs_raw_celda_repr",
        "diferencia_sum",
        "diferencia_absoluta_sum",
        "diferencia_relativa_sum",
        "diferencia_fsum",
        "diferencia_absoluta_fsum",
        "diferencia_relativa_fsum",
        "diferencia_sum_vs_fsum",
        "formato_excel_agregado",
        "decimales_mostrados",
        "suma_redondeada_half_up",
        "total_redondeado_half_up",
        "coincide_precision_mostrada",
    ]

    with CSV_CONTROLES.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as archivo:
        escritor = csv.DictWriter(
            archivo,
            fieldnames=columnas,
        )

        escritor.writeheader()

        for fila in controles:
            escritor.writerow(
                fila
            )


# ============================================================
# 20. GUARDAR CSV DE BLOQUES / CELDAS RAW
# ============================================================


def guardar_csv_bloques(
    bloques_raw
):
    columnas = [
        "indice_bloque",
        "banco_original",
        "es_agregado",
        "columna_mn",
        "columna_me",
        "columna_total",
        "encabezado_mn_raw",
        "encabezado_me_raw",
        "encabezado_total_raw",
        "nombre_celda_1_raw",
        "nombre_celda_2_raw",
        "nombre_celda_3_raw",
        "mn_raw",
        "mn_raw_repr",
        "mn_ctype",
        "mn_xf_index",
        "mn_format_key",
        "mn_format_str",
        "mn_decimales",
        "me_raw",
        "me_raw_repr",
        "me_ctype",
        "me_xf_index",
        "me_format_key",
        "me_format_str",
        "me_decimales",
        "total_raw",
        "total_raw_repr",
        "total_ctype",
        "total_xf_index",
        "total_format_key",
        "total_format_str",
        "total_decimales",
    ]

    with CSV_BLOQUES.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as archivo:
        escritor = csv.DictWriter(
            archivo,
            fieldnames=columnas,
        )

        escritor.writeheader()

        for fila in bloques_raw:
            escritor.writerow(
                fila
            )


# ============================================================
# 21. MAIN
# ============================================================


def main():
    try:
        titulo(
            "DIAGNÓSTICO LOCAL Y / B-2359 / 2017-07"
        )

        escribir(
            "No se realizará ninguna descarga."
        )

        escribir(
            "No se modifica el script de producción."
        )

        escribir(
            "No se modifica DOLCRED."
        )

        escribir(
            "No se modifica el parser."
        )

        escribir(
            "No se ejecuta X1."
        )

        escribir()

        escribir(
            f"Archivo esperado = "
            f"{ARCHIVO_LOCAL}"
        )

        if not ARCHIVO_LOCAL.exists():
            raise ErrorDiagnosticoY(
                "No existe el archivo crudo local indicado: "
                f"{ARCHIVO_LOCAL}"
            )

        if (
            ARCHIVO_LOCAL.stat().st_size
            <=
            0
        ):
            raise ErrorDiagnosticoY(
                "El archivo crudo local existe "
                "pero tiene 0 bytes."
            )

        escribir(
            f"Tamaño bytes = "
            f"{ARCHIVO_LOCAL.stat().st_size}"
        )

        escribir(
            f"SHA-256 = "
            f"{sha256_archivo(ARCHIVO_LOCAL)}"
        )

        formato = detectar_formato(
            ARCHIVO_LOCAL
        )

        escribir(
            f"Formato real = "
            f"{formato}"
        )

        # ----------------------------------------------------
        # Reproducir exactamente el parser Y de producción.
        # ----------------------------------------------------

        resultado = extraer_y_creditos(
            ARCHIVO_LOCAL
        )

        hoja_normal, estructura = (
            recuperar_estructura_seleccionada(
                ARCHIVO_LOCAL,
                resultado,
            )
        )

        bloques_raw = leer_celdas_bloques(
            ARCHIVO_LOCAL,
            resultado,
            estructura,
        )

        bancos, agregado, controles = (
            diagnosticar_controles(
                resultado,
                bloques_raw,
            )
        )

        auditoria = auditar_estructura(
            resultado,
            estructura,
            bloques_raw,
            hoja_normal,
        )

        # ----------------------------------------------------
        # Mostrar resultados.
        # ----------------------------------------------------

        mostrar_controles(
            controles
        )

        mostrar_valores_mn(
            bloques_raw
        )

        mostrar_auditoria_estructural(
            resultado,
            estructura,
            bloques_raw,
            auditoria,
        )

        # ----------------------------------------------------
        # Controles de correspondencia parser/raw.
        # ----------------------------------------------------

        titulo(
            "CONTROL DE CORRESPONDENCIA "
            "PARSER VS CELDAS RAW"
        )

        if (
            len(resultado["registros"])
            !=
            len(bloques_raw)
        ):
            escribir(
                "ERROR: cantidad de registros "
                "del parser distinta de bloques raw."
            )

        else:
            escribir(
                "Cantidad de registros parser "
                "y bloques raw coincide: "
                f"{len(bloques_raw)}"
            )

        discrepancias = []

        for registro, raw in zip(
            resultado["registros"],
            bloques_raw,
        ):
            controles_valor = (
                (
                    "MN",
                    registro[
                        "mn_soles_miles"
                    ],
                    convertir_numero(
                        raw["mn_raw"]
                    ),
                ),
                (
                    "ME",
                    registro[
                        "me_usd_miles"
                    ],
                    convertir_numero(
                        raw["me_raw"]
                    ),
                ),
                (
                    "TOTAL",
                    registro[
                        "total_soles_miles"
                    ],
                    convertir_numero(
                        raw["total_raw"]
                    ),
                ),
            )

            if (
                registro[
                    "banco_original"
                ]
                !=
                raw[
                    "banco_original"
                ]
            ):
                discrepancias.append(
                    (
                        raw[
                            "indice_bloque"
                        ],
                        "NOMBRE",
                        registro[
                            "banco_original"
                        ],
                        raw[
                            "banco_original"
                        ],
                    )
                )

            for (
                variable,
                valor_parser,
                valor_raw_convertido,
            ) in controles_valor:
                if (
                    valor_parser
                    !=
                    valor_raw_convertido
                ):
                    discrepancias.append(
                        (
                            raw[
                                "indice_bloque"
                            ],
                            variable,
                            valor_parser,
                            valor_raw_convertido,
                        )
                    )

        if discrepancias:
            escribir(
                "Discrepancias parser/raw:"
            )

            for item in discrepancias:
                escribir(
                    f"  * {item}"
                )

        else:
            escribir(
                "Discrepancias parser/raw: ninguna"
            )

        # ----------------------------------------------------
        # Mostrar específicamente agregado raw solicitado.
        # ----------------------------------------------------

        titulo(
            "TOTAL BANCA MÚLTIPLE - VALORES RAW"
        )

        raw_agregado = [
            b
            for b in bloques_raw
            if b["es_agregado"]
        ]

        if len(raw_agregado) == 1:
            raw_agregado = raw_agregado[
                0
            ]

            escribir(
                f"Nombre = "
                f"{raw_agregado['banco_original']!r}"
            )

            escribir(
                f"MN raw = "
                f"{raw_agregado['mn_raw_repr']}"
            )

            escribir(
                f"ME raw = "
                f"{raw_agregado['me_raw_repr']}"
            )

            escribir(
                f"TOTAL raw = "
                f"{raw_agregado['total_raw_repr']}"
            )

            escribir(
                f"MN formato = "
                f"{raw_agregado['mn_format_str']!r}"
            )

            escribir(
                f"ME formato = "
                f"{raw_agregado['me_format_str']!r}"
            )

            escribir(
                f"TOTAL formato = "
                f"{raw_agregado['total_format_str']!r}"
            )

        else:
            escribir(
                "No se pudo aislar exactamente "
                "un bloque agregado raw."
            )

        # ----------------------------------------------------
        # Guardar CSV auxiliares.
        # ----------------------------------------------------

        guardar_csv_controles(
            controles
        )

        guardar_csv_bloques(
            bloques_raw
        )

        # ----------------------------------------------------
        # Clasificación DIAGNÓSTICA, no validación de Y.
        # ----------------------------------------------------

        titulo(
            "CLASIFICACIÓN DEL DIAGNÓSTICO"
        )

        escribir(
            "ESTADO = "
            "DIAGNOSTICO_Y_B2359_2017_07_COMPLETADO"
        )

        escribir()

        escribir(
            "Este estado NO significa que "
            "Y / 2017-07 esté validada."
        )

        escribir(
            "No se modificó producción "
            "ni ninguna tolerancia."
        )

    except Exception as error:
        titulo(
            "ERROR DEL DIAGNÓSTICO"
        )

        escribir(
            f"Tipo = "
            f"{type(error).__name__}"
        )

        escribir(
            f"Detalle = "
            f"{repr(error)}"
        )

        escribir(
            "ESTADO = "
            "DIAGNOSTICO_Y_B2359_2017_07_ERROR"
        )

        raise

    finally:
        TXT_RESULTADO.write_text(
            "\n".join(
                lineas_reporte
            ),
            encoding="utf-8",
        )

        print()
        print(
            "=" * 122
        )

        print(
            "DIAGNÓSTICO FINALIZADO"
        )

        print(
            "=" * 122
        )

        print()
        print(
            "TXT:"
        )

        print(
            TXT_RESULTADO
        )

        print()
        print(
            "CSV controles:"
        )

        print(
            CSV_CONTROLES
        )

        print()
        print(
            "CSV bloques/raw:"
        )

        print(
            CSV_BLOQUES
        )


# ============================================================
# 22. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":
    main()