# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Diagnóstico estructural piloto - Movimiento de los Depósitos SBS B-2318

from pathlib import Path
import re

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. CONFIGURACIÓN
# ============================================================

RAIZ = Path(__file__).resolve().parent.parent

CARPETA_PILOTO = (
    RAIZ
    / "datos_crudos"
    / "sbs_piloto"
)

ARCHIVOS_CONTROL = [
    "depositos_2008_06.xls",
    "depositos_2010_12.xls",
    "depositos_2011_12.xls",
    "depositos_2015_06.xls",
    "depositos_2020_06.xls",
    "depositos_2025_12.xls",
]


# ============================================================
# 2. NORMALIZAR TEXTO
# ============================================================

def normalizar_texto(valor):

    if valor is None:
        return ""

    texto = str(valor)

    texto = texto.replace("\n", " ")
    texto = texto.replace("\r", " ")

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


# ============================================================
# 3. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato(ruta):

    with open(ruta, "rb") as archivo:
        cabecera = archivo.read(8)

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:
        return "XLS_OLE"

    if cabecera.startswith(firmas_zip):
        return "XLSX_ZIP"

    return "DESCONOCIDO"


# ============================================================
# 4. LEER XLS/OLE
# ============================================================

def leer_xls(ruta):

    libro = xlrd.open_workbook(
        filename=str(ruta),
        on_demand=True
    )

    hojas = []

    for nombre in libro.sheet_names():

        hoja = libro.sheet_by_name(nombre)

        filas = []

        for i in range(hoja.nrows):

            fila = [
                hoja.cell_value(i, j)
                for j in range(hoja.ncols)
            ]

            filas.append(fila)

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.release_resources()

    return hojas


# ============================================================
# 5. LEER XLSX/ZIP
# ============================================================

def leer_xlsx(ruta):

    with open(ruta, "rb") as archivo_binario:

        libro = load_workbook(
            archivo_binario,
            read_only=True,
            data_only=True
        )

        hojas = []

        for nombre in libro.sheetnames:

            hoja = libro[nombre]

            filas = [
                list(fila)
                for fila in hoja.iter_rows(
                    values_only=True
                )
            ]

            hojas.append({
                "nombre": nombre,
                "filas": filas,
            })

        libro.close()

    return hojas


# ============================================================
# 6. LEER ARCHIVO
# ============================================================

def leer_archivo(ruta):

    formato = detectar_formato(ruta)

    if formato == "XLS_OLE":
        hojas = leer_xls(ruta)

    elif formato == "XLSX_ZIP":
        hojas = leer_xlsx(ruta)

    else:
        raise ValueError(
            f"Formato no reconocido: {ruta.name}"
        )

    return formato, hojas


# ============================================================
# 7. PALABRAS CLAVE QUE NOS INTERESAN
# ============================================================

PALABRAS_CLAVE = [
    "moneda nacional",
    "moneda extranjera",
    "empresas",
    "empresa",
    "saldo anterior",
    "abonos",
    "intereses",
    "retiros",
    "cargos",
    "saldo final",
    "tipo de cambio",
    "total banca",
]


# ============================================================
# 8. BUSCAR PALABRAS CLAVE
# ============================================================

def buscar_coincidencias(filas):

    coincidencias = []

    for numero_fila, fila in enumerate(filas):

        encontrados = []

        for numero_columna, valor in enumerate(fila):

            texto = normalizar_texto(valor)

            if not texto:
                continue

            for palabra in PALABRAS_CLAVE:

                if palabra in texto:

                    encontrados.append({
                        "columna": numero_columna,
                        "valor": str(valor).strip(),
                        "palabra": palabra,
                    })

                    break

        if encontrados:

            coincidencias.append({
                "fila": numero_fila,
                "encontrados": encontrados,
            })

    return coincidencias


# ============================================================
# 9. MOSTRAR FILA COMPACTA
# ============================================================

def mostrar_fila(numero_fila, fila):

    celdas = []

    for numero_columna, valor in enumerate(fila):

        texto = str(valor).strip()

        if texto:

            celdas.append(
                f"C{numero_columna}={repr(texto)}"
            )

    if celdas:

        print(
            f"Fila {numero_fila}: "
            + " | ".join(celdas)
        )


# ============================================================
# 10. DIAGNOSTICAR UNA HOJA
# ============================================================

def diagnosticar_hoja(nombre, filas):

    max_columnas = max(
        (len(fila) for fila in filas),
        default=0
    )

    print()
    print("-" * 100)

    print(
        "HOJA:",
        repr(nombre)
    )

    print(
        "DIMENSIÓN:",
        f"{len(filas)} x {max_columnas}"
    )

    coincidencias = buscar_coincidencias(
        filas
    )

    print()
    print("FILAS CON PALABRAS CLAVE")
    print("-" * 100)

    if not coincidencias:

        print(
            "No se encontraron palabras clave."
        )

    else:

        for item in coincidencias:

            numero_fila = item["fila"]

            mostrar_fila(
                numero_fila,
                filas[numero_fila]
            )

    print()
    print("PRIMERAS 15 FILAS NO VACÍAS")
    print("-" * 100)

    contador = 0

    for numero_fila, fila in enumerate(filas):

        tiene_contenido = any(
            normalizar_texto(valor)
            for valor in fila
        )

        if not tiene_contenido:
            continue

        mostrar_fila(
            numero_fila,
            fila
        )

        contador += 1

        if contador >= 15:
            break


# ============================================================
# 11. DIAGNOSTICAR ARCHIVO
# ============================================================

def diagnosticar_archivo(ruta):

    print()
    print("=" * 100)

    print(
        "ARCHIVO:",
        ruta.name
    )

    print("=" * 100)

    formato, hojas = leer_archivo(
        ruta
    )

    print(
        "FORMATO REAL:",
        formato
    )

    print(
        "NÚMERO DE HOJAS:",
        len(hojas)
    )

    print(
        "NOMBRES DE HOJAS:",
        [
            hoja["nombre"]
            for hoja in hojas
        ]
    )

    for hoja in hojas:

        diagnosticar_hoja(
            hoja["nombre"],
            hoja["filas"]
        )


# ============================================================
# 12. FUNCIÓN PRINCIPAL
# ============================================================

def main():

    print()
    print("=" * 100)

    print(
        "DIAGNÓSTICO B-2318"
    )

    print(
        "MOVIMIENTO DE LOS DEPÓSITOS"
    )

    print("=" * 100)

    print()
    print(
        "Esta prueba NO descarga archivos."
    )

    print(
        "Esta prueba NO modifica "
        "los archivos crudos."
    )

    encontrados = 0

    for nombre in ARCHIVOS_CONTROL:

        ruta = (
            CARPETA_PILOTO
            / nombre
        )

        if not ruta.exists():

            print()
            print(
                "[ERROR] No existe:",
                ruta
            )

            continue

        diagnosticar_archivo(ruta)

        encontrados += 1

    print()
    print("=" * 100)

    print(
        "ARCHIVOS DIAGNOSTICADOS:",
        encontrados,
        "de",
        len(ARCHIVOS_CONTROL)
    )

    print("=" * 100)


# ============================================================
# 13. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":
    main()