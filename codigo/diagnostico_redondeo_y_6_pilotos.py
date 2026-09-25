# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Fecha de extracción: 2026-09-24

from pathlib import Path
import csv
import hashlib
import math
import re
import unicodedata

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. OBJETIVO
# ============================================================

"""
DIAGNÓSTICO LOCAL DE REDONDEO / PRECISIÓN
Y = Dolarización del crédito (%)
Reporte SBS B-2359

Pilotos:

    2008-06
    2010-12
    2011-12
    2015-06
    2020-06
    2025-12

NO:

    - descarga archivos;
    - consulta Internet;
    - modifica el parser;
    - modifica DOLCRED;
    - modifica tolerancias de producción;
    - declara Y validada;
    - ejecuta X1;
    - interpola;
    - homologa nombres;
    - selecciona bancos.

Para cada mes y para MN, ME y Total registra:

    - cantidad de bancos individuales;
    - sum();
    - math.fsum();
    - Total Banca Múltiple SBS;
    - diferencia absoluta;
    - diferencia relativa;
    - formato Excel del agregado;
    - decimales mostrados;
    - round(suma, decimales) vs
      round(total SBS, decimales).
"""


# ============================================================
# 2. CONFIGURACIÓN
# ============================================================

PILOTOS = [
    (2008, 6),
    (2010, 12),
    (2011, 12),
    (2015, 6),
    (2020, 6),
    (2025, 12),
]

CODIGO_MES_SBS = {
    1: "en",
    2: "fe",
    3: "ma",
    4: "ab",
    5: "my",
    6: "jn",
    7: "jl",
    8: "ag",
    9: "se",
    10: "oc",
    11: "no",
    12: "di",
}


# ============================================================
# 3. RUTAS
# ============================================================

RAIZ = Path(
    __file__
).resolve().parent.parent

CARPETA_DATOS_CRUDOS = (
    RAIZ
    / "datos_crudos"
)

CARPETA_DIAGNOSTICOS = (
    RAIZ
    / "salidas"
    / "diagnosticos"
)

CARPETA_DIAGNOSTICOS.mkdir(
    parents=True,
    exist_ok=True
)

TXT_RESULTADO = (
    CARPETA_DIAGNOSTICOS
    / "diagnostico_redondeo_Y_6_pilotos.txt"
)

CSV_RESULTADO = (
    CARPETA_DIAGNOSTICOS
    / "diagnostico_redondeo_Y_6_pilotos.csv"
)


# ============================================================
# 4. RUTAS MANUALES OPCIONALES
#
# Déjalas en None si la búsqueda automática
# reconoce correctamente los archivos.
# ============================================================

RUTAS_MANUALES = {
    "2008-06": None,
    "2010-12": None,
    "2011-12": None,
    "2015-06": None,
    "2020-06": None,
    "2025-12": None,
}


# ============================================================
# 5. REPORTE
# ============================================================

lineas_reporte = []


def escribir(texto=""):

    texto = str(
        texto
    )

    print(
        texto
    )

    lineas_reporte.append(
        texto
    )


def titulo(texto):

    escribir()
    escribir("=" * 122)
    escribir(texto)
    escribir("=" * 122)


def subtitulo(texto):

    escribir()
    escribir("-" * 122)
    escribir(texto)
    escribir("-" * 122)


# ============================================================
# 6. UTILIDADES DE TEXTO
#
# MISMAS BASES DEL PARSER VALIDADO.
# ============================================================

def normalizar_texto(valor):

    if valor is None:
        return ""

    texto = str(
        valor
    )

    texto = texto.replace(
        "\n",
        " "
    )

    texto = texto.replace(
        "\r",
        " "
    )

    texto = texto.replace(
        "\xa0",
        " "
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


def normalizar_comparacion(valor):

    texto = normalizar_texto(
        valor
    )

    texto = unicodedata.normalize(
        "NFKD",
        texto
    )

    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(
            caracter
        )
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip()


def sin_tildes(valor):

    return normalizar_comparacion(
        valor
    )


# ============================================================
# 7. CONVERSIÓN NUMÉRICA
#
# MISMO CRITERIO DEL PARSER.
# ============================================================

def convertir_numero(valor):

    if valor is None:
        return None

    if isinstance(
        valor,
        (int, float)
    ):

        return float(
            valor
        )

    texto = str(
        valor
    ).strip()

    if texto == "":
        return None

    if texto.lower() in {
        "-",
        "s.i.",
        "s.i",
        "n.d.",
        "n.d",
        "nd",
        "n/a",
        "na",
    }:

        return None

    try:

        return float(
            texto.replace(
                ",",
                ""
            )
        )

    except ValueError:

        return None


# ============================================================
# 8. HASH LOCAL
# ============================================================

def sha256_archivo(ruta):

    hash_obj = hashlib.sha256()

    with ruta.open(
        "rb"
    ) as archivo:

        while True:

            bloque = archivo.read(
                1024 * 1024
            )

            if not bloque:
                break

            hash_obj.update(
                bloque
            )

    return hash_obj.hexdigest()


# ============================================================
# 9. IDENTIFICADORES DEL MES
# ============================================================

def clave_mes(
    year,
    month
):

    return (
        f"{year:04d}-{month:02d}"
    )


def patrones_nombre_mes(
    year,
    month
):

    codigo = CODIGO_MES_SBS[
        month
    ]

    return [
        f"b-2359-{codigo}{year}",
        f"b2359-{codigo}{year}",
        f"b-2359_{year}_{month:02d}",
        f"b-2359-{year}-{month:02d}",
        f"b2359_{year}_{month:02d}",
        f"b2359-{year}-{month:02d}",
        f"b-2359_{year}{month:02d}",
        f"b2359_{year}{month:02d}",

        # ====================================================
        # NUEVO PATRÓN PARA TUS ARCHIVOS PILOTO LOCALES
        # Ejemplo:
        # creditos_2015_06.xls
        # ====================================================
        f"creditos_{year}_{month:02d}.xls",
    ]


# ============================================================
# 10. BUSCAR ARCHIVO PILOTO LOCAL
# ============================================================

def buscar_archivo_piloto(
    year,
    month
):

    mes = clave_mes(
        year,
        month
    )

    ruta_manual = RUTAS_MANUALES.get(
        mes
    )

    if ruta_manual:

        ruta = Path(
            ruta_manual
        )

        if not ruta.is_absolute():

            ruta = (
                RAIZ
                /
                ruta
            )

        if not ruta.exists():

            raise RuntimeError(
                f"{mes}: la ruta manual no existe: "
                f"{ruta}"
            )

        return ruta

    patrones = patrones_nombre_mes(
        year,
        month
    )

    candidatos = []

    if not CARPETA_DATOS_CRUDOS.exists():

        raise RuntimeError(
            "No existe la carpeta datos_crudos."
        )

    for ruta in CARPETA_DATOS_CRUDOS.rglob(
        "*"
    ):

        if not ruta.is_file():
            continue

        if ruta.suffix.lower() not in {
            ".xls",
            ".xlsx",
        }:
            continue

        nombre = ruta.name.lower()

        if any(
            patron in nombre
            for patron in patrones
        ):

            candidatos.append(
                ruta
            )

    candidatos = sorted(
        set(
            candidatos
        )
    )

    if len(
        candidatos
    ) == 0:

        raise RuntimeError(
            f"{mes}: no se encontró localmente "
            "el archivo piloto B-2359. "
            "Si tiene otro nombre, colócalo "
            "en RUTAS_MANUALES."
        )

    if len(
        candidatos
    ) == 1:

        return candidatos[
            0
        ]

    # ========================================================
    # Si existen varias COPIAS idénticas,
    # no hay ambigüedad de contenido.
    # ========================================================

    por_hash = {}

    for ruta in candidatos:

        hash_archivo = sha256_archivo(
            ruta
        )

        por_hash.setdefault(
            hash_archivo,
            []
        ).append(
            ruta
        )

    if len(
        por_hash
    ) == 1:

        rutas_iguales = next(
            iter(
                por_hash.values()
            )
        )

        seleccionada = sorted(
            rutas_iguales,
            key=lambda x:
                (
                    len(
                        str(x)
                    ),
                    str(x)
                )
        )[0]

        escribir(
            f"{mes}: existen "
            f"{len(rutas_iguales)} copias "
            "idénticas por SHA-256."
        )

        escribir(
            f"  Se usará: {seleccionada}"
        )

        return seleccionada

    detalle = "\n".join(
        f"  - {ruta}"
        for ruta in candidatos
    )

    raise RuntimeError(
        f"{mes}: se encontraron varios "
        "archivos candidatos B-2359 "
        "con contenido distinto.\n"
        f"{detalle}"
    )


# ============================================================
# 11. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato(
    ruta
):

    with ruta.open(
        "rb"
    ) as archivo:

        cabecera = archivo.read(
            8
        )

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:

        return "XLS_OLE"

    if cabecera.startswith(
        firmas_zip
    ):

        return "XLSX_ZIP"

    return "DESCONOCIDO"


# ============================================================
# 12. LECTURA NORMAL XLS
# ============================================================

def leer_xls(
    ruta
):

    libro = xlrd.open_workbook(
        filename=str(
            ruta
        ),
        on_demand=True
    )

    hojas = []

    for nombre in libro.sheet_names():

        hoja = libro.sheet_by_name(
            nombre
        )

        filas = []

        for i in range(
            hoja.nrows
        ):

            fila = [
                hoja.cell_value(
                    i,
                    j
                )
                for j in range(
                    hoja.ncols
                )
            ]

            filas.append(
                fila
            )

        hojas.append({
            "nombre":
                nombre,

            "filas":
                filas,
        })

    libro.release_resources()

    return hojas


# ============================================================
# 13. LECTURA NORMAL XLSX
# ============================================================

def leer_xlsx(
    ruta
):

    with ruta.open(
        "rb"
    ) as archivo_binario:

        libro = load_workbook(
            archivo_binario,
            read_only=True,
            data_only=True
        )

        hojas = []

        for nombre in libro.sheetnames:

            hoja = libro[
                nombre
            ]

            filas = [
                list(
                    fila
                )
                for fila in hoja.iter_rows(
                    values_only=True
                )
            ]

            hojas.append({
                "nombre":
                    nombre,

                "filas":
                    filas,
            })

        libro.close()

    return hojas


# ============================================================
# 14. LEER ARCHIVO
# ============================================================

def leer_archivo(
    ruta
):

    formato = detectar_formato(
        ruta
    )

    if formato == "XLS_OLE":

        hojas = leer_xls(
            ruta
        )

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(
            ruta
        )

    else:

        raise RuntimeError(
            f"Formato no reconocido: {ruta}"
        )

    return (
        formato,
        hojas,
    )


# ============================================================
# 15. RECONOCER MN / ME / TOTAL
#
# MISMO PARSER VALIDADO.
# ============================================================

def es_mn(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith(
            "mn"
        )
        and
        "mil" in texto
    )


def es_me(valor):

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith(
            "me"
        )
        and
        "mil" in texto
    )


def es_total(valor):

    texto = normalizar_texto(
        valor
    )

    return texto.startswith(
        "total"
    )


# ============================================================
# 16. FILA DE ENCABEZADOS
# ============================================================

def encontrar_fila_encabezados_creditos(
    filas
):

    candidatos = []

    for numero_fila, fila in enumerate(
        filas
    ):

        bloques = []

        for columna in range(
            max(
                0,
                len(fila) - 2
            )
        ):

            if (
                es_mn(
                    fila[
                        columna
                    ]
                )
                and
                es_me(
                    fila[
                        columna + 1
                    ]
                )
                and
                es_total(
                    fila[
                        columna + 2
                    ]
                )
            ):

                bloques.append(
                    columna
                )

        if bloques:

            candidatos.append({
                "fila":
                    numero_fila,

                "bloques":
                    bloques,
            })

    if not candidatos:

        return None

    return max(
        candidatos,
        key=lambda x:
        len(
            x[
                "bloques"
            ]
        )
    )


# ============================================================
# 17. NOMBRE ORIGINAL DEL BANCO
# ============================================================

def obtener_nombre_banco_creditos(
    filas,
    fila_bancos,
    columna_inicio
):

    if (
        fila_bancos < 0
        or
        fila_bancos >= len(
            filas
        )
    ):

        return ""

    fila = filas[
        fila_bancos
    ]

    for desplazamiento in (
        0,
        1,
        2,
    ):

        columna = (
            columna_inicio
            +
            desplazamiento
        )

        if columna >= len(
            fila
        ):
            continue

        valor = fila[
            columna
        ]

        if normalizar_texto(
            valor
        ):

            return str(
                valor
            ).strip()

    return ""


# ============================================================
# 18. ESTRUCTURA DE CRÉDITOS
# ============================================================

def detectar_estructura_creditos(
    filas
):

    encabezado = encontrar_fila_encabezados_creditos(
        filas
    )

    if encabezado is None:

        return None

    fila_encabezados = encabezado[
        "fila"
    ]

    fila_bancos = (
        fila_encabezados
        -
        1
    )

    bloques = []

    for columna_inicio in encabezado[
        "bloques"
    ]:

        banco = obtener_nombre_banco_creditos(
            filas,
            fila_bancos,
            columna_inicio
        )

        bloques.append({
            "banco_original":
                banco,

            "columna_mn":
                columna_inicio,

            "columna_me":
                columna_inicio + 1,

            "columna_total":
                columna_inicio + 2,
        })

    return {
        "fila_bancos":
            fila_bancos,

        "fila_encabezados":
            fila_encabezados,

        "bloques":
            bloques,
    }


# ============================================================
# 19. ELEGIR HOJA PRINCIPAL
# ============================================================

def elegir_hoja_principal_creditos(
    hojas
):

    candidatos = []

    for hoja in hojas:

        estructura = detectar_estructura_creditos(
            hoja[
                "filas"
            ]
        )

        if estructura is None:
            continue

        candidatos.append({
            "hoja":
                hoja,

            "estructura":
                estructura,
        })

    if not candidatos:

        return None

    return max(
        candidatos,
        key=lambda x:
        len(
            x[
                "estructura"
            ][
                "bloques"
            ]
        )
    )


# ============================================================
# 20. FILA TOTAL CRÉDITOS
# ============================================================

def buscar_fila_total_creditos(
    filas
):

    for numero_fila, fila in enumerate(
        filas
    ):

        for valor in fila:

            texto = sin_tildes(
                valor
            )

            if (
                "total creditos"
                in texto
            ):

                return numero_fila

    return None


# ============================================================
# 21. TOTAL BANCA MÚLTIPLE
# ============================================================

def es_agregado_creditos(
    nombre
):

    return (
        "total banca multiple"
        in
        sin_tildes(
            nombre
        )
    )


# ============================================================
# 22. PARSER Y
#
# MISMO PARSER Y MISMA FÓRMULA DOLCRED.
# ============================================================

def extraer_y_creditos(
    ruta
):

    formato, hojas = leer_archivo(
        ruta
    )

    seleccionado = elegir_hoja_principal_creditos(
        hojas
    )

    if seleccionado is None:

        raise RuntimeError(
            "No se encontró la estructura "
            "MN|ME|Total."
        )

    hoja = seleccionado[
        "hoja"
    ]

    estructura = seleccionado[
        "estructura"
    ]

    filas = hoja[
        "filas"
    ]

    fila_total = buscar_fila_total_creditos(
        filas
    )

    if fila_total is None:

        raise RuntimeError(
            "No se encontró "
            "'Total Créditos:'."
        )

    fila = filas[
        fila_total
    ]

    registros = []

    for bloque in estructura[
        "bloques"
    ]:

        banco = bloque[
            "banco_original"
        ]

        col_mn = bloque[
            "columna_mn"
        ]

        col_me = bloque[
            "columna_me"
        ]

        col_total = bloque[
            "columna_total"
        ]

        mn = (
            convertir_numero(
                fila[
                    col_mn
                ]
            )
            if col_mn < len(
                fila
            )
            else None
        )

        me = (
            convertir_numero(
                fila[
                    col_me
                ]
            )
            if col_me < len(
                fila
            )
            else None
        )

        total = (
            convertir_numero(
                fila[
                    col_total
                ]
            )
            if col_total < len(
                fila
            )
            else None
        )

        problemas = []

        if not banco:

            problemas.append(
                "BANCO_SIN_NOMBRE"
            )

        if mn is None:

            problemas.append(
                "MN_FALTANTE"
            )

        if me is None:

            problemas.append(
                "ME_FALTANTE"
            )

        if total is None:

            problemas.append(
                "TOTAL_FALTANTE"
            )

        if (
            total is not None
            and
            total <= 0
        ):

            problemas.append(
                "TOTAL_NO_POSITIVO"
            )

        dolcred = None
        tc_implicito = None

        # ====================================================
        # MISMA FÓRMULA VALIDADA
        # ====================================================

        if (
            total is not None
            and
            total > 0
            and
            mn is not None
        ):

            dolcred = (
                (
                    total
                    -
                    mn
                )
                /
                total
                *
                100
            )

        if (
            mn is not None
            and
            me is not None
            and
            total is not None
            and
            me > 0
        ):

            tc_implicito = (
                (
                    total
                    -
                    mn
                )
                /
                me
            )

        registros.append({
            "banco_original":
                banco,

            "mn_soles_miles":
                mn,

            "me_usd_miles":
                me,

            "total_soles_miles":
                total,

            "dolcred_pct":
                dolcred,

            "tc_implicito":
                tc_implicito,

            "es_agregado":
                es_agregado_creditos(
                    banco
                ),

            "problemas":
                problemas,

            "columna_mn":
                col_mn,

            "columna_me":
                col_me,

            "columna_total":
                col_total,
        })

    return {
        "formato":
            formato,

        "hoja":
            hoja[
                "nombre"
            ],

        "fila_encabezados":
            estructura[
                "fila_encabezados"
            ],

        "fila_bancos":
            estructura[
                "fila_bancos"
            ],

        "fila_total_creditos":
            fila_total,

        "cantidad_bloques":
            len(
                estructura[
                    "bloques"
                ]
            ),

        "registros":
            registros,
    }


# ============================================================
# 23. INFERIR DECIMALES MOSTRADOS
# ============================================================

def inferir_decimales_formato_excel(
    formato
):

    if formato is None:
        return None

    texto = str(
        formato
    ).strip()

    if not texto:
        return None

    if texto.lower() == "general":
        return None

    seccion = texto.split(
        ";"
    )[0]

    seccion = re.sub(
        r'"[^"]*"',
        "",
        seccion
    )

    seccion = re.sub(
        r"\[[^\]]*\]",
        "",
        seccion
    )

    coincidencia = re.search(
        r"\.([0#?]+)",
        seccion
    )

    if coincidencia:

        return len(
            coincidencia.group(1)
        )

    if re.search(
        r"[0#?]",
        seccion
    ):

        return 0

    return None


# ============================================================
# 24. FORMATO DE CELDA DEL AGREGADO - XLS/OLE
# ============================================================

def obtener_formato_agregado_xls(
    ruta,
    nombre_hoja,
    fila_total,
    agregado
):

    libro = xlrd.open_workbook(
        filename=str(
            ruta
        ),
        formatting_info=True,
        on_demand=True
    )

    hoja = libro.sheet_by_name(
        nombre_hoja
    )

    salida = {}

    for variable, columna in (
        (
            "MN",
            agregado[
                "columna_mn"
            ]
        ),
        (
            "ME",
            agregado[
                "columna_me"
            ]
        ),
        (
            "TOTAL",
            agregado[
                "columna_total"
            ]
        ),
    ):

        celda = hoja.cell(
            fila_total,
            columna
        )

        xf_index = hoja.cell_xf_index(
            fila_total,
            columna
        )

        format_key = None
        format_str = None

        if (
            0 <= xf_index < len(
                libro.xf_list
            )
        ):

            xf = libro.xf_list[
                xf_index
            ]

            format_key = getattr(
                xf,
                "format_key",
                None
            )

            if format_key is not None:

                formato_obj = libro.format_map.get(
                    format_key
                )

                if formato_obj is not None:

                    format_str = getattr(
                        formato_obj,
                        "format_str",
                        None
                    )

        salida[
            variable
        ] = {
            "valor_raw":
                celda.value,

            "valor_raw_repr":
                repr(
                    celda.value
                ),

            "xf_index":
                xf_index,

            "format_key":
                format_key,

            "format_str":
                format_str,

            "decimales_mostrados":
                inferir_decimales_formato_excel(
                    format_str
                ),
        }

    libro.release_resources()

    return salida


# ============================================================
# 25. FORMATO DEL AGREGADO - XLSX
# ============================================================

def obtener_formato_agregado_xlsx(
    ruta,
    nombre_hoja,
    fila_total,
    agregado
):

    with ruta.open(
        "rb"
    ) as archivo_binario:

        libro = load_workbook(
            archivo_binario,
            read_only=False,
            data_only=True
        )

        hoja = libro[
            nombre_hoja
        ]

        salida = {}

        for variable, columna in (
            (
                "MN",
                agregado[
                    "columna_mn"
                ]
            ),
            (
                "ME",
                agregado[
                    "columna_me"
                ]
            ),
            (
                "TOTAL",
                agregado[
                    "columna_total"
                ]
            ),
        ):

            celda = hoja.cell(
                row=fila_total + 1,
                column=columna + 1
            )

            formato = celda.number_format

            salida[
                variable
            ] = {
                "valor_raw":
                    celda.value,

                "valor_raw_repr":
                    repr(
                        celda.value
                    ),

                "xf_index":
                    celda.style_id,

                "format_key":
                    "",

                "format_str":
                    formato,

                "decimales_mostrados":
                    inferir_decimales_formato_excel(
                        formato
                    ),
            }

        libro.close()

    return salida


# ============================================================
# 26. FORMATO DEL AGREGADO
# ============================================================

def obtener_formatos_agregado(
    ruta,
    resultado,
    agregado
):

    if resultado[
        "formato"
    ] == "XLS_OLE":

        return obtener_formato_agregado_xls(
            ruta,
            resultado[
                "hoja"
            ],
            resultado[
                "fila_total_creditos"
            ],
            agregado
        )

    if resultado[
        "formato"
    ] == "XLSX_ZIP":

        return obtener_formato_agregado_xlsx(
            ruta,
            resultado[
                "hoja"
            ],
            resultado[
                "fila_total_creditos"
            ],
            agregado
        )

    raise RuntimeError(
        "Formato Excel no compatible "
        "con inspección de estilos."
    )


# ============================================================
# 27. DIAGNÓSTICO DE UN MES
# ============================================================

def diagnosticar_mes(
    year,
    month
):

    mes = clave_mes(
        year,
        month
    )

    titulo(
        f"PILOTO {mes}"
    )

    ruta = buscar_archivo_piloto(
        year,
        month
    )

    escribir(
        f"Archivo local = {ruta}"
    )

    escribir(
        f"SHA-256 = {sha256_archivo(ruta)}"
    )

    resultado = extraer_y_creditos(
        ruta
    )

    escribir(
        f"Formato real = "
        f"{resultado['formato']}"
    )

    escribir(
        f"Hoja = "
        f"{resultado['hoja']!r}"
    )

    escribir(
        f"Fila Total Créditos = "
        f"{resultado['fila_total_creditos']}"
    )

    escribir(
        f"Bloques detectados = "
        f"{resultado['cantidad_bloques']}"
    )

    registros = resultado[
        "registros"
    ]

    agregados = [
        registro
        for registro in registros
        if registro[
            "es_agregado"
        ]
    ]

    bancos = [
        registro
        for registro in registros
        if not registro[
            "es_agregado"
        ]
    ]

    if len(
        agregados
    ) != 1:

        raise RuntimeError(
            f"{mes}: se esperaba exactamente "
            "un Total Banca Múltiple, "
            f"pero se encontraron "
            f"{len(agregados)}."
        )

    agregado = agregados[
        0
    ]

    escribir(
        f"Cantidad de bancos individuales = "
        f"{len(bancos)}"
    )

    formatos = obtener_formatos_agregado(
        ruta,
        resultado,
        agregado
    )

    filas_resultado = []

    configuracion = [
        (
            "MN",
            "mn_soles_miles"
        ),
        (
            "ME",
            "me_usd_miles"
        ),
        (
            "TOTAL",
            "total_soles_miles"
        ),
    ]

    for variable, campo in configuracion:

        valores = [
            registro[
                campo
            ]
            for registro in bancos
        ]

        faltantes = [
            registro[
                "banco_original"
            ]
            for registro in bancos
            if registro[
                campo
            ]
            is None
        ]

        if faltantes:

            raise RuntimeError(
                f"{mes} / {variable}: "
                "existen valores faltantes y "
                "no puede calcularse el control. "
                f"Bancos: {faltantes}"
            )

        total_sbs = agregado[
            campo
        ]

        if total_sbs is None:

            raise RuntimeError(
                f"{mes} / {variable}: "
                "Total Banca Múltiple está vacío."
            )

        suma_sum = sum(
            valores
        )

        suma_fsum = math.fsum(
            valores
        )

        diferencia_sum = (
            suma_sum
            -
            total_sbs
        )

        diferencia_fsum = (
            suma_fsum
            -
            total_sbs
        )

        diferencia_absoluta_sum = abs(
            diferencia_sum
        )

        diferencia_absoluta_fsum = abs(
            diferencia_fsum
        )

        if total_sbs != 0:

            diferencia_relativa_sum = (
                diferencia_absoluta_sum
                /
                abs(
                    total_sbs
                )
            )

            diferencia_relativa_fsum = (
                diferencia_absoluta_fsum
                /
                abs(
                    total_sbs
                )
            )

        else:

            diferencia_relativa_sum = None
            diferencia_relativa_fsum = None

        metadata = formatos[
            variable
        ]

        decimales = metadata[
            "decimales_mostrados"
        ]

        if decimales is None:

            round_sum = None
            round_fsum = None
            round_total = None
            coincide_sum = None
            coincide_fsum = None

        else:

            round_sum = round(
                suma_sum,
                decimales
            )

            round_fsum = round(
                suma_fsum,
                decimales
            )

            round_total = round(
                total_sbs,
                decimales
            )

            coincide_sum = (
                round_sum
                ==
                round_total
            )

            coincide_fsum = (
                round_fsum
                ==
                round_total
            )

        fila = {
            "mes":
                mes,

            "variable":
                variable,

            "cantidad_bancos":
                len(
                    bancos
                ),

            "suma_sum":
                suma_sum,

            "suma_fsum":
                suma_fsum,

            "total_sbs":
                total_sbs,

            "diferencia_sum":
                diferencia_sum,

            "diferencia_absoluta_sum":
                diferencia_absoluta_sum,

            "diferencia_relativa_sum":
                diferencia_relativa_sum,

            "diferencia_fsum":
                diferencia_fsum,

            "diferencia_absoluta_fsum":
                diferencia_absoluta_fsum,

            "diferencia_relativa_fsum":
                diferencia_relativa_fsum,

            "diferencia_sum_vs_fsum":
                (
                    suma_sum
                    -
                    suma_fsum
                ),

            "formato_excel_agregado":
                metadata[
                    "format_str"
                ],

            "decimales_mostrados":
                decimales,

            "valor_raw_agregado":
                metadata[
                    "valor_raw"
                ],

            "valor_raw_agregado_repr":
                metadata[
                    "valor_raw_repr"
                ],

            "xf_index_agregado":
                metadata[
                    "xf_index"
                ],

            "format_key_agregado":
                metadata[
                    "format_key"
                ],

            "round_sum":
                round_sum,

            "round_fsum":
                round_fsum,

            "round_total_sbs":
                round_total,

            "coincide_round_sum":
                coincide_sum,

            "coincide_round_fsum":
                coincide_fsum,

            "coincide_precision_mostrada":
                (
                    coincide_sum
                    and
                    coincide_fsum
                    if (
                        coincide_sum is not None
                        and
                        coincide_fsum is not None
                    )
                    else None
                ),

            "archivo":
                str(
                    ruta
                ),

            "sha256":
                sha256_archivo(
                    ruta
                ),
        }

        filas_resultado.append(
            fila
        )

        subtitulo(
            f"{mes} / {variable}"
        )

        escribir(
            f"Cantidad bancos = "
            f"{fila['cantidad_bancos']}"
        )

        escribir(
            f"sum() = "
            f"{fila['suma_sum']!r}"
        )

        escribir(
            f"math.fsum() = "
            f"{fila['suma_fsum']!r}"
        )

        escribir(
            f"Total Banca Múltiple SBS = "
            f"{fila['total_sbs']!r}"
        )

        escribir(
            f"Diferencia absoluta sum = "
            f"{fila['diferencia_absoluta_sum']!r}"
        )

        escribir(
            f"Diferencia relativa sum = "
            f"{fila['diferencia_relativa_sum']!r}"
        )

        escribir(
            f"Diferencia absoluta fsum = "
            f"{fila['diferencia_absoluta_fsum']!r}"
        )

        escribir(
            f"Diferencia relativa fsum = "
            f"{fila['diferencia_relativa_fsum']!r}"
        )

        escribir(
            f"sum() - math.fsum() = "
            f"{fila['diferencia_sum_vs_fsum']!r}"
        )

        escribir(
            f"Formato Excel agregado = "
            f"{fila['formato_excel_agregado']!r}"
        )

        escribir(
            f"Decimales mostrados = "
            f"{fila['decimales_mostrados']}"
        )

        escribir(
            f"Valor raw agregado = "
            f"{fila['valor_raw_agregado_repr']}"
        )

        escribir(
            f"round(sum, decimales) = "
            f"{fila['round_sum']!r}"
        )

        escribir(
            f"round(fsum, decimales) = "
            f"{fila['round_fsum']!r}"
        )

        escribir(
            f"round(total SBS, decimales) = "
            f"{fila['round_total_sbs']!r}"
        )

        escribir(
            "Coincide a precisión mostrada "
            f"(sum) = "
            f"{fila['coincide_round_sum']}"
        )

        escribir(
            "Coincide a precisión mostrada "
            f"(fsum) = "
            f"{fila['coincide_round_fsum']}"
        )

    return filas_resultado


# ============================================================
# 28. GUARDAR CSV
# ============================================================

def guardar_csv(
    filas
):

    columnas = [
        "mes",
        "variable",
        "cantidad_bancos",
        "suma_sum",
        "suma_fsum",
        "total_sbs",
        "diferencia_sum",
        "diferencia_absoluta_sum",
        "diferencia_relativa_sum",
        "diferencia_fsum",
        "diferencia_absoluta_fsum",
        "diferencia_relativa_fsum",
        "diferencia_sum_vs_fsum",
        "formato_excel_agregado",
        "decimales_mostrados",
        "valor_raw_agregado",
        "valor_raw_agregado_repr",
        "xf_index_agregado",
        "format_key_agregado",
        "round_sum",
        "round_fsum",
        "round_total_sbs",
        "coincide_round_sum",
        "coincide_round_fsum",
        "coincide_precision_mostrada",
        "archivo",
        "sha256",
    ]

    with CSV_RESULTADO.open(
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as archivo:

        escritor = csv.DictWriter(
            archivo,
            fieldnames=columnas
        )

        escritor.writeheader()

        for fila in filas:

            escritor.writerow(
                fila
            )


# ============================================================
# 29. RESUMEN DE LOS 6 PILOTOS
# ============================================================

def mostrar_resumen(
    filas
):

    titulo(
        "RESUMEN DE LOS 6 PILOTOS"
    )

    escribir(
        "mes | variable | bancos | "
        "abs_diff_fsum | rel_diff_fsum | "
        "formato | decimales | "
        "round coincide"
    )

    escribir(
        "-" * 122
    )

    for fila in filas:

        escribir(
            f"{fila['mes']} | "
            f"{fila['variable']} | "
            f"{fila['cantidad_bancos']} | "
            f"{fila['diferencia_absoluta_fsum']} | "
            f"{fila['diferencia_relativa_fsum']} | "
            f"{fila['formato_excel_agregado']!r} | "
            f"{fila['decimales_mostrados']} | "
            f"{fila['coincide_precision_mostrada']}"
        )

    esperadas = (
        len(
            PILOTOS
        )
        *
        3
    )

    escribir()
    escribir(
        f"Controles esperados = "
        f"{esperadas}"
    )

    escribir(
        f"Controles obtenidos = "
        f"{len(filas)}"
    )

    comparables = [
        fila
        for fila in filas
        if fila[
            "coincide_precision_mostrada"
        ]
        is not None
    ]

    coincidentes = [
        fila
        for fila in comparables
        if fila[
            "coincide_precision_mostrada"
        ]
        is True
    ]

    no_coincidentes = [
        fila
        for fila in comparables
        if fila[
            "coincide_precision_mostrada"
        ]
        is False
    ]

    escribir(
        f"Controles con precisión interpretable = "
        f"{len(comparables)}"
    )

    escribir(
        f"Coinciden a precisión mostrada = "
        f"{len(coincidentes)}"
    )

    escribir(
        f"No coinciden a precisión mostrada = "
        f"{len(no_coincidentes)}"
    )

    escribir()
    escribir(
        "RESULTADO POR MES:"
    )

    todos_meses_coinciden = True

    for year, month in PILOTOS:

        mes = clave_mes(
            year,
            month
        )

        filas_mes = [
            fila
            for fila in filas
            if fila[
                "mes"
            ]
            == mes
        ]

        mes_ok = (
            len(
                filas_mes
            )
            == 3
            and
            all(
                fila[
                    "coincide_precision_mostrada"
                ]
                is True
                for fila in filas_mes
            )
        )

        escribir(
            f"  {mes}: "
            f"{'COINCIDE' if mes_ok else 'NO COINCIDE / REVISAR'}"
        )

        if not mes_ok:

            todos_meses_coinciden = False

    escribir()
    escribir(
        "TODOS_LOS_6_MESES_COINCIDEN_"
        "A_PRECISION_MOSTRADA = "
        f"{todos_meses_coinciden}"
    )

    titulo(
        "CLASIFICACIÓN"
    )

    escribir(
        "ESTADO = "
        "DIAGNOSTICO_REDONDEO_Y_6_PILOTOS_COMPLETADO"
    )

    escribir()
    escribir(
        "Este estado NO significa "
        "que Y esté validada."
    )

    escribir(
        "No se modificó ninguna tolerancia."
    )

    escribir(
        "No se modificó el parser."
    )

    escribir(
        "No se modificó DOLCRED."
    )

    escribir(
        "X1 no fue ejecutada ni modificada."
    )


# ============================================================
# 30. MAIN
# ============================================================

def main():

    filas_totales = []

    try:

        titulo(
            "DIAGNÓSTICO LOCAL DE REDONDEO "
            "Y / B-2359 - 6 PILOTOS"
        )

        escribir(
            "No se realizará ninguna descarga."
        )

        escribir(
            f"Carpeta de búsqueda local: "
            f"{CARPETA_DATOS_CRUDOS}"
        )

        escribir(
            "Parser B-2359: sin cambios."
        )

        escribir(
            "DOLCRED: sin cambios."
        )

        escribir(
            "X1: no se ejecuta."
        )

        escribir()

        escribir(
            "Pilotos:"
        )

        for year, month in PILOTOS:

            escribir(
                f"  - "
                f"{clave_mes(year, month)}"
            )

        for year, month in PILOTOS:

            filas_mes = diagnosticar_mes(
                year,
                month
            )

            filas_totales.extend(
                filas_mes
            )

        if len(
            filas_totales
        ) != 18:

            raise RuntimeError(
                "Se esperaban 18 controles "
                "(6 meses × 3 variables), "
                f"pero se obtuvieron "
                f"{len(filas_totales)}."
            )

        guardar_csv(
            filas_totales
        )

        mostrar_resumen(
            filas_totales
        )

    except Exception as error:

        titulo(
            "ERROR DEL DIAGNÓSTICO"
        )

        escribir(
            f"Tipo = "
            f"{type(error).__name__}"
        )

        escribir(
            f"Detalle = "
            f"{repr(error)}"
        )

        escribir(
            "ESTADO = "
            "DIAGNOSTICO_REDONDEO_Y_6_PILOTOS_ERROR"
        )

        if filas_totales:

            try:

                guardar_csv(
                    filas_totales
                )

            except Exception:

                pass

        raise

    finally:

        TXT_RESULTADO.write_text(
            "\n".join(
                lineas_reporte
            ),
            encoding="utf-8"
        )

        print()
        print("=" * 122)
        print("DIAGNÓSTICO FINALIZADO")
        print("=" * 122)

        print()
        print(
            "TXT:"
        )

        print(
            TXT_RESULTADO
        )

        print()
        print(
            "CSV:"
        )

        print(
            CSV_RESULTADO
        )


# ============================================================
# 31. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":

    main()