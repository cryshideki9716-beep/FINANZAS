# BRICEÑO LEON CRYSTELL HIDEKI
# Código: 2024200485D
# Tema 4: Dolarización del crédito y de los depósitos en el sistema financiero peruano
# Prueba del parser robusto por contenido - Créditos SBS

from pathlib import Path
import re

import xlrd
from openpyxl import load_workbook


# ============================================================
# 1. CARPETAS
# ============================================================

RAIZ = Path(__file__).resolve().parent.parent

CARPETA_PILOTO = (
    RAIZ
    / "datos_crudos"
    / "sbs_piloto"
)


# ============================================================
# 2. ARCHIVOS DE CONTROL
# ============================================================

ARCHIVOS_CONTROL = [
    "creditos_2008_06.xls",
    "creditos_2010_12.xls",
    "creditos_2011_12.xls",
    "creditos_2015_06.xls",
    "creditos_2020_06.xls",
    "creditos_2025_12.xls",
]


# ============================================================
# 3. NORMALIZAR TEXTO PARA COMPARACIONES
# ============================================================

def normalizar_texto(valor):
    """
    Normalización utilizada únicamente para detectar
    encabezados.

    NO modifica el valor almacenado en el archivo original.
    """

    if valor is None:
        return ""

    texto = str(valor)

    texto = texto.replace("\n", " ")
    texto = texto.replace("\r", " ")

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip().lower()


# ============================================================
# 4. DETECTAR FORMATO REAL
# ============================================================

def detectar_formato(ruta):

    with open(ruta, "rb") as archivo:
        cabecera = archivo.read(8)

    firma_ole = bytes.fromhex(
        "D0CF11E0A1B11AE1"
    )

    firmas_zip = (
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
    )

    if cabecera == firma_ole:
        return "XLS_OLE"

    if cabecera.startswith(firmas_zip):
        return "XLSX_ZIP"

    return "DESCONOCIDO"


# ============================================================
# 5. LEER XLS/OLE
# ============================================================

def leer_xls(ruta):

    libro = xlrd.open_workbook(
        filename=str(ruta),
        on_demand=True
    )

    hojas = []

    for nombre in libro.sheet_names():

        hoja = libro.sheet_by_name(nombre)

        filas = []

        for i in range(hoja.nrows):

            fila = [
                hoja.cell_value(i, j)
                for j in range(hoja.ncols)
            ]

            filas.append(fila)

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.release_resources()

    return hojas


# ============================================================
# 6. LEER XLSX/ZIP
# ============================================================

def leer_xlsx(ruta):

    archivo_binario = open(
        ruta,
        "rb"
    )

    libro = load_workbook(
        archivo_binario,
        read_only=True,
        data_only=True
    )

    hojas = []

    for nombre in libro.sheetnames:

        hoja = libro[nombre]

        filas = []

        for fila in hoja.iter_rows(
            values_only=True
        ):

            filas.append(
                list(fila)
            )

        hojas.append({
            "nombre": nombre,
            "filas": filas,
        })

    libro.close()
    archivo_binario.close()

    return hojas


# ============================================================
# 7. LEER ARCHIVO SEGÚN FORMATO REAL
# ============================================================

def leer_archivo(ruta):

    formato = detectar_formato(
        ruta
    )

    if formato == "XLS_OLE":

        hojas = leer_xls(
            ruta
        )

    elif formato == "XLSX_ZIP":

        hojas = leer_xlsx(
            ruta
        )

    else:

        raise ValueError(
            f"Formato no reconocido: {ruta.name}"
        )

    return formato, hojas


# ============================================================
# 8. RECONOCER MN
# ============================================================

def es_mn(valor):
    """
    Reconoce encabezados de Moneda Nacional.

    Ejemplos esperados:
        MN (S/. Miles)
        MN (S/ Miles)
    """

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith("mn")
        and "mil" in texto
    )


# ============================================================
# 9. RECONOCER ME
# ============================================================

def es_me(valor):
    """
    Reconoce encabezados de Moneda Extranjera.

    Ejemplo:
        ME (US$ miles)
    """

    texto = normalizar_texto(
        valor
    )

    return (
        texto.startswith("me")
        and "mil" in texto
    )


# ============================================================
# 10. RECONOCER TOTAL
# ============================================================

def es_total(valor):
    """
    Reconoce encabezados Total.
    """

    texto = normalizar_texto(
        valor
    )

    return texto.startswith(
        "total"
    )


# ============================================================
# 11. BUSCAR FILA DE ENCABEZADOS MN-ME-TOTAL
# ============================================================

def encontrar_fila_encabezados(filas):
    """
    Busca automáticamente una fila que contenga
    varios bloques consecutivos:

        MN | ME | Total

    No depende de que sea exactamente la fila 5.
    """

    candidatos = []

    for numero_fila, fila in enumerate(
        filas
    ):

        bloques = []

        for columna in range(
            len(fila) - 2
        ):

            if (
                es_mn(fila[columna])
                and es_me(fila[columna + 1])
                and es_total(fila[columna + 2])
            ):

                bloques.append(
                    columna
                )

        if bloques:

            candidatos.append({
                "fila": numero_fila,
                "bloques": bloques,
            })

    if not candidatos:

        return None

    # Elegimos la fila que tenga más bloques MN-ME-Total.
    mejor = max(
        candidatos,
        key=lambda x: len(x["bloques"])
    )

    return mejor


# ============================================================
# 12. OBTENER NOMBRE DEL BANCO
# ============================================================

def obtener_nombre_banco(
    filas,
    fila_bancos,
    columna_inicio
):
    """
    Obtiene el nombre ORIGINAL del banco situado
    encima del bloque MN-ME-Total.

    Se conserva exactamente como aparece en SBS.
    """

    if fila_bancos < 0:
        return ""

    fila = filas[
        fila_bancos
    ]

    # Primero comprobamos la columna inicial del bloque.
    if columna_inicio < len(fila):

        valor = fila[
            columna_inicio
        ]

        if normalizar_texto(valor):

            return str(
                valor
            ).strip()

    # Algunas hojas pueden representar encabezados
    # combinados de manera diferente.
    # Revisamos las otras dos columnas del bloque.
    for desplazamiento in [
        1,
        2
    ]:

        columna = (
            columna_inicio
            + desplazamiento
        )

        if columna < len(fila):

            valor = fila[
                columna
            ]

            if normalizar_texto(valor):

                return str(
                    valor
                ).strip()

    return ""


# ============================================================
# 13. DETECTAR BLOQUES DE BANCOS
# ============================================================

def detectar_bloques_bancos(
    filas
):

    resultado = encontrar_fila_encabezados(
        filas
    )

    if resultado is None:

        return None

    fila_encabezados = resultado[
        "fila"
    ]

    fila_bancos = (
        fila_encabezados - 1
    )

    bloques = []

    for columna_inicio in resultado[
        "bloques"
    ]:

        banco_original = obtener_nombre_banco(
            filas,
            fila_bancos,
            columna_inicio
        )

        bloques.append({
            "banco_original": banco_original,
            "columna_mn": columna_inicio,
            "columna_me": columna_inicio + 1,
            "columna_total": columna_inicio + 2,
        })

    return {
        "fila_bancos": fila_bancos,
        "fila_encabezados": fila_encabezados,
        "bloques": bloques,
    }


# ============================================================
# 14. ANALIZAR UN ARCHIVO
# ============================================================

def analizar_archivo(
    ruta
):

    print()
    print("=" * 90)

    print(
        "ARCHIVO:",
        ruta.name
    )

    print("=" * 90)

    formato, hojas = leer_archivo(
        ruta
    )

    print(
        "Formato real:",
        formato
    )

    print(
        "Hojas:",
        [
            hoja["nombre"]
            for hoja in hojas
        ]
    )

    candidatos = []

    # --------------------------------------------------------
    # No suponemos el nombre de la hoja.
    # Probamos TODAS las hojas.
    # --------------------------------------------------------

    for hoja in hojas:

        resultado = detectar_bloques_bancos(
            hoja["filas"]
        )

        if resultado is not None:

            candidatos.append({
                "hoja": hoja["nombre"],
                "resultado": resultado,
            })

    if not candidatos:

        print()
        print(
            "[ERROR] No se encontró ningún "
            "patrón MN-ME-Total."
        )

        return False

    # --------------------------------------------------------
    # Si hubiera más de una hoja candidata,
    # elegimos la que tenga más bloques.
    # --------------------------------------------------------

    mejor = max(
        candidatos,
        key=lambda x: len(
            x["resultado"]["bloques"]
        )
    )

    nombre_hoja = mejor[
        "hoja"
    ]

    resultado = mejor[
        "resultado"
    ]

    print()
    print(
        "Hoja seleccionada:",
        nombre_hoja
    )

    print(
        "Fila de bancos:",
        resultado["fila_bancos"]
    )

    print(
        "Fila MN/ME/Total:",
        resultado["fila_encabezados"]
    )

    print(
        "Bloques detectados:",
        len(resultado["bloques"])
    )

    print()
    print(
        "BANCOS Y COLUMNAS DETECTADAS"
    )

    print("-" * 90)

    bancos_validos = 0

    for numero, bloque in enumerate(
        resultado["bloques"],
        start=1
    ):

        banco = bloque[
            "banco_original"
        ]

        print()
        print(
            f"Bloque {numero}"
        )

        print(
            "  Banco original:",
            repr(banco)
        )

        print(
            "  Columna MN:",
            bloque["columna_mn"]
        )

        print(
            "  Columna ME:",
            bloque["columna_me"]
        )

        print(
            "  Columna Total:",
            bloque["columna_total"]
        )

        if banco:
            bancos_validos += 1

    print()
    print(
        "Bloques con nombre de banco:",
        bancos_validos,
        "de",
        len(resultado["bloques"])
    )

    if (
        bancos_validos
        != len(resultado["bloques"])
    ):

        print()
        print(
            "[ADVERTENCIA] Existen bloques "
            "sin nombre de banco."
        )

        return False

    print()
    print(
        "[OK] Estructura reconocida."
    )

    return True


# ============================================================
# 15. FUNCIÓN PRINCIPAL
# ============================================================

def main():

    print()
    print("=" * 90)

    print(
        "PRUEBA DEL PARSER DE CRÉDITOS SBS"
    )

    print(
        "PILOTO 2008-2025"
    )

    print("=" * 90)

    print()
    print(
        "Esta prueba NO descarga archivos."
    )

    print(
        "Solamente analiza los 6 XLS "
        "que ya existen en datos_crudos."
    )

    correctos = 0

    for nombre_archivo in ARCHIVOS_CONTROL:

        ruta = (
            CARPETA_PILOTO
            / nombre_archivo
        )

        if not ruta.exists():

            print()
            print(
                "[ERROR] No existe:",
                ruta
            )

            continue

        if analizar_archivo(
            ruta
        ):

            correctos += 1

    print()
    print("=" * 90)

    print(
        "RESULTADO DEL PARSER"
    )

    print("=" * 90)

    print()
    print(
        "Archivos reconocidos correctamente:",
        correctos,
        "de",
        len(ARCHIVOS_CONTROL)
    )

    if correctos == len(
        ARCHIVOS_CONTROL
    ):

        print()
        print(
            "[OK] Los seis puntos de control "
            "fueron reconocidos."
        )

        print()
        print(
            "El patrón MN-ME-Total puede "
            "detectarse por contenido."
        )

    else:

        print()
        print(
            "[ADVERTENCIA] Hay estructuras "
            "que todavía requieren revisión."
        )


# ============================================================
# 16. PUNTO DE ENTRADA
# ============================================================

if __name__ == "__main__":
    main()